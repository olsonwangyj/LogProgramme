"""Regression tests for the boundary-aware TXT-plus-log-folder axis activity analyzer."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from backend.log_axis_activity_analyzer.config import (
    CONTROL_PWM_COMMANDS,
    DEFAULT_EVENT_RULES,
    DIAGNOSTIC_SUMMARY_COLUMNS,
    DETAIL_COLUMNS,
    DURATION_STATUS_NOT_APPLICABLE,
    DURATION_STATUS_TOO_LONG,
    DURATION_STATUS_VALID,
    OVERALL_STATUS_BOUNDARY_CLOSED,
    OVERALL_STATUS_CLOSED_BY_NEW_START,
    OVERALL_STATUS_INITIALIZATION_FAILED,
    OVERALL_STATUS_OK,
    OVERALL_STATUS_PWM_WARNING,
    OVERALL_STATUS_UNMATCHED,
    PWM_STATUS_CONFLICT,
    PWM_STATUS_LATEST_BEFORE_TOO_FAR,
    PWM_STATUS_MATCHED_CARRY_FORWARD,
    PWM_STATUS_MATCHED_CONTAINING,
    PWM_STATUS_MATCHED_LATEST_BEFORE,
    PWM_STATUS_MATCHED_NEAREST,
    PWM_STATUS_MATCHED_NEAREST_FUTURE,
    PWM_STATUS_NO_EARLIER_PWM_FOR_AXIS,
    PWM_STATUS_NO_CONTROL_LOGS_AVAILABLE,
    PWM_STATUS_NO_RELEVANT_LOG_FILE,
    PWM_STATUS_NO_SAME_AXIS_IN_FOLDER,
    PWM_STATUS_RELEVANT_LOG_FILE_LACKS_AXIS_PWM,
    STATUS_CLOSED_BY_NEW_START,
    STATUS_INITIALIZATION_FAILED,
    STATUS_MATCHED,
    STATUS_CLOSED_BY_BOUNDARY,
    STATUS_DURATION_TOO_LONG_CANDIDATE,
    STATUS_PARSE_WARNING,
    STATUS_UNMATCHED_END,
    STATUS_UNMATCHED_START,
    build_control_pwm_pattern,
)
import backend.log_axis_activity_analyzer.log_b_parser as log_b_parser
from backend.log_axis_activity_analyzer.duty_cycle_associator import DutyCycleAssociator
from backend.log_axis_activity_analyzer.file_loader import TextFileLoader
from backend.log_axis_activity_analyzer.log_a_parser import MainLogParser
from backend.log_axis_activity_analyzer.log_b_parser import DutyCycleLogParser
from backend.log_axis_activity_analyzer.log_folder_scanner import LogFolderScanner
from backend.log_axis_activity_analyzer.matcher import EventMatcher
from backend.log_axis_activity_analyzer.models import ActivityRecord
from backend.log_axis_activity_analyzer.service import LogAnalysisService
from backend.log_axis_activity_analyzer.summary import SummaryGenerator
from backend.log_axis_activity_analyzer.time_utils import calculate_duration_values, parse_log_timestamp
from backend.log_axis_activity_analyzer.validation import ActivityValidator


def _write_lines(path: Path, lines: list[str]) -> Path:
    """Write helper text content with trailing newlines preserved."""

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _parse_main_log(tmp_path: Path, lines: list[str]):
    """Build a parsed main-log result from synthetic lines."""

    path = _write_lines(tmp_path / "sample.txt", lines)
    parser = MainLogParser(TextFileLoader())
    return parser.parse(path)


def test_parse_log_timestamp_preserves_milliseconds() -> None:
    """Millisecond precision should survive timestamp parsing."""

    parsed = parse_log_timestamp("2026-03-31 09:39:40:554")
    assert parsed == datetime(2026, 3, 31, 9, 39, 40, 554000)


def test_calculate_duration_values_returns_expected_result() -> None:
    """Duration must be computed as end minus start with millisecond precision."""

    start_time = datetime(2026, 3, 31, 9, 39, 40, 554000)
    end_time = datetime(2026, 3, 31, 9, 39, 47, 373000)

    duration_ms, duration_s = calculate_duration_values(start_time, end_time)

    assert duration_ms == 6819
    assert duration_s == 6.819


def test_valid_same_axis_match_produces_expected_duration(tmp_path: Path) -> None:
    """A same-axis start/end pair should match with the correct rule and duration."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-03-31 09:39:40:554 MCU   @[Y] start clearing: -19.50",
            "2026-03-31 09:39:47:373 MCU   @[Y] motor cleared",
            "2026-03-31 09:39:47:373 MCU   @[Y] min: -29.51",
            "2026-03-31 09:39:48:000 MCU   @[WRN] Node 3 (GETVER) response timeout (2016). Try again.",
        ],
    )

    matcher = EventMatcher(DEFAULT_EVENT_RULES)
    records = ActivityValidator().validate(matcher.build_activity_records(parse_result.timeline))

    matched = [record for record in records if record.match_status == STATUS_MATCHED]
    assert len(matched) == 1
    record = matched[0]
    assert record.axis == "Y"
    assert record.rule_id == "clear_motor"
    assert record.duration_ms == 6819
    assert record.duration_s == 6.819
    assert record.status == OVERALL_STATUS_OK


def test_different_axis_events_do_not_match(tmp_path: Path) -> None:
    """A start on one axis must not match an end on another axis."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-03-31 09:39:40:554 MCU   @[Y] start clearing: -19.50",
            "2026-03-31 09:39:47:373 MCU   @[Z] motor cleared",
        ],
    )

    matcher = EventMatcher(DEFAULT_EVENT_RULES)
    records = ActivityValidator().validate(matcher.build_activity_records(parse_result.timeline))

    assert any(record.match_status == STATUS_UNMATCHED_START and record.axis == "Y" for record in records)
    assert any(record.match_status == STATUS_UNMATCHED_END and record.axis == "Z" for record in records)
    assert not any(record.match_status == STATUS_MATCHED for record in records)


def test_initialization_failure_closes_pending_start(tmp_path: Path) -> None:
    """Initialization failure must close the stale start before a later successful attempt."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-04-10 08:37:43:959 MCU   @[Z] start clearing: -50.00",
            "2026-04-10 08:38:14:121 MCU   @[ERR] [Z] cannot reach the target (30 seconds timeout). initialization failed.",
            "2026-04-10 09:20:23:541 MCU   @[Z] start clearing: -50.00",
            "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared",
            "2026-04-10 09:20:46:317 MCU   @[Z] min: -50.00",
        ],
    )

    matcher = EventMatcher(DEFAULT_EVENT_RULES)
    records = ActivityValidator().validate(matcher.build_activity_records(parse_result.timeline))

    stale_rows = [
        record
        for record in records
        if record.axis == "Z" and record.start_time == datetime(2026, 4, 10, 8, 37, 43, 959000)
    ]
    assert len(stale_rows) == 1
    assert stale_rows[0].match_status == STATUS_INITIALIZATION_FAILED
    assert stale_rows[0].end_time is None

    corrected = [
        record
        for record in records
        if record.axis == "Z"
        and record.start_time == datetime(2026, 4, 10, 9, 20, 23, 541000)
        and record.end_time == datetime(2026, 4, 10, 9, 20, 46, 317000)
    ]
    assert len(corrected) == 1
    assert corrected[0].match_status == STATUS_MATCHED
    assert corrected[0].duration_ms == 22776
    assert corrected[0].duration_s == 22.776

    assert not any(
        record.match_status == STATUS_MATCHED
        and record.axis == "Z"
        and record.start_time == datetime(2026, 4, 10, 8, 37, 43, 959000)
        and record.end_time == datetime(2026, 4, 10, 9, 20, 46, 317000)
        for record in records
    )


def test_main_log_diagnostics_are_not_parse_warnings(tmp_path: Path) -> None:
    """Known ERR/WRN diagnostic lines should be structured diagnostics, not malformed rows."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-04-10 09:15:50:471 MCU   @[WRN] Node 3 (GETVER) response timeout (2016). Try again.",
            "2026-04-10 09:15:52:494 MCU   @[ERR] Node 3 (GETVER) response timeout (2015). Check connection.",
            "2026-04-10 09:05:44:965 MCU   @[ERR] motor Z RIGHT sensor cut",
        ],
    )

    assert len(parse_result.warnings) == 0
    assert [event.diagnostic_type for event in parse_result.diagnostics] == [
        "NodeResponseTimeout",
        "NodeResponseTimeout",
        "SensorCut",
    ]
    assert parse_result.diagnostics[0].severity == "WRN"
    assert parse_result.diagnostics[0].node_id == 3
    assert parse_result.diagnostics[1].severity == "ERR"
    assert parse_result.diagnostics[2].axis == "Z"
    assert parse_result.diagnostics[2].flush_pending is True


def test_amd_amx_partially_corrupted_is_diagnostic_not_parse_warning(tmp_path: Path) -> None:
    """AMD AMX partially corrupted lines should be structured diagnostics."""

    raw_line = "2026-04-23 12:34:39:688 MCU   @[AMD] AMX partially corrupted [ 25 A5 0C 01 31 05 E6 3D 77 00 ]"
    parse_result = _parse_main_log(tmp_path, [raw_line])

    assert len(parse_result.warnings) == 0
    assert len(parse_result.diagnostics) == 1
    diagnostic = parse_result.diagnostics[0]
    assert diagnostic.severity == "AMD"
    assert diagnostic.diagnostic_type == "AMXPartiallyCorrupted"
    assert diagnostic.axis is None
    assert diagnostic.node_id is None
    assert diagnostic.message == "MCU   @[AMD] AMX partially corrupted [ 25 A5 0C 01 31 05 E6 3D 77 00 ]"
    assert diagnostic.raw_line == raw_line
    assert diagnostic.flush_pending is False


def test_amd_partial_amx_amended_maps_node_to_axis_and_is_not_parse_warning(tmp_path: Path) -> None:
    """AMD partial AMX amended lines should capture node IDs and optional axis mapping."""

    raw_line = (
        "2026-04-23 12:34:39:711 MCU   @[AMD] [N12] partial AMX amended "
        "[ 25 A5 0C 01 31 05 E6 3D 77 00 00 DD C0 ]"
    )
    parse_result = _parse_main_log(tmp_path, [raw_line])

    assert len(parse_result.warnings) == 0
    assert len(parse_result.diagnostics) == 1
    diagnostic = parse_result.diagnostics[0]
    assert diagnostic.severity == "AMD"
    assert diagnostic.diagnostic_type == "PartialAMXAmended"
    assert diagnostic.node_id == 12
    assert diagnostic.axis == "Z"
    assert diagnostic.raw_line == raw_line
    assert diagnostic.flush_pending is False


def test_initialization_failed_remains_boundary_not_diagnostic(tmp_path: Path) -> None:
    """Initialization failure should still be a flushing boundary event."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-04-10 08:38:14:121 MCU   @[ERR] [Z] cannot reach the target (30 seconds timeout). initialization failed.",
        ],
    )

    assert len(parse_result.boundary_events) == 1
    assert parse_result.boundary_events[0].boundary_type == "InitializationFailed"
    assert parse_result.boundary_events[0].axis == "Z"
    assert len(parse_result.diagnostics) == 0
    assert len(parse_result.warnings) == 0


def test_new_initialization_closes_old_pending_start(tmp_path: Path) -> None:
    """A new initialization boundary must flush any pending starts from the prior attempt."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-04-10 08:37:43:959 MCU   @[Z] start clearing: -50.00",
            "2026-04-10 08:38:51:254 User  @on_startInitRobot_clicked",
            "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared",
        ],
    )

    matcher = EventMatcher(DEFAULT_EVENT_RULES)
    records = ActivityValidator().validate(matcher.build_activity_records(parse_result.timeline))

    assert any(
        record.match_status == STATUS_CLOSED_BY_BOUNDARY
        and record.axis == "Z"
        and record.start_time == datetime(2026, 4, 10, 8, 37, 43, 959000)
        for record in records
    )
    assert any(
        record.match_status == STATUS_UNMATCHED_END
        and record.axis == "Z"
        and record.end_time == datetime(2026, 4, 10, 9, 20, 46, 317000)
        for record in records
    )


def test_abort_stop_and_exit_boundaries_close_pending_start(tmp_path: Path) -> None:
    """Workflow interruption boundaries should close pending starts immediately."""

    cases = [
        ("2026-04-10 09:16:55:572 User  @on_functionAbortMotor_clicked", "MotorAbortClicked"),
        ("2026-04-10 09:16:55:586 MCU   @robot moving stopped", "RobotMovingStopped"),
        ("2026-04-10 09:16:58:200 User  @system menu selected (Exit)", "SystemExitSelected"),
    ]

    for boundary_line, boundary_type in cases:
        parse_result = _parse_main_log(
            tmp_path,
            [
                "2026-04-10 09:16:54:000 MCU   @[Z] start clearing: -50.00",
                boundary_line,
            ],
        )
        records = ActivityValidator().validate(EventMatcher(DEFAULT_EVENT_RULES).build_activity_records(parse_result.timeline))

        closed = next(record for record in records if record.axis == "Z")
        assert closed.match_status == STATUS_CLOSED_BY_BOUNDARY
        assert closed.duration_status == DURATION_STATUS_NOT_APPLICABLE
        assert closed.closed_by_boundary_type == boundary_type
        assert closed.closed_by_boundary_line_text == boundary_line
        assert closed.status == OVERALL_STATUS_BOUNDARY_CLOSED


def test_long_duration_is_rejected(tmp_path: Path) -> None:
    """A long-gap candidate pair must not survive as a normal matched row."""

    parse_result = _parse_main_log(
        tmp_path,
        [
            "2026-04-10 08:37:43:959 MCU   @[Z] start clearing: -50.00",
            "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared",
        ],
    )

    matcher = EventMatcher(DEFAULT_EVENT_RULES)
    records = ActivityValidator().validate(matcher.build_activity_records(parse_result.timeline))

    assert not any(record.match_status == STATUS_MATCHED for record in records)
    rejected = next(record for record in records if record.match_status == STATUS_DURATION_TOO_LONG_CANDIDATE)
    assert rejected.duration_status == DURATION_STATUS_TOO_LONG
    assert rejected.candidate_duration_ms == 2582358
    assert rejected.candidate_duration_s == 2582.358
    assert rejected.duration_ms is None
    assert not any(record.duration_ms == 2582358 and record.match_status == STATUS_MATCHED for record in records)


def test_pwm_signed_value_is_normalized(tmp_path: Path) -> None:
    """Signed PWM values should preserve raw sign while normalizing the percent magnitude."""

    path = _write_lines(
        tmp_path / "signed.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] RUN 255 176 (-80)",
            "2026-04-10 08:37:33:000 [IN ] sample",
            "                              [N6:H] RUN 'S' 255 176 (-80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)

    assert len(result.pwm_events) == 2
    event = result.pwm_events[0]
    assert event.axis == "H"
    assert event.pwm_raw_value == -80
    assert event.pwm_percent == 80
    assert event.direction == "Reverse"
    assert event.is_confirmed_by_status_line is True


def test_pwm_command_types_are_configurable_defaults(tmp_path: Path) -> None:
    """Default PWM command parsing should be driven by the configured command set."""

    assert {"RUN", "VEL"} <= CONTROL_PWM_COMMANDS
    path = _write_lines(
        tmp_path / "commands.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] RUN 255 176 (-80)",
            "2026-04-10 08:37:33:000 [OUT] sample",
            "                              [N6:H] VEL 0 80 (80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)

    assert [event.command_type for event in result.pwm_events] == ["RUN", "VEL"]


def test_pwm_command_type_extension_parses_without_parser_rewrite(tmp_path: Path, monkeypatch) -> None:
    """Extending CONTROL_PWM_COMMANDS should be enough for new PWM command names."""

    monkeypatch.setattr(log_b_parser, "CONTROL_PWM_COMMANDS", {"RUN", "VEL", "PWM"})
    monkeypatch.setattr(log_b_parser, "CONTROL_PWM_PATTERN", build_control_pwm_pattern({"RUN", "VEL", "PWM"}))
    path = _write_lines(
        tmp_path / "extended-command.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] PWM 0 80 (80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)

    assert len(result.pwm_events) == 1
    assert result.pwm_events[0].command_type == "PWM"
    assert result.pwm_events[0].pwm_percent == 80


def test_unsupported_pwm_command_is_ignored_without_warning(tmp_path: Path) -> None:
    """Unknown command names should not be treated as malformed configured PWM commands."""

    path = _write_lines(
        tmp_path / "unsupported-command.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] SPEED 0 80 (80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)

    assert result.pwm_events == []
    assert result.warnings == []


def test_pwm_direction_change_is_not_a_conflict(tmp_path: Path) -> None:
    """Raw -80 and 80 should be one normalized PWM profile with direction-change metadata."""

    path = _write_lines(
        tmp_path / "direction-change.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] RUN 255 176 (-80)",
            "2026-04-10 08:37:33:000 [IN ] sample",
            "                              [N6:H] RUN 'S' 255 176 (-80)",
            "2026-04-10 08:37:34:000 [OUT] sample",
            "                              [N6:H] VEL 0 80 (80)",
            "2026-04-10 08:37:35:000 [IN ] sample",
            "                              [N6:H] VEL 'S' 0 80 (80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)
    profile = result.axis_profiles["H"]

    assert profile.pwm_percent == 80
    assert profile.pwm_percent_values == [80]
    assert profile.pwm_raw_values == [-80, 80]
    assert profile.direction_values == ["Forward", "Reverse"]
    assert profile.conflict is False
    assert profile.direction_changed is True
    assert "Direction changed" in profile.notes


def test_pwm_percent_change_is_a_conflict(tmp_path: Path) -> None:
    """Different normalized PWM percentages in one source file should remain a conflict."""

    path = _write_lines(
        tmp_path / "percent-conflict.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N4:Y] RUN 0 60 (60)",
            "2026-04-10 08:37:33:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(path)
    profile = result.axis_profiles["Y"]

    assert profile.pwm_percent_values == [60, 80]
    assert profile.conflict is True
    assert profile.conflict_reason == "Conflicting normalized PWM percentages in file for axis Y: 60, 80"


def test_direction_changed_profile_does_not_create_pwm_warning(tmp_path: Path) -> None:
    """Direction changes should keep the real PWM match status and stay non-warning."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "direction-change.log",
        [
            "2026-04-10 08:37:32:000 [OUT] sample",
            "                              [N6:H] RUN 255 176 (-80)",
            "2026-04-10 08:37:33:000 [OUT] sample",
            "                              [N6:H] VEL 0 80 (80)",
            "2026-04-10 08:37:40:000 [OUT] sample",
        ],
    )
    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="H",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 4, 10, 8, 37, 35),
        end_time=datetime(2026, 4, 10, 8, 37, 36),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_percent == 80
    assert record.pwm_match_status == PWM_STATUS_MATCHED_CONTAINING
    assert record.pwm_direction_changed is True
    assert record.pwm_conflict is False
    assert record.pwm_warning is False
    assert record.status == OVERALL_STATUS_OK
    assert "Direction changed within source file" in record.notes


def test_summary_uses_normalized_pwm_percent_for_mode() -> None:
    """Summary PWM mode should not split signed raw values into separate duty categories."""

    records = ActivityValidator().validate(
        [
            ActivityRecord(
                axis="H",
                rule_id="clear_motor",
                event_type="start clearing",
                start_event="start clearing",
                end_event="motor cleared",
                start_time=datetime(2026, 1, 1, 0, 0, 0),
                end_time=datetime(2026, 1, 1, 0, 0, 1),
                match_status=STATUS_MATCHED,
                pwm_percent=80,
                pwm_raw_value=-80,
            ),
            ActivityRecord(
                axis="H",
                rule_id="clear_motor",
                event_type="start clearing",
                start_event="start clearing",
                end_event="motor cleared",
                start_time=datetime(2026, 1, 1, 0, 0, 2),
                end_time=datetime(2026, 1, 1, 0, 0, 3),
                match_status=STATUS_MATCHED,
                pwm_percent=80,
                pwm_raw_value=80,
            ),
        ]
    )

    frames = SummaryGenerator().build_report_frames(records, [])

    assert frames.axis_summary.iloc[0]["Most Common PWM (%)"] == 80


def test_log_folder_scanner_keeps_per_file_pwm_profiles(tmp_path: Path) -> None:
    """Different files should retain separate PWM profiles for the same axis."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "one.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 60 (60)",
            "2026-01-01 00:00:02:000 [IN ] sample",
            "                              [N4:Y] RUN 'S' 0 60 (60)",
        ],
    )
    _write_lines(
        folder / "two.log",
        [
            "2026-01-01 00:05:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:05:02:000 [IN ] sample",
            "                              [N4:Y] RUN 'S' 0 80 (80)",
        ],
    )

    scanner = LogFolderScanner(DutyCycleLogParser(TextFileLoader()))
    result = scanner.scan(folder)

    assert len(result.files) == 2
    file_profiles = {file_result.source_path.name: file_result.axis_profiles["Y"] for file_result in result.files}
    assert file_profiles["one.log"].pwm_percent == 60
    assert file_profiles["two.log"].pwm_percent == 80
    assert file_profiles["one.log"].pwm_raw_values == [60]
    assert file_profiles["two.log"].pwm_raw_values == [80]


def test_pwm_association_prefers_containing_log_file(tmp_path: Path) -> None:
    """PWM association should choose the relevant containing log file, not a faraway global value."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "near.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 60 (60)",
            "2026-01-01 00:00:02:000 [IN ] sample",
            "                              [N4:Y] RUN 'S' 0 60 (60)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    _write_lines(
        folder / "far.log",
        [
            "2026-01-01 00:10:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:10:02:000 [IN ] sample",
            "                              [N4:Y] RUN 'S' 0 80 (80)",
            "2026-01-01 00:10:10:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_percent == 60
    assert record.pwm_source_file == "near.log"
    assert record.pwm_match_status == PWM_STATUS_MATCHED_CONTAINING


def test_pwm_association_falls_back_when_containing_file_lacks_axis(tmp_path: Path) -> None:
    """A containing file without the target axis should not block a nearby same-axis profile."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "contains-without-y.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N1:X] RUN 0 50 (50)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    _write_lines(
        folder / "near-y.log",
        [
            "2026-01-01 00:00:11:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:00:12:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_percent == 80
    assert record.pwm_source_file == "near-y.log"
    assert record.pwm_match_status == PWM_STATUS_MATCHED_NEAREST_FUTURE
    assert record.pwm_warning is True
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_pwm_association_does_not_attach_far_latest_before(tmp_path: Path) -> None:
    """Latest-before fallback should stay blank when the same-axis source is too old."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "old-y.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:00:02:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 10, 0),
        end_time=datetime(2026, 1, 1, 0, 10, 1),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_percent is None
    assert record.pwm_source_file == ""
    assert record.pwm_match_status == PWM_STATUS_LATEST_BEFORE_TOO_FAR
    assert record.pwm_missing_reason.startswith("The latest PWM record before the activity for axis Y was too far away")
    assert record.pwm_time_delta_ms == 599000
    assert record.pwm_warning is True
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_latest_before_strategy_does_not_attach_future_pwm(tmp_path: Path) -> None:
    """Latest-before semantics must never relabel a future PWM as prior evidence."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "future-y.log",
        [
            "2026-01-01 00:05:00:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:05:01:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        end_time=datetime(2026, 1, 1, 0, 0, 1),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files, strategy="latest_before_start"))

    assert record.pwm_percent is None
    assert record.pwm_source_file == ""
    assert record.pwm_match_status == PWM_STATUS_NO_EARLIER_PWM_FOR_AXIS
    assert record.pwm_warning is True


def test_latest_known_strategy_uses_only_prior_pwm(tmp_path: Path) -> None:
    """Latest-known should select earlier PWM and ignore future-only evidence."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "past-y.log",
        [
            "2026-01-01 00:00:59:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:01:00:000 [OUT] sample",
        ],
    )
    _write_lines(
        folder / "future-y.log",
        [
            "2026-01-01 00:05:00:000 [OUT] sample",
            "                              [N4:Y] RUN 0 60 (60)",
            "2026-01-01 00:05:01:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 1, 30),
        end_time=datetime(2026, 1, 1, 0, 1, 31),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files, strategy="latest_known"))

    assert record.pwm_percent == 80
    assert record.pwm_source_file == "past-y.log"
    assert record.pwm_match_status == PWM_STATUS_MATCHED_LATEST_BEFORE
    assert record.pwm_time_delta_ms == 31000


def test_pwm_carry_forward_is_explicit_and_warned_when_enabled(tmp_path: Path) -> None:
    """Carry-forward PWM is opt-in and clearly marked as a warning source."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "old-y.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:00:02:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 10, 0),
        end_time=datetime(2026, 1, 1, 0, 10, 1),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files, allow_carry_forward=True))

    assert record.pwm_percent == 80
    assert record.pwm_match_status == PWM_STATUS_MATCHED_CARRY_FORWARD
    assert record.pwm_match_method == "CarryForwardWithinSession"
    assert record.pwm_warning is True
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_pwm_association_reports_no_control_logs_available() -> None:
    """Missing control-log coverage should be distinguishable from axis-specific misses."""

    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], []))

    assert record.pwm_match_status == PWM_STATUS_NO_CONTROL_LOGS_AVAILABLE
    assert record.pwm_match_method == "NoControlLogsAvailable"
    assert record.pwm_missing_reason == "No control-log files were available for PWM matching."
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_pwm_association_reports_no_same_axis_pwm_in_folder(tmp_path: Path) -> None:
    """A folder with control logs but no target-axis PWM should say so directly."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "x-only.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N3:X] RUN 0 60 (60)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_match_status == PWM_STATUS_NO_SAME_AXIS_IN_FOLDER
    assert record.pwm_match_method == "NoSameAxisPWMInFolder"
    assert record.pwm_missing_reason == "No PWM profile was found for axis Y in the scanned control-log folder."
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_pwm_association_reports_containing_file_lacks_axis_when_no_safe_fallback(tmp_path: Path) -> None:
    """A containing log without the target axis should be distinguishable when fallback fails."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "contains-without-y.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N3:X] RUN 0 50 (50)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    _write_lines(
        folder / "future-y.log",
        [
            "2026-01-01 00:20:00:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:20:01:000 [OUT] sample",
        ],
    )
    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_match_status == PWM_STATUS_RELEVANT_LOG_FILE_LACKS_AXIS_PWM
    assert record.pwm_missing_reason.startswith("1 containing log file(s) lacked axis Y")
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_pwm_association_reports_normalized_conflict(tmp_path: Path) -> None:
    """Only normalized percent conflicts should raise PWMConflictInSourceFile."""

    folder = tmp_path / "logs"
    folder.mkdir()
    _write_lines(
        folder / "conflict-y.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N4:Y] RUN 0 60 (60)",
            "2026-01-01 00:00:02:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )

    files = LogFolderScanner(DutyCycleLogParser(TextFileLoader())).scan(folder).files
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 0, 0, 5),
        end_time=datetime(2026, 1, 1, 0, 0, 6),
        match_status=STATUS_MATCHED,
    )

    ActivityValidator().validate(DutyCycleAssociator().attach([record], files))

    assert record.pwm_match_status == PWM_STATUS_CONFLICT
    assert record.pwm_conflict is True
    assert record.pwm_warning is True
    assert record.status == OVERALL_STATUS_PWM_WARNING


def test_overall_status_keeps_match_and_duration_status_separate() -> None:
    """PWM warnings should not overwrite match or duration correctness."""

    matched_with_pwm_warning = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        end_time=datetime(2026, 1, 1, 0, 0, 1),
        match_status=STATUS_MATCHED,
        pwm_match_status=PWM_STATUS_NO_RELEVANT_LOG_FILE,
    )
    matched_clean = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        end_time=datetime(2026, 1, 1, 0, 0, 1),
        match_status=STATUS_MATCHED,
    )
    boundary_closed = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        match_status=STATUS_CLOSED_BY_BOUNDARY,
        duration_status=DURATION_STATUS_NOT_APPLICABLE,
        pwm_match_status=PWM_STATUS_NO_CONTROL_LOGS_AVAILABLE,
    )
    unmatched_start = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        match_status=STATUS_UNMATCHED_START,
        duration_status=DURATION_STATUS_NOT_APPLICABLE,
        pwm_match_status=PWM_STATUS_NO_RELEVANT_LOG_FILE,
    )
    unmatched_end = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        end_time=datetime(2026, 1, 1, 0, 0, 1),
        match_status=STATUS_UNMATCHED_END,
        duration_status=DURATION_STATUS_NOT_APPLICABLE,
        pwm_match_status=PWM_STATUS_NO_RELEVANT_LOG_FILE,
    )
    closed_by_new_start = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        match_status=STATUS_CLOSED_BY_NEW_START,
        duration_status=DURATION_STATUS_NOT_APPLICABLE,
        pwm_match_status=PWM_STATUS_NO_RELEVANT_LOG_FILE,
    )
    initialization_failed = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        start_time=datetime(2026, 1, 1, 0, 0, 0),
        match_status=STATUS_INITIALIZATION_FAILED,
        duration_status=DURATION_STATUS_NOT_APPLICABLE,
        pwm_match_status=PWM_STATUS_NO_RELEVANT_LOG_FILE,
    )

    ActivityValidator().validate(
        [
            matched_with_pwm_warning,
            matched_clean,
            boundary_closed,
            unmatched_start,
            unmatched_end,
            closed_by_new_start,
            initialization_failed,
        ]
    )

    assert matched_with_pwm_warning.match_status == STATUS_MATCHED
    assert matched_with_pwm_warning.duration_status == DURATION_STATUS_VALID
    assert matched_with_pwm_warning.status == OVERALL_STATUS_PWM_WARNING
    assert matched_clean.match_status == STATUS_MATCHED
    assert matched_clean.duration_status == DURATION_STATUS_VALID
    assert matched_clean.status == OVERALL_STATUS_OK
    assert boundary_closed.match_status == STATUS_CLOSED_BY_BOUNDARY
    assert boundary_closed.duration_status == DURATION_STATUS_NOT_APPLICABLE
    assert boundary_closed.status == OVERALL_STATUS_BOUNDARY_CLOSED
    assert unmatched_start.status == OVERALL_STATUS_UNMATCHED
    assert unmatched_end.status == OVERALL_STATUS_UNMATCHED
    assert closed_by_new_start.status == OVERALL_STATUS_CLOSED_BY_NEW_START
    assert initialization_failed.status == OVERALL_STATUS_INITIALIZATION_FAILED


def test_closed_by_new_start_is_counted_in_summary_and_run_result(tmp_path: Path) -> None:
    """A replaced pending start should have its own summary and run-result count."""

    txt_path = _write_lines(
        tmp_path / "sample.txt",
        [
            "2026-01-01 00:00:01:000 MCU   @[Z] start clearing: -50.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] start clearing: -50.00",
            "2026-01-01 00:00:03:000 MCU   @[Z] motor cleared",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "analysis.xlsx"

    result = LogAnalysisService().run_analysis(txt_path, log_folder, output_path)

    assert result.closed_by_new_start_count == 1
    workbook = load_workbook(output_path, data_only=True)
    details_sheet = workbook["Details"]
    details_headers = [cell.value for cell in details_sheet[1]]
    details_index = {header: index for index, header in enumerate(details_headers)}
    detail_rows = list(details_sheet.iter_rows(min_row=2, values_only=True))
    assert any(
        row[details_index["Match Status"]] == STATUS_CLOSED_BY_NEW_START
        and row[details_index["Overall Status"]] == OVERALL_STATUS_CLOSED_BY_NEW_START
        for row in detail_rows
    )
    assert any(
        row[details_index["Match Status"]] == STATUS_MATCHED
        and row[details_index["Duration (ms)"]] == 1000
        for row in detail_rows
    )

    summary_sheet = workbook["Summary"]
    metadata = {
        row[0]: row[1]
        for row in summary_sheet.iter_rows(min_row=1, max_col=2, values_only=True)
        if row[0]
    }
    assert metadata["Closed By New Start Count"] == 1

    summary_rows = list(summary_sheet.iter_rows(values_only=True))
    header_row = next(row for row in summary_rows if row and row[0] == "Axis")
    header_index = {header: index for index, header in enumerate(header_row)}
    z_summary = next(row for row in summary_rows if row and row[0] == "Z" and row[1] == "start clearing")
    assert z_summary[header_index["Matched Count"]] == 1
    assert z_summary[header_index["Closed By New Start Count"]] == 1


def test_excel_export_validation_with_synthetic_service_run(tmp_path: Path) -> None:
    """A synthetic end-to-end run should produce workbook sheets and numeric duration values."""

    txt_path = _write_lines(
        tmp_path / "sample.txt",
        [
            "2026-03-31 09:39:40:554 MCU   @[Y] start clearing: -19.50",
            "2026-03-31 09:39:47:373 MCU   @[Y] motor cleared",
            "2026-03-31 09:39:47:373 MCU   @[Y] min: -29.51",
            "2026-03-31 09:39:48:000 MCU   @[WRN] Node 3 (GETVER) response timeout (2016). Try again.",
            "2026-03-31 09:39:49:000 MCU   @[AMD] AMX partially corrupted [ 25 A5 0C 01 31 05 E6 3D 77 00 ]",
            "2026-03-31 09:39:49:100 MCU   @[AMD] [N12] partial AMX amended [ 25 A5 0C 01 31 05 E6 3D 77 00 00 DD C0 ]",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "initialization.log",
        [
            "2026-03-31 09:39:32:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-03-31 09:39:33:000 [IN ] sample",
            "                              [N4:Y] RUN 'S' 0 80 (80)",
            "2026-03-31 09:39:34:000 [OUT] sample",
            "                              [N6:H] RUN 255 176 (-80)",
            "2026-03-31 09:39:35:000 [IN ] sample",
            "                              [N6:H] RUN 'S' 255 176 (-80)",
            "2026-03-31 09:39:50:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "analysis.xlsx"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
    )

    assert result.detail_count > 0
    assert result.diagnostic_count == 3
    assert result.parse_warning_count == 0
    workbook = load_workbook(output_path, data_only=True)
    assert {"Details", "Summary", "PWM Sources", "Diagnostics", "Diagnostics Summary"} <= set(workbook.sheetnames)

    details_sheet = workbook["Details"]
    headers = [cell.value for cell in details_sheet[1]]
    header_index = {header: index for index, header in enumerate(headers)}
    assert headers == DETAIL_COLUMNS

    rows = list(details_sheet.iter_rows(min_row=2, values_only=True))
    assert rows
    target_row = next(
        row
        for row in rows
        if row[header_index["Axis"]] == "Y"
        and row[header_index["Start Event"]] == "start clearing"
        and row[header_index["End Event"]] == "motor cleared"
    )
    assert isinstance(target_row[header_index["Duration (ms)"]], int)
    assert target_row[header_index["Duration (ms)"]] == 6819
    assert target_row[header_index["Duration (s)"]] == 6.819
    assert target_row[header_index["PWM (%)"]] == 80
    assert target_row[header_index["PWM Raw Value"]] == 80
    assert target_row[header_index["PWM Source File"]] == "initialization.log"
    assert target_row[header_index["PWM Match Status"]] == PWM_STATUS_MATCHED_CONTAINING
    assert isinstance(target_row[header_index["PWM Source Line Number"]], int)
    assert "RUN" in target_row[header_index["PWM Source Line Text"]]
    assert target_row[header_index["Source TXT Start Line Number"]] == 1
    assert "start clearing" in target_row[header_index["Source TXT Start Line Text"]]
    assert target_row[header_index["Source TXT End Line Number"]] == 2
    assert "motor cleared" in target_row[header_index["Source TXT End Line Text"]]
    assert target_row[header_index["Overall Status"]] == OVERALL_STATUS_OK
    assert not any(
        row[header_index["Match Status"]] == STATUS_PARSE_WARNING
        and (
            "AMX partially corrupted" in str(row[header_index["Source TXT Start Line Text"]])
            or "partial AMX amended" in str(row[header_index["Source TXT Start Line Text"]])
        )
        for row in rows
    )

    summary_sheet = workbook["Summary"]
    summary_header_rows = [
        row
        for row in summary_sheet.iter_rows(values_only=True)
        if row and row[0] == "Axis"
    ]
    assert summary_header_rows
    assert all("Diagnostic Count" not in row for row in summary_header_rows)

    pwm_sources_sheet = workbook["PWM Sources"]
    pwm_rows = list(pwm_sources_sheet.iter_rows(min_row=2, values_only=True))
    negative_row = next(row for row in pwm_rows if row[3] == "H")
    assert negative_row[4] == 80
    assert negative_row[5] == -80
    assert negative_row[6] == "Reverse"

    diagnostics_sheet = workbook["Diagnostics"]
    diagnostic_rows = list(diagnostics_sheet.iter_rows(min_row=2, values_only=True))
    assert len(diagnostic_rows) == 3
    diagnostics_by_type = {row[2]: row for row in diagnostic_rows}
    assert diagnostics_by_type["NodeResponseTimeout"][1] == "WRN"
    assert diagnostics_by_type["NodeResponseTimeout"][4] == 3
    assert diagnostics_by_type["AMXPartiallyCorrupted"][1] == "AMD"
    assert "AMX partially corrupted" in diagnostics_by_type["AMXPartiallyCorrupted"][6]
    assert diagnostics_by_type["PartialAMXAmended"][1] == "AMD"
    assert diagnostics_by_type["PartialAMXAmended"][3] == "Z"
    assert diagnostics_by_type["PartialAMXAmended"][4] == 12

    diagnostics_summary_sheet = workbook["Diagnostics Summary"]
    diagnostics_summary_headers = [cell.value for cell in diagnostics_summary_sheet[1]]
    assert diagnostics_summary_headers == DIAGNOSTIC_SUMMARY_COLUMNS
    diagnostics_summary_rows = list(diagnostics_summary_sheet.iter_rows(min_row=2, values_only=True))
    summary_by_type = {row[0]: row for row in diagnostics_summary_rows}
    assert summary_by_type["AMXPartiallyCorrupted"][1] == "AMD"
    assert summary_by_type["AMXPartiallyCorrupted"][4] == 1
    assert summary_by_type["PartialAMXAmended"][1] == "AMD"
    assert summary_by_type["PartialAMXAmended"][2] == "Z"
    assert summary_by_type["PartialAMXAmended"][3] == 12
    assert summary_by_type["PartialAMXAmended"][4] == 1
