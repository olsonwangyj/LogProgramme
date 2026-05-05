"""Tests for normal distribution analysis, charts, and workbook export."""

from __future__ import annotations

import math
import shutil
import subprocess
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.log_axis_activity_analyzer.chart_generator import NormalDistributionChartGenerator
from backend.log_axis_activity_analyzer.config import (
    AXIS_RAW_SCALE,
    DEFAULT_EVENT_RULES,
    DISTRIBUTION_CHART_BLOCK_HEIGHT,
    DISTRIBUTION_IMAGE_INDEX_COLUMNS,
    DISTRIBUTION_IMAGE_STATISTICS_COLUMNS,
    DISTRIBUTION_SUMMARY_COLUMNS,
    DURATION_STATUS_VALID,
    HARDWARE_STATUS_MATCHED_NEAREST_FUTURE,
    HARDWARE_STATUS_MATCHED_OVERLAP,
    HARDWARE_STATUS_MULTIPLE_CANDIDATES,
    HARDWARE_STATUS_NO_SEGMENT,
    HARDWARE_STATUS_REFERENCE_NOT_APPLICABLE,
    HARDWARE_REFERENCE_STATUS_FOUND,
    HARDWARE_REFERENCE_STATUS_NO_EVIDENCE,
    HARDWARE_REFERENCE_STATUS_NOT_APPLICABLE,
    LOG_COVERAGE_GAPS_COLUMNS,
    LOG_COVERAGE_SUMMARY_COLUMNS,
    OVERALL_STATUS_HARDWARE_WARNING,
    OVERALL_STATUS_OK,
    PWM_STATUS_LATEST_BEFORE_TOO_FAR,
    PWM_STATUS_MATCHED_CONTAINING,
    PWM_STATUS_MATCHED_NEAREST,
    PWM_STATUS_MATCHED_NEAREST_FUTURE,
    STATUS_MATCHED,
)
from backend.log_axis_activity_analyzer.distribution import DistributionAnalyzer, ReferenceDurationAnalyzer
from backend.log_axis_activity_analyzer.excel_exporter import ExcelExporter
from backend.log_axis_activity_analyzer.file_loader import TextFileLoader
from backend.log_axis_activity_analyzer.hardware_motion_associator import HardwareMotionAssociator
from backend.log_axis_activity_analyzer.image_gallery_exporter import DistributionImageGalleryExporter
from backend.log_axis_activity_analyzer.log_a_parser import MainLogParser
from backend.log_axis_activity_analyzer.log_b_parser import DutyCycleLogParser
from backend.log_axis_activity_analyzer.matcher import EventMatcher
from backend.log_axis_activity_analyzer.models import ActivityRecord, DutyCycleLogFileResult, HardwareMotionSegment
from backend.log_axis_activity_analyzer.service import LogAnalysisService
from backend.log_axis_activity_analyzer.summary import SummaryGenerator
from backend.log_axis_activity_analyzer.validation import ActivityValidator


def _write_lines(path: Path, lines: list[str]) -> Path:
    """Write helper text content with trailing newlines preserved."""

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


_NODE_BY_AXIS = {
    "X": "N3",
    "Y": "N4",
    "V": "N5",
    "H": "N6",
    "N": "N7",
    "R": "N8",
    "P": "N9",
    "Z": "N12",
}


def _control_timestamp(value: datetime) -> str:
    """Format a control-log timestamp line."""

    return f"{value.strftime('%Y-%m-%d %H:%M:%S:%f')[:-3]} [OUT] synthetic"


def _tpos_lines(axis: str, start_time: datetime, end_time: datetime, distance: float) -> list[str]:
    """Build a simple hardware TPOS Start/End pair for synthetic tests."""

    raw_end = int(round(abs(distance) * abs(AXIS_RAW_SCALE[axis])))
    node = _NODE_BY_AXIS[axis]
    sign = -1 if axis in {"X", "Y", "Z", "N", "P", "R"} else 1
    raw_end *= sign
    return [
        _control_timestamp(start_time),
        f"                              [{node}:{axis}] TPOS 'S' 0 0 0 0 (0)",
        _control_timestamp(end_time),
        f"                              [{node}:{axis}] TPOS 'E' 0 0 0 0 ({raw_end})",
    ]


def _record(
    *,
    seconds: float = 10.0,
    txt_offset: int = 0,
    axis: str = "Z",
    pwm_percent: float = 80,
    start_value: float | None = -50.0,
    end_value: float | None = None,
    rule_id: str = "clear_motor",
    start_event: str = "start clearing",
    end_event: str = "motor cleared",
    movement_start_position: float | None = 0.0,
    movement_target_position: float | None = None,
    movement_end_position: float | None = None,
    movement_commanded_distance: float | None = None,
    movement_actual_distance: float | None = None,
    movement_distance: float | None = None,
    movement_distance_source: str = "CommandTargetPosition",
    movement_distance_method: str = "KnownStartPositionToTarget",
    pwm_match_status: str = PWM_STATUS_MATCHED_CONTAINING,
) -> ActivityRecord:
    """Build a valid matched activity record for distribution tests."""

    start_time = datetime(2026, 1, 1, 0, 0, 0) + timedelta(seconds=txt_offset)
    end_time = start_time + timedelta(seconds=seconds)
    target_position = movement_target_position if movement_target_position is not None else start_value
    if movement_commanded_distance is None and movement_start_position is not None and target_position is not None:
        movement_commanded_distance = abs(target_position - movement_start_position)
    if movement_actual_distance is None and movement_start_position is not None and movement_end_position is not None:
        movement_actual_distance = abs(movement_end_position - movement_start_position)
    if movement_distance is None:
        movement_distance = movement_actual_distance if movement_actual_distance is not None else movement_commanded_distance
    if movement_actual_distance is not None:
        movement_distance_source = "ActualEndPosition"
        movement_distance_method = "KnownStartPositionToActualEnd"
    hardware_actual_distance = movement_distance
    hardware_commanded_distance = movement_commanded_distance
    hardware_start_position = movement_start_position
    hardware_end_position = movement_end_position
    if hardware_end_position is None and hardware_start_position is not None and hardware_actual_distance is not None:
        hardware_end_position = hardware_start_position + abs(hardware_actual_distance)
    hardware_target_position = target_position
    return ActivityRecord(
        axis=axis,
        rule_id=rule_id,
        event_type=start_event,
        start_event=start_event,
        end_event=end_event,
        start_time=start_time,
        end_time=end_time,
        duration_ms=int(seconds * 1000),
        duration_s=seconds,
        start_value=start_value,
        end_value=end_value,
        movement_start_position=movement_start_position,
        movement_target_position=target_position,
        movement_end_position=movement_end_position,
        movement_commanded_distance=movement_commanded_distance,
        movement_actual_distance=movement_actual_distance,
        movement_distance=movement_distance,
        movement_distance_source=movement_distance_source,
        movement_distance_method=movement_distance_method,
        hardware_motion_match_status="MatchedByOverlappingTime" if hardware_actual_distance is not None else "NoHardwareSegmentFound",
        hardware_motion_source_file="hardware.log" if hardware_actual_distance is not None else "",
        hardware_raw_start_position=0 if hardware_actual_distance is not None else None,
        hardware_raw_end_position=int(abs(hardware_actual_distance) * 1000) if hardware_actual_distance is not None else None,
        hardware_start_position=hardware_start_position,
        hardware_end_position=hardware_end_position,
        hardware_target_position=hardware_target_position,
        hardware_actual_distance=hardware_actual_distance,
        hardware_commanded_distance=hardware_commanded_distance,
        hardware_start_line_number=20 if hardware_actual_distance is not None else None,
        hardware_end_line_number=21 if hardware_actual_distance is not None else None,
        hardware_start_line_text="[N12:Z] TPOS 'S' 0 0 0 0 (0)" if hardware_actual_distance is not None else "",
        hardware_end_line_text=f"[N12:Z] TPOS 'E' 0 0 0 0 ({int(abs(hardware_actual_distance) * 1000)})" if hardware_actual_distance is not None else "",
        pwm_percent=pwm_percent,
        pwm_match_status=pwm_match_status,
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
        start_line_number=txt_offset + 1,
        end_line_number=txt_offset + 2,
        pwm_source_file="pwm.log",
        pwm_source_line="RUN 0 80 (80)",
        pwm_line_number=10,
        pwm_source_time=start_time - timedelta(seconds=1),
    )


def _reference_record(
    *,
    seconds: float = 10.0,
    txt_offset: int = 0,
    axis: str = "Y",
    hardware_reference_status: str = HARDWARE_REFERENCE_STATUS_FOUND,
) -> ActivityRecord:
    """Build a matched search-reference record for reference-duration tests."""

    start_time = datetime(2026, 1, 1, 1, 0, 0) + timedelta(seconds=txt_offset)
    return ActivityRecord(
        axis=axis,
        rule_id="search_reference",
        event_type="start searching reference",
        start_event="start searching reference",
        end_event="reference found",
        start_time=start_time,
        end_time=start_time + timedelta(seconds=seconds),
        duration_ms=int(seconds * 1000),
        duration_s=seconds,
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
        hardware_motion_match_status=HARDWARE_STATUS_REFERENCE_NOT_APPLICABLE,
        hardware_reference_match_status=hardware_reference_status,
        hardware_reference_source_file="reference.log" if hardware_reference_status == HARDWARE_REFERENCE_STATUS_FOUND else "",
        hardware_reference_zero_sensor_raw_value=76 if hardware_reference_status == HARDWARE_REFERENCE_STATUS_FOUND else None,
        hardware_reference_line_text="[N4:Y] TPOS 'Z' 76 (76)"
        if hardware_reference_status == HARDWARE_REFERENCE_STATUS_FOUND
        else "",
        movement_distance=None,
        movement_distance_source="",
        movement_distance_method="DistanceNotApplicableForReference",
        status=OVERALL_STATUS_OK
        if hardware_reference_status == HARDWARE_REFERENCE_STATUS_FOUND
        else OVERALL_STATUS_HARDWARE_WARNING,
        start_line_number=txt_offset + 1,
        end_line_number=txt_offset + 2,
        source_txt_start_line="@[Y] start searching reference",
        source_txt_end_line="@[Y] reference found (SENSOR)",
    )


def _matched_records_from_txt(tmp_path: Path, lines: list[str]) -> list[ActivityRecord]:
    """Parse, match, and validate synthetic TXT lines."""

    txt_path = _write_lines(tmp_path / "movement.txt", lines)
    parse_result = MainLogParser(TextFileLoader()).parse(txt_path)
    records = EventMatcher(DEFAULT_EVENT_RULES).build_activity_records(parse_result.timeline)
    validated = ActivityValidator().validate(records)
    return validated


def _apply_distance_selection(records: list[ActivityRecord], distance_source: str = "hardware_actual") -> None:
    """Apply the same runtime movement-distance selection the service uses for Details."""

    analyzer = DistributionAnalyzer(distance_source=distance_source)
    for record in records:
        movement = analyzer.derive_movement_distance(record)
        record.movement_distance = movement.movement_distance
        record.movement_distance_source = movement.movement_distance_source
        record.movement_distance_method = movement.movement_distance_method
        record.movement_distance_notes = movement.movement_distance_notes


def _stats_for(records: list[ActivityRecord]):
    """Return analyzer, rows, grouped rows, and stats for records."""

    analyzer = DistributionAnalyzer()
    rows = analyzer.build_input_rows(records, "UroBiopsy_20260410.txt")
    grouped = analyzer.group_rows(rows)
    stats = analyzer.compute_group_stats(grouped)
    return analyzer, rows, grouped, stats


def _distribution_result_for(records: list[ActivityRecord]):
    """Return a full distribution result for gallery-export tests."""

    return DistributionAnalyzer().analyze(records, "sample.txt")


def _write_gallery_with_charts(
    tmp_path: Path,
    records: list[ActivityRecord],
    max_images: int = 500,
    include_skipped_groups: bool = False,
    layout: str = "vertical",
):
    """Generate charts and export a gallery workbook for synthetic records."""

    result = _distribution_result_for(records)
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    gallery_path = tmp_path / "gallery.xlsx"
    export_result = DistributionImageGalleryExporter(
        max_images=max_images,
        include_skipped_groups=include_skipped_groups,
        layout=layout,
    ).export(gallery_path, result)
    return result, export_result, load_workbook(gallery_path, data_only=True)


def test_matcher_keeps_raw_distance_components_until_runtime_selection(tmp_path: Path) -> None:
    """Matcher should not bake the global distribution distance source into Details fields."""

    txt_path = _write_lines(
        tmp_path / "raw.txt",
        [
            "2026-04-30 09:50:25:829 MCU   @[Y] max: 0.00",
            "2026-04-30 09:50:27:029 MCU   @[Y] start clearing: -19.50",
            "2026-04-30 09:50:33:666 MCU   @[Y] motor cleared",
            "2026-04-30 09:50:33:667 MCU   @[Y] min: -29.51",
        ],
    )
    parse_result = MainLogParser(TextFileLoader()).parse(txt_path)
    records = EventMatcher(DEFAULT_EVENT_RULES).build_activity_records(parse_result.timeline)

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_commanded_distance == pytest.approx(19.50)
    assert record.movement_actual_distance == pytest.approx(29.51)
    assert record.movement_distance is None
    assert record.movement_distance_source == ""


def test_distribution_grouping_key_splits_only_on_configured_fields() -> None:
    """Groups should be keyed by TXT file, PWM, axis, rounded distance, and rule."""

    _, _, grouped, _ = _stats_for([_record(txt_offset=0, pwm_percent=-80), _record(txt_offset=20, pwm_percent=80)])
    assert len(grouped) == 1

    _, _, grouped, _ = _stats_for([_record(txt_offset=0), _record(txt_offset=20, pwm_percent=60)])
    assert len(grouped) == 2

    _, _, grouped, _ = _stats_for([_record(txt_offset=0), _record(txt_offset=20, axis="Y")])
    assert len(grouped) == 2

    _, _, grouped, _ = _stats_for(
        [
            _record(txt_offset=0),
            _record(txt_offset=20, start_value=-60, movement_target_position=-60, movement_distance=60),
        ]
    )
    assert len(grouped) == 2

    _, _, grouped, _ = _stats_for(
        [
            _record(txt_offset=0),
            _record(
                txt_offset=20,
                rule_id="move_to_home",
                start_event="start moving to home",
                end_event="motor homed",
            ),
        ]
    )
    assert len(grouped) == 2


def test_control_log_parser_parses_tpos_start_end_and_builds_hardware_segment(tmp_path: Path) -> None:
    """Control logs should expose TPOS S/E records and hardware actual distance."""

    log_path = _write_lines(
        tmp_path / "hardware.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N3:X] TPOS 'S' 0 0 0 0 (0)",
            "2026-01-01 00:00:03:000 [OUT] sample",
            "                              [N3:X] TPOS 'E' 255 236 72 0 (-1292112)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)

    assert [event.position_kind for event in result.hardware_position_events] == ["Start", "End"]
    assert result.hardware_position_events[0].axis == "X"
    assert result.hardware_position_events[1].raw_position == -1292112
    assert len(result.hardware_motion_segments) == 1
    segment = result.hardware_motion_segments[0]
    assert segment.axis == "X"
    assert segment.raw_start_position == 0
    assert segment.raw_end_position == -1292112
    assert segment.hardware_actual_distance == pytest.approx(22.01, abs=0.01)


def test_control_log_parser_tpos_zero_sensor_is_reference_not_position(tmp_path: Path) -> None:
    """TPOS Z should be reference evidence, not a raw motor-position sample."""

    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'Z' 76 (76)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)

    assert len(result.hardware_position_events) == 1
    event = result.hardware_position_events[0]
    assert event.axis == "Y"
    assert event.position_kind == "ZeroSensor"
    assert event.raw_position is None
    assert event.reference_raw_value == 76
    assert event.hardware_status_value == 76
    assert event.physical_position is None
    assert result.hardware_motion_segments == []
    assert len(result.hardware_reference_events) == 1
    evidence = result.hardware_reference_events[0]
    assert evidence.evidence_type == "ZeroSensor"
    assert evidence.raw_value == 76


def test_control_log_parser_tpos_reset_init_is_reference_evidence(tmp_path: Path) -> None:
    """TPOS I should be preserved as reset/init reference evidence."""

    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'I'",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)

    assert len(result.hardware_position_events) == 1
    event = result.hardware_position_events[0]
    assert event.position_kind == "ResetOrInit"
    assert event.raw_position is None
    assert event.reference_raw_value is None
    assert event.physical_position is None
    assert result.hardware_motion_segments == []
    assert len(result.hardware_reference_events) == 1
    assert result.hardware_reference_events[0].evidence_type == "ResetOrInit"


def test_hardware_target_attach_rejects_stale_target(tmp_path: Path) -> None:
    """Hardware commanded distance should not reuse an old TPOS target for a later Start."""

    log_path = _write_lines(
        tmp_path / "hardware.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N3:X] TPOS 0 0 0 0 (-1292112)",
            "2026-01-01 00:10:00:000 [OUT] sample",
            "                              [N3:X] TPOS 'S' 0 0 0 0 (0)",
            "2026-01-01 00:10:02:000 [OUT] sample",
            "                              [N3:X] TPOS 'E' 255 236 72 0 (-1292112)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)
    segment = result.hardware_motion_segments[0]

    assert segment.hardware_actual_distance == pytest.approx(22.01, abs=0.01)
    assert segment.raw_target_position is None
    assert segment.hardware_commanded_distance is None
    assert "stale" in segment.notes


def test_hardware_target_attach_accepts_fresh_target(tmp_path: Path) -> None:
    """Hardware target TPOS close to the Start should remain available for audit."""

    log_path = _write_lines(
        tmp_path / "hardware.log",
        [
            "2026-01-01 00:00:00:900 [OUT] sample",
            "                              [N3:X] TPOS 0 0 0 0 (-1292112)",
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N3:X] TPOS 'S' 0 0 0 0 (0)",
            "2026-01-01 00:00:03:000 [OUT] sample",
            "                              [N3:X] TPOS 'E' 255 236 72 0 (-1292112)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)
    segment = result.hardware_motion_segments[0]

    assert segment.raw_target_position == -1292112
    assert segment.hardware_commanded_distance == pytest.approx(22.01, abs=0.01)


def test_hardware_motion_segments_include_incomplete_and_unmatched_end(tmp_path: Path) -> None:
    """Hardware segment audit data should keep incomplete starts and unmatched ends."""

    log_path = _write_lines(
        tmp_path / "hardware.log",
        [
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N3:X] TPOS 'S' 0 0 0 0 (0)",
            "2026-01-01 00:00:03:000 [OUT] sample",
            "                              [N4:Y] TPOS 'E' 255 236 72 0 (-1145360)",
        ],
    )

    result = DutyCycleLogParser(TextFileLoader()).parse(log_path)
    statuses = {segment.match_status for segment in result.hardware_motion_segments}

    assert "HardwareSegmentIncomplete" in statuses
    assert "UnmatchedHardwareEnd" in statuses


def test_reference_evidence_matches_search_reference_without_distance_warning(tmp_path: Path) -> None:
    """Search-reference activities should use TPOS Z/I evidence instead of requiring S/E distance."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="Y",
        rule_id="search_reference",
        start_event="start searching reference",
        end_event="reference found",
        start_time=start,
        end_time=start + timedelta(seconds=3),
        match_status=STATUS_MATCHED,
    )
    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'Z' 76 (76)",
            "2026-01-01 09:00:02:000 [OUT] sample",
            "                              [N4:Y] TPOS 'I'",
        ],
    )
    log_result = DutyCycleLogParser(TextFileLoader()).parse(log_path)

    HardwareMotionAssociator().attach([record], [log_result])
    ActivityValidator().validate([record])

    assert record.hardware_reference_match_status == HARDWARE_REFERENCE_STATUS_FOUND
    assert record.hardware_motion_match_status == HARDWARE_STATUS_REFERENCE_NOT_APPLICABLE
    assert record.hardware_reference_source_file == "reference.log"
    assert record.hardware_reference_time_delta_ms == 0
    assert record.hardware_reference_zero_sensor_raw_value == 76
    assert record.hardware_actual_distance is None
    assert record.movement_distance is None
    assert record.movement_distance_method == "DistanceNotApplicableForReference"
    assert record.hardware_warning is False
    assert record.status == OVERALL_STATUS_OK
    evidence_statuses = {event.evidence_type: event.match_status for event in log_result.hardware_reference_events}
    assert evidence_statuses == {
        "ZeroSensor": HARDWARE_REFERENCE_STATUS_FOUND,
        "ResetOrInit": HARDWARE_REFERENCE_STATUS_FOUND,
    }


def test_summary_exports_hardware_reference_events_rows(tmp_path: Path) -> None:
    """Workbook frames should expose TPOS Z/I reference evidence for audit."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="Y",
        rule_id="search_reference",
        start_event="start searching reference",
        end_event="reference found",
        start_time=start,
        end_time=start + timedelta(seconds=3),
        match_status=STATUS_MATCHED,
    )
    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'Z' 76 (76)",
            "2026-01-01 09:00:02:000 [OUT] sample",
            "                              [N4:Y] TPOS 'I'",
        ],
    )
    log_result = DutyCycleLogParser(TextFileLoader()).parse(log_path)
    HardwareMotionAssociator().attach([record], [log_result])
    ActivityValidator().validate([record])

    frames = SummaryGenerator().build_report_frames([record], [log_result])
    rows = frames.hardware_reference_events.to_dict("records")

    assert len(rows) == 2
    zero_row = next(row for row in rows if row["Evidence Type"] == "ZeroSensor")
    reset_row = next(row for row in rows if row["Evidence Type"] == "ResetOrInit")
    assert zero_row["Raw / Status Value"] == 76
    assert zero_row["Matched TXT Rule ID"] == "search_reference"
    assert zero_row["Match Status"] == HARDWARE_REFERENCE_STATUS_FOUND
    assert reset_row["Matched TXT Rule ID"] == "search_reference"


def test_search_reference_without_reference_evidence_is_hardware_warning() -> None:
    """Search-reference rows still need Z/I reference evidence for a clean hardware status."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="Y",
        rule_id="search_reference",
        start_event="start searching reference",
        end_event="reference found",
        start_time=start,
        end_time=start + timedelta(seconds=3),
        match_status=STATUS_MATCHED,
    )

    HardwareMotionAssociator().attach([record], [DutyCycleLogFileResult(source_path=Path("empty.log"))])
    ActivityValidator().validate([record])

    assert record.hardware_reference_match_status == HARDWARE_REFERENCE_STATUS_NO_EVIDENCE
    assert record.hardware_motion_match_status == HARDWARE_STATUS_REFERENCE_NOT_APPLICABLE
    assert record.movement_distance is None
    assert record.hardware_warning is True
    assert record.status == OVERALL_STATUS_HARDWARE_WARNING


def test_tpos_reference_evidence_is_not_used_as_normal_motion_distance(tmp_path: Path) -> None:
    """TPOS Z/I evidence must not satisfy a clear_motor hardware-distance requirement."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="Y",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=start,
        end_time=start + timedelta(seconds=3),
        match_status=STATUS_MATCHED,
    )
    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'Z' 76 (76)",
            "2026-01-01 09:00:02:000 [OUT] sample",
            "                              [N4:Y] TPOS 'I'",
        ],
    )
    log_result = DutyCycleLogParser(TextFileLoader()).parse(log_path)

    HardwareMotionAssociator().attach([record], [log_result])
    ActivityValidator().validate([record])

    assert record.hardware_motion_match_status == HARDWARE_STATUS_NO_SEGMENT
    assert record.hardware_reference_match_status == HARDWARE_REFERENCE_STATUS_NOT_APPLICABLE
    assert record.hardware_actual_distance is None
    assert record.movement_distance is None
    assert record.hardware_warning is True
    assert record.status == OVERALL_STATUS_HARDWARE_WARNING


def test_duplicate_hardware_segments_are_deduped_for_matching_but_kept_for_audit() -> None:
    """Exact copied hardware segments should not create ambiguous matching by themselves."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    first = HardwareMotionSegment(
        source_path=Path("copy-a.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )
    second = HardwareMotionSegment(
        source_path=Path("copy-b.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [
            DutyCycleLogFileResult(source_path=Path("copy-a.log"), hardware_motion_segments=[first]),
            DutyCycleLogFileResult(source_path=Path("copy-b.log"), hardware_motion_segments=[second]),
        ],
    )

    assert record.hardware_motion_match_status == HARDWARE_STATUS_MATCHED_OVERLAP
    assert record.movement_distance == pytest.approx(22.01)
    assert first.possible_duplicate_hardware_segment is True
    assert second.possible_duplicate_hardware_segment is True
    assert first.duplicate_segment_count == 2
    assert "copy-a.log" in first.possible_duplicate_source_files
    assert "copy-b.log" in first.possible_duplicate_source_files
    assert sum(segment.effective_segment_used_for_matching for segment in (first, second)) == 1
    assert sum(segment.matched_txt_activity_count for segment in (first, second)) == 1


def test_unique_hardware_segment_duplicate_fields_are_blank_and_effective() -> None:
    """Normal non-duplicate hardware rows should not display a misleading duplicate count of one."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    segment = HardwareMotionSegment(
        source_path=Path("unique.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("unique.log"), hardware_motion_segments=[segment])],
    )

    assert segment.possible_duplicate_hardware_segment is False
    assert segment.duplicate_segment_count == 0
    assert segment.effective_segment_used_for_matching is True
    assert segment.matched_txt_activity_count == 1

    frames = SummaryGenerator().build_report_frames([record], [DutyCycleLogFileResult(source_path=Path("unique.log"), hardware_motion_segments=[segment])])
    segment_row = frames.hardware_motion_segments.iloc[0].to_dict()
    assert segment_row["Duplicate Segment Count"] is None
    assert segment_row["Possible Duplicate Hardware Segment"] is False
    assert segment_row["Effective Segment Used For Matching"] is True


def test_duplicate_hardware_dedupe_preserves_target_information() -> None:
    """The effective duplicate representative should keep target/commanded-distance audit fields."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    without_target = HardwareMotionSegment(
        source_path=Path("copy-a.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_start_position=0.0,
        hardware_end_position=-22.01,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )
    with_target = HardwareMotionSegment(
        source_path=Path("copy-b.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        target_time=start - timedelta(milliseconds=100),
        raw_start_position=0,
        raw_end_position=-1292112,
        raw_target_position=-1292112,
        hardware_start_position=0.0,
        hardware_end_position=-22.01,
        hardware_target_position=-22.01,
        hardware_actual_distance=22.01,
        hardware_commanded_distance=22.01,
        target_line_number=7,
        target_line_text="[N3:X] TPOS 0 0 0 0 (-1292112)",
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [
            DutyCycleLogFileResult(source_path=Path("copy-a.log"), hardware_motion_segments=[without_target]),
            DutyCycleLogFileResult(source_path=Path("copy-b.log"), hardware_motion_segments=[with_target]),
        ],
    )

    assert record.hardware_raw_target_position == -1292112
    assert record.hardware_commanded_distance == pytest.approx(22.01)
    assert record.hardware_target_line_text == "[N3:X] TPOS 0 0 0 0 (-1292112)"


def test_hardware_nearest_exact_start_is_selected_before_later_segment() -> None:
    """A zero-millisecond nearest delta should sort as the best candidate, not infinity."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_value=-1.0,
        movement_commanded_distance=1.0,
        start_time=start,
        end_time=start + timedelta(milliseconds=500),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    exact = HardwareMotionSegment(
        source_path=Path("exact.log"),
        axis="X",
        start_time=start,
        end_time=start - timedelta(milliseconds=1),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )
    later = HardwareMotionSegment(
        source_path=Path("later.log"),
        axis="X",
        start_time=start + timedelta(seconds=1),
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-58708,
        hardware_actual_distance=1.0,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("hardware.log"), hardware_motion_segments=[later, exact])],
    )

    assert record.hardware_motion_source_file == "exact.log"
    assert record.hardware_motion_time_delta_ms == 0
    assert record.hardware_actual_distance == pytest.approx(22.01)


def test_hardware_nearest_future_status_is_distinct() -> None:
    """Nearest fallback should show whether the hardware segment is in the future."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_value=-1.0,
        movement_commanded_distance=1.0,
        start_time=start,
        end_time=start + timedelta(milliseconds=500),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    future = HardwareMotionSegment(
        source_path=Path("future.log"),
        axis="X",
        start_time=start + timedelta(seconds=1),
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-58708,
        hardware_actual_distance=1.0,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("hardware.log"), hardware_motion_segments=[future])],
    )

    assert record.hardware_motion_match_status == HARDWARE_STATUS_MATCHED_NEAREST_FUTURE
    assert record.hardware_actual_distance == pytest.approx(1.0)
    assert record.movement_distance is None
    assert record.hardware_distance_consistency_status == "CandidateOnlyNotValidated"
    ActivityValidator().validate([record])
    assert record.hardware_warning is True
    assert record.status == OVERALL_STATUS_HARDWARE_WARNING


def test_details_split_selected_and_candidate_hardware_distances() -> None:
    """Reliable overlap rows use selected distance; weak rows expose only candidate distance."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    overlap = ActivityRecord(
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    nearest = ActivityRecord(
        axis="X",
        start_time=start + timedelta(seconds=10),
        end_time=start + timedelta(seconds=11),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    overlap_segment = HardwareMotionSegment(
        source_path=Path("overlap.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )
    future_segment = HardwareMotionSegment(
        source_path=Path("future.log"),
        axis="X",
        start_time=start + timedelta(seconds=11, milliseconds=100),
        end_time=start + timedelta(seconds=12),
        raw_start_position=0,
        raw_end_position=-58708,
        hardware_actual_distance=1.0,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator(match_window_ms=2000).attach(
        [overlap, nearest],
        [
            DutyCycleLogFileResult(
                source_path=Path("hardware.log"),
                hardware_motion_segments=[overlap_segment, future_segment],
            )
        ],
    )
    frames = SummaryGenerator().build_report_frames([overlap, nearest], [])
    detail_rows = frames.details.to_dict("records")

    overlap_row = next(row for row in detail_rows if row["Hardware Motion Match Status"] == HARDWARE_STATUS_MATCHED_OVERLAP)
    nearest_row = next(row for row in detail_rows if row["Hardware Motion Match Status"] == HARDWARE_STATUS_MATCHED_NEAREST_FUTURE)
    assert overlap_row["Selected Hardware Actual Distance"] == pytest.approx(22.01)
    assert math.isnan(overlap_row["Candidate Hardware Actual Distance"])
    assert math.isnan(nearest_row["Selected Hardware Actual Distance"])
    assert nearest_row["Candidate Hardware Actual Distance"] == pytest.approx(1.0)


def test_multiple_hardware_candidates_are_candidates_not_selected_distance() -> None:
    """Ambiguous overlapping hardware candidates should not fill selected movement distance."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    segments = [
        HardwareMotionSegment(
            source_path=Path("a.log"),
            axis="X",
            start_time=start,
            end_time=start + timedelta(seconds=2),
            raw_start_position=0,
            raw_end_position=-1292112,
            hardware_actual_distance=22.01,
            match_status="HardwareSegmentComplete",
        ),
        HardwareMotionSegment(
            source_path=Path("b.log"),
            axis="X",
            start_time=start + timedelta(milliseconds=10),
            end_time=start + timedelta(seconds=2, milliseconds=10),
            raw_start_position=0,
            raw_end_position=-1292112,
            hardware_actual_distance=22.01,
            match_status="HardwareSegmentComplete",
        ),
    ]

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("hardware.log"), hardware_motion_segments=segments)],
    )
    ActivityValidator().validate([record])

    assert record.hardware_motion_match_status == HARDWARE_STATUS_MULTIPLE_CANDIDATES
    assert record.hardware_actual_distance == pytest.approx(22.01)
    assert record.movement_distance is None
    assert record.status == OVERALL_STATUS_HARDWARE_WARNING


def test_hardware_distance_consistency_check_uses_txt_target_not_reported_min() -> None:
    """Consistency diagnostics compare hardware distance to TXT target-derived distance only."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_value=-22.0,
        movement_commanded_distance=22.0,
        movement_actual_distance=32.01,
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    segment = HardwareMotionSegment(
        source_path=Path("x.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1292112,
        hardware_actual_distance=22.01,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("x.log"), hardware_motion_segments=[segment])],
    )

    assert record.hardware_distance_consistency_status == "ConsistentWithTXTTarget"
    assert record.hardware_distance_consistency_delta == pytest.approx(0.01, abs=0.02)


def test_hardware_distance_consistency_flags_large_target_delta() -> None:
    """A hardware distance far from the TXT target-derived command should be flagged."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="X",
        start_value=-22.0,
        movement_commanded_distance=22.0,
        start_time=start,
        end_time=start + timedelta(seconds=2),
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
    )
    segment = HardwareMotionSegment(
        source_path=Path("x.log"),
        axis="X",
        start_time=start,
        end_time=start + timedelta(seconds=2),
        raw_start_position=0,
        raw_end_position=-1879083,
        hardware_actual_distance=32.01,
        match_status="HardwareSegmentComplete",
    )

    HardwareMotionAssociator().attach(
        [record],
        [DutyCycleLogFileResult(source_path=Path("x.log"), hardware_motion_segments=[segment])],
    )

    assert record.hardware_distance_consistency_status == "DistanceDiffTooLarge"
    assert record.hardware_distance_consistency_delta == pytest.approx(10.01)


def test_distribution_excludes_when_hardware_actual_distance_is_missing() -> None:
    """Software distances must not be used when no hardware S/E segment is attached."""

    record = _record(movement_distance=None, movement_commanded_distance=None, movement_start_position=None)
    record.movement_commanded_distance = 32.01

    analyzer, rows, grouped, stats = _stats_for([record])

    assert rows == []
    assert grouped == {}
    assert stats == []
    assert analyzer.exclusion_counts["NoHardwareSegmentFound"] == 1
    ActivityValidator().validate([record])
    assert record.hardware_warning is True
    assert record.status == OVERALL_STATUS_HARDWARE_WARNING


def test_distribution_exclusion_counts_split_hardware_from_secondary_pwm() -> None:
    """A missing-hardware primary exclusion should not be reported as only PWM reliability."""

    repeated_note = (
        "No complete same-axis hardware TPOS Start/End segment was matched; "
        "software TXT positions are not used as selected distance."
    )
    record = ActivityRecord(
        axis="Z",
        rule_id="clear_motor",
        start_event="start clearing",
        end_event="motor cleared",
        start_time=datetime(2026, 1, 1, 9, 0, 0),
        end_time=datetime(2026, 1, 1, 9, 0, 2),
        duration_ms=2000,
        duration_s=2.0,
        match_status=STATUS_MATCHED,
        duration_status=DURATION_STATUS_VALID,
        hardware_motion_match_status=HARDWARE_STATUS_NO_SEGMENT,
        pwm_match_status=PWM_STATUS_LATEST_BEFORE_TOO_FAR,
        pwm_percent=None,
        notes=repeated_note,
        movement_distance_notes=repeated_note,
        pwm_missing_reason=repeated_note,
        source_txt_start_line="@[Z] start clearing: -50.00",
    )

    result = DistributionAnalyzer().analyze([record], "sample.txt")
    exclusion = result.exclusions[0]
    breakdown = LogAnalysisService()._distribution_exclusion_breakdown(result)

    assert exclusion.reason == "NoHardwareSegmentFound"
    assert exclusion.secondary_reasons == "LatestBeforeStartTooFar"
    assert exclusion.notes.count(repeated_note) == 1
    assert breakdown["missing_hardware"] == 1
    assert breakdown["missing_pwm"] == 0
    assert breakdown["unreliable_pwm"] == 0
    assert breakdown["multiple_reasons"] == 1


def test_distribution_excludes_nearest_hardware_match_by_default() -> None:
    """Nearest hardware segment fallback must not be used for distribution unless allowed."""

    record = _record(movement_distance=22.01)
    record.hardware_motion_match_status = HARDWARE_STATUS_MATCHED_NEAREST_FUTURE

    analyzer, rows, _, _ = _stats_for([record])

    assert rows == []
    assert analyzer.exclusion_counts["NearestHardwareSegmentNotAllowedForDistribution"] == 1


def test_distribution_can_include_nearest_hardware_match_when_allowed() -> None:
    """Explicit configuration can include nearest hardware matches for diagnostics-oriented runs."""

    record = _record(movement_distance=22.01)
    record.hardware_motion_match_status = HARDWARE_STATUS_MATCHED_NEAREST_FUTURE
    analyzer = DistributionAnalyzer(
        allowed_hardware_motion_match_statuses={
            HARDWARE_STATUS_MATCHED_OVERLAP,
            HARDWARE_STATUS_MATCHED_NEAREST_FUTURE,
        }
    )

    rows = analyzer.build_input_rows([record], "sample.txt")

    assert len(rows) == 1


def test_distribution_excludes_multiple_hardware_candidates_by_default() -> None:
    """Ambiguous overlapping hardware matches must be excluded by default."""

    record = _record(movement_distance=22.01)
    record.hardware_motion_match_status = HARDWARE_STATUS_MULTIPLE_CANDIDATES

    analyzer, rows, _, _ = _stats_for([record])

    assert rows == []
    assert analyzer.exclusion_counts["MultipleHardwareCandidatesNotAllowedForDistribution"] == 1


def test_distribution_grouping_uses_hardware_actual_distance_not_software_distance() -> None:
    """Software-reported distance should not affect distribution grouping."""

    same_hardware_a = _record(txt_offset=0, movement_distance=22.01)
    same_hardware_b = _record(txt_offset=20, movement_distance=22.01)
    same_hardware_a.movement_actual_distance = 32.01
    same_hardware_b.movement_actual_distance = 99.99

    _, _, grouped, _ = _stats_for([same_hardware_a, same_hardware_b])
    assert len(grouped) == 1

    different_hardware_a = _record(txt_offset=40, movement_distance=22.01)
    different_hardware_b = _record(txt_offset=60, movement_distance=32.01)
    different_hardware_a.movement_actual_distance = 32.01
    different_hardware_b.movement_actual_distance = 32.01

    _, _, grouped, _ = _stats_for([different_hardware_a, different_hardware_b])
    assert len(grouped) == 2


def test_distribution_uses_hardware_actual_distance() -> None:
    """Distribution should select hardware actual distance, not software target values."""

    analyzer = DistributionAnalyzer()
    movement = analyzer.derive_movement_distance(
        _record(
            start_value=-25.16,
            movement_start_position=-49.02,
            movement_target_position=-25.16,
            movement_distance=23.86,
        )
    )
    assert movement.movement_start_position == -49.02
    assert movement.movement_target_position == -25.16
    assert movement.movement_distance == pytest.approx(23.86)
    assert movement.movement_distance_method == "TPOSStartEndRawDifference"
    assert movement.movement_distance_source == "HardwareActualDistance"

    movement = analyzer.derive_movement_distance(
        _record(
            start_value=-30.0,
            end_value=-10.0,
            movement_start_position=None,
            movement_target_position=-30,
            movement_distance=None,
            movement_distance_method="MissingStartPosition",
        )
    )
    assert movement.movement_distance is None
    assert movement.movement_distance_method == "MissingHardwareActualDistance"


def test_move_to_home_distance_uses_previous_position_not_target_absolute(tmp_path: Path) -> None:
    """H home movement should use min position to target, not abs(target)."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:08:24:786 MCU   @[H] motor cleared",
            "2026-04-10 09:08:24:786 MCU   @[H] min: -49.02",
            "2026-04-10 09:08:24:987 MCU   @[H] start moving to home: -25.16",
            "2026-04-10 09:08:29:226 MCU   @[H] motor homed",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position == -49.02
    assert record.movement_target_position == -25.16
    assert record.movement_commanded_distance == pytest.approx(23.86)
    assert record.movement_distance is None


def test_move_to_max_distance_uses_previous_min_position(tmp_path: Path) -> None:
    """X max movement should use previous min position, not target absolute value."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:08:18:557 MCU   @[X] motor cleared",
            "2026-04-10 09:08:18:558 MCU   @[X] min: -32.01",
            "2026-04-10 09:08:18:760 MCU   @[X] start moving to max pos: -3.00",
            "2026-04-10 09:08:28:623 MCU   @[X] motor reached max pos",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_commanded_distance == pytest.approx(29.01)
    assert record.movement_distance is None


def test_z_move_to_home_distance_uses_previous_min_position(tmp_path: Path) -> None:
    """Z home movement should use previous min position, not target absolute value."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:22:44:008 MCU   @[Z] motor cleared",
            "2026-04-10 09:22:44:009 MCU   @[Z] min: -50.00",
            "2026-04-10 09:22:44:031 MCU   @[Z] start moving to home: -30.10",
            "2026-04-10 09:22:53:786 MCU   @[Z] motor homed",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_commanded_distance == pytest.approx(19.90)
    assert record.movement_distance is None


def test_end_value_search_ignores_unrelated_axes_until_companion_line(tmp_path: Path) -> None:
    """Same-axis companion min/max lines should not be missed because other axes interleave."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:50:25:829 MCU   @[Y] max: 0.00",
            "2026-04-30 09:50:27:029 MCU   @[Y] start clearing: -19.50",
            "2026-04-30 09:50:33:666 MCU   @[Y] motor cleared",
            "2026-04-30 09:50:33:667 MCU   @[X] max: 0.00",
            "2026-04-30 09:50:33:668 MCU   @[Z] max: 0.00",
            "2026-04-30 09:50:33:669 MCU   @[H] max: 0.00",
            "2026-04-30 09:50:33:670 MCU   @[N] max: 69.51",
            "2026-04-30 09:50:33:671 MCU   @[P] max: 0.00",
            "2026-04-30 09:50:33:672 MCU   @[Y] min: -29.51",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_end_position == -29.51
    assert record.movement_actual_distance == pytest.approx(29.51)
    assert record.movement_distance is None


def test_reset_physical_position_to_zero_updates_position_state(tmp_path: Path) -> None:
    """A reset physical position line should act as a zero-position hint."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:50:25:829 MCU   @[H] reset physical position to 0",
            "2026-04-30 09:50:27:029 MCU   @[H] start clearing: -49.00",
            "2026-04-30 09:50:33:666 MCU   @[H] motor cleared",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position == 0.0
    assert record.movement_target_position == -49.0
    assert record.movement_commanded_distance == pytest.approx(49.0)
    assert record.movement_actual_distance is None
    assert record.movement_distance is None
    assert record.movement_distance_source == ""


def test_position_state_resets_on_new_initialization(tmp_path: Path) -> None:
    """A hard workflow boundary must prevent stale positions from leaking into the next cycle."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:49:00:000 MCU   @[X] min: -32.01",
            "2026-04-30 09:50:00:000 User  @on_startInitRobot_clicked",
            "2026-04-30 09:50:27:029 MCU   @[X] start moving to max pos: -3.00",
            "2026-04-30 09:50:33:666 MCU   @[X] motor reached max pos",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position is None
    assert record.movement_target_position == -3.0
    assert record.movement_distance is None
    assert record.movement_distance_method == "MissingStartPosition"


def test_tool_menu_boundary_does_not_reset_position_state_by_default(tmp_path: Path) -> None:
    """Soft menu boundaries should not discard physical position hints by default."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:49:00:000 MCU   @[H] min: -49.03",
            "2026-04-30 09:49:01:000 User  @tool menu selected",
            "2026-04-30 09:49:02:000 MCU   @[H] start moving to home: -25.16",
            "2026-04-30 09:49:06:000 MCU   @[H] motor homed",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position == -49.03
    assert record.movement_commanded_distance == pytest.approx(23.87)
    assert record.movement_distance is None


def test_clearing_from_known_zero_prefers_actual_end_when_available(tmp_path: Path) -> None:
    """Clearing from known zero should prefer the actual end companion position."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:20:23:500 MCU   @[Z] max: 0.00",
            "2026-04-10 09:20:23:541 MCU   @[Z] start clearing: -50.00",
            "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared",
            "2026-04-10 09:20:46:317 MCU   @[Z] min: -50.00",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position == 0.0
    assert record.movement_target_position == -50.0
    assert record.movement_end_position == -50.0
    assert record.movement_commanded_distance == pytest.approx(50.0)
    assert record.movement_actual_distance == pytest.approx(50.0)
    assert record.movement_distance is None


def test_y_clear_motor_prefers_actual_end_over_command_target(tmp_path: Path) -> None:
    """Y clearing should use min position after motor cleared, not the commanded target."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:50:25:829 MCU   @[Y] max: 0.00",
            "2026-04-30 09:50:27:029 MCU   @[Y] start clearing: -19.50",
            "2026-04-30 09:50:33:666 MCU   @[Y] motor cleared",
            "2026-04-30 09:50:33:667 MCU   @[Y] min: -29.51",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_start_position == 0.0
    assert record.movement_target_position == -19.50
    assert record.movement_end_position == -29.51
    assert record.movement_commanded_distance == pytest.approx(19.50)
    assert record.movement_actual_distance == pytest.approx(29.51)
    assert record.movement_distance is None


def test_x_clear_motor_prefers_actual_end_over_command_target(tmp_path: Path) -> None:
    """X clearing should use actual min position after completion."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:50:25:829 MCU   @[X] max: 0.00",
            "2026-04-30 09:50:27:029 MCU   @[X] start clearing: -22.00",
            "2026-04-30 09:50:33:666 MCU   @[X] motor cleared",
            "2026-04-30 09:50:33:668 MCU   @[X] min: -32.01",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_commanded_distance == pytest.approx(22.00)
    assert record.movement_actual_distance == pytest.approx(32.01)
    assert record.movement_distance is None


def test_h_clear_motor_prefers_actual_end_when_available(tmp_path: Path) -> None:
    """H clearing should preserve both commanded and actual distance values."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-30 09:50:25:829 MCU   @[H] max: 0.00",
            "2026-04-30 09:50:27:029 MCU   @[H] start clearing: -49.00",
            "2026-04-30 09:50:33:666 MCU   @[H] motor cleared",
            "2026-04-30 09:50:33:666 MCU   @[H] min: -49.03",
        ],
    )

    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    assert record.movement_commanded_distance == pytest.approx(49.00)
    assert record.movement_actual_distance == pytest.approx(49.03)
    assert record.movement_distance is None


def test_completed_target_updates_next_movement_start_position(tmp_path: Path) -> None:
    """A later clear should start from the completed home target, not stale min position."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:08:24:786 MCU   @[H] motor cleared",
            "2026-04-10 09:08:24:786 MCU   @[H] min: -49.02",
            "2026-04-10 09:08:24:987 MCU   @[H] start moving to home: -25.16",
            "2026-04-10 09:08:29:226 MCU   @[H] motor homed",
            "2026-04-10 09:08:30:000 MCU   @[H] start clearing: -49.02",
            "2026-04-10 09:08:38:000 MCU   @[H] motor cleared",
            "2026-04-10 09:08:38:000 MCU   @[H] min: -49.02",
        ],
    )

    clear_record = next(
        item
        for item in records
        if item.match_status == STATUS_MATCHED and item.rule_id == "clear_motor"
    )
    assert clear_record.movement_start_position == -25.16
    assert clear_record.movement_target_position == -49.02
    assert clear_record.movement_commanded_distance == pytest.approx(23.86)
    assert clear_record.movement_distance is None


def test_missing_start_position_is_not_used_for_distribution_by_default(tmp_path: Path) -> None:
    """Rows without true movement distance should be excluded when fallback is disabled."""

    records = _matched_records_from_txt(
        tmp_path,
        [
            "2026-04-10 09:22:44:031 MCU   @[Z] start moving to home: -30.10",
            "2026-04-10 09:22:53:786 MCU   @[Z] motor homed",
        ],
    )
    record = next(item for item in records if item.match_status == STATUS_MATCHED)
    record.pwm_percent = 80
    record.pwm_match_status = PWM_STATUS_MATCHED_CONTAINING

    analyzer, rows, grouped, stats = _stats_for([record])

    assert record.movement_distance is None
    assert record.movement_distance_method == "MissingStartPosition"
    assert rows == []
    assert grouped == {}
    assert stats == []
    assert analyzer.exclusion_counts["MissingHardwareActualDistance"] == 1


def test_distribution_grouping_uses_corrected_distance() -> None:
    """Home motions should group by true travel distance, not target absolute value."""

    _, rows, grouped, stats = _stats_for(
        [
            _record(
                seconds=4,
                txt_offset=0,
                axis="H",
                start_value=-25.16,
                movement_start_position=-49.02,
                movement_target_position=-25.16,
                movement_distance=23.86,
                rule_id="move_to_home",
                start_event="start moving to home",
                end_event="motor homed",
            ),
            _record(
                seconds=5,
                txt_offset=20,
                axis="H",
                start_value=-25.16,
                movement_start_position=-49.02,
                movement_target_position=-25.16,
                movement_distance=23.86,
                rule_id="move_to_home",
                start_event="start moving to home",
                end_event="motor homed",
            ),
        ]
    )

    assert len(rows) == 2
    assert len(grouped) == 1
    assert stats[0].movement_distance_group_value == pytest.approx(23.9)
    assert stats[0].movement_distance_rounded == pytest.approx(23.9)
    assert stats[0].movement_distance_raw_example == pytest.approx(23.86)


def test_distance_binning_groups_small_float_variations_but_keeps_raw_values() -> None:
    """Default 0.1 distance bins should group near-identical physical movements."""

    _, rows, grouped, stats = _stats_for(
        [
            _record(txt_offset=0, start_value=-23.86, movement_target_position=-23.86),
            _record(txt_offset=20, start_value=-23.87, movement_target_position=-23.87),
            _record(txt_offset=40, start_value=-25.16, movement_target_position=-25.16),
        ]
    )

    assert len(rows) == 3
    assert len(grouped) == 2
    group_values = sorted(item.movement_distance_group_value for item in stats)
    assert group_values == pytest.approx([23.9, 25.2])
    assert sorted(row.movement_distance for row in rows) == pytest.approx([23.86, 23.87, 25.16])
    near_group = next(item for item in stats if item.movement_distance_group_value == pytest.approx(23.9))
    assert near_group.sample_count == 2
    assert near_group.position_values_mixed is True


def test_distribution_groups_pwm_by_absolute_percent_for_gallery_stats(tmp_path: Path) -> None:
    """Raw -80 and 80 should share one normalized distribution image-statistics group."""

    reverse = _record(seconds=10, txt_offset=0, pwm_percent=-80)
    reverse.pwm_raw_value = -80
    reverse.pwm_direction = "Reverse"
    forward = _record(seconds=12, txt_offset=20, pwm_percent=80)
    forward.pwm_raw_value = 80
    forward.pwm_direction = "Forward"

    result, _, workbook = _write_gallery_with_charts(tmp_path, [reverse, forward])

    assert len(result.stats) == 1
    assert result.stats[0].pwm_percent == 80
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("PWM (%)")] == 80
    assert row[headers.index("PWM Raw Value Example")] == -80
    assert row[headers.index("PWM Raw Values Seen")] == "-80; 80"
    assert row[headers.index("PWM Directions Seen")] == "Forward; Reverse"
    assert row[headers.index("PWM Direction Mixed")] is True
    assert row[headers.index("PWM Direction Example")] == "Reverse"
    index_sheet = workbook["Image Index"]
    index_headers = [cell.value for cell in index_sheet[1]]
    index_row = next(index_sheet.iter_rows(min_row=2, values_only=True))
    assert "PWM Raw Values Seen" in index_headers
    assert "Movement Distance Grouping Mode" in index_headers
    assert "Normal Fit Mean (s)" in index_headers
    assert "Image Insert Error" in index_headers
    assert index_row[index_headers.index("PWM Raw Values Seen")] == "-80; 80"
    assert index_row[index_headers.index("PWM Directions Seen")] == "Forward; Reverse"
    assert index_row[index_headers.index("Movement Distance Grouping Mode")] == "bin"


def test_distribution_image_statistics_single_pwm_direction_is_not_mixed(tmp_path: Path) -> None:
    """A one-direction group should still show raw PWM audit fields without a mixed flag."""

    records = [_record(seconds=value, txt_offset=index * 20, pwm_percent=80) for index, value in enumerate([10, 12])]
    for record in records:
        record.pwm_raw_value = 80
        record.pwm_direction = "Forward"

    result, _, workbook = _write_gallery_with_charts(tmp_path, records)

    assert len(result.stats) == 1
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("PWM Raw Values Seen")] == "80"
    assert row[headers.index("PWM Directions Seen")] == "Forward"
    assert row[headers.index("PWM Direction Mixed")] is False


def test_distribution_groups_distance_by_absolute_value_for_gallery_stats(tmp_path: Path) -> None:
    """Negative and positive selected distances should share the same positive group value."""

    negative = _record(
        seconds=10,
        txt_offset=0,
        start_value=None,
        movement_start_position=None,
        movement_target_position=None,
        movement_commanded_distance=-29.51,
        movement_distance=-29.51,
    )
    positive = _record(
        seconds=12,
        txt_offset=20,
        start_value=None,
        movement_start_position=None,
        movement_target_position=None,
        movement_commanded_distance=29.51,
        movement_distance=29.51,
    )

    result, _, workbook = _write_gallery_with_charts(tmp_path, [negative, positive])

    assert len(result.stats) == 1
    assert result.stats[0].movement_distance_group_value == pytest.approx(29.5)
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Selected Group Distance")] == pytest.approx(29.5)
    assert row[headers.index("Hardware Actual Distance Group Value")] == pytest.approx(29.5)
    assert row[headers.index("Hardware Actual Distance Group Display")] == "29.5"
    assert row[headers.index("Movement Distance Grouping Mode")] == "bin"
    assert row[headers.index("Movement Distance Bin Size")] == pytest.approx(0.1)
    assert row[headers.index("Hardware Actual Distance Raw Example")] == pytest.approx(29.51)
    assert row[headers.index("Hardware Actual Distance Min")] == pytest.approx(29.51)
    assert row[headers.index("Hardware Actual Distance Max")] == pytest.approx(29.51)


def test_distribution_image_statistics_exports_distance_grouping_details(tmp_path: Path) -> None:
    """Image Statistics should explain binned group values without hiding raw distances."""

    records = [
        _record(
            seconds=10,
            txt_offset=0,
            start_value=-25.16,
            movement_start_position=-49.02,
            movement_target_position=-25.16,
            movement_commanded_distance=23.86,
            movement_distance=23.86,
        ),
        _record(
            seconds=12,
            txt_offset=20,
            start_value=-25.16,
            movement_start_position=-49.03,
            movement_target_position=-25.16,
            movement_commanded_distance=23.87,
            movement_distance=23.87,
        ),
    ]

    result, _, workbook = _write_gallery_with_charts(tmp_path, records)

    assert len(result.stats) == 1
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Hardware Actual Distance Group Value")] == pytest.approx(23.9)
    assert row[headers.index("Hardware Actual Distance Group Display")] == "23.9"
    assert row[headers.index("Movement Distance Grouping Mode")] == "bin"
    assert row[headers.index("Movement Distance Bin Size")] == pytest.approx(0.1)
    assert row[headers.index("Hardware Actual Distance Raw Example")] == pytest.approx(23.86)
    assert row[headers.index("Hardware Actual Distance Min")] == pytest.approx(23.86)
    assert row[headers.index("Hardware Actual Distance Max")] == pytest.approx(23.87)
    assert row[headers.index("Position Values Mixed")] is True


def test_exact_distance_grouping_does_not_round_group_keys() -> None:
    """Exact mode should keep near-but-different distances as separate groups."""

    analyzer = DistributionAnalyzer(distance_grouping_mode="exact")
    rows = analyzer.build_input_rows(
        [
            _record(txt_offset=0, movement_distance=23.861, movement_commanded_distance=23.861),
            _record(txt_offset=20, movement_distance=23.862, movement_commanded_distance=23.862),
        ],
        "sample.txt",
    )
    grouped = analyzer.group_rows(rows)
    stats = analyzer.compute_group_stats(grouped)

    assert len(grouped) == 2
    assert sorted(item.movement_distance_group_value for item in stats) == pytest.approx([23.861, 23.862])
    assert "23.861" in stats[0].group_id or "23.862" in stats[0].group_id


def test_round_digits_distance_grouping_can_merge_exact_variations() -> None:
    """Round-digits mode should group values that round to the same configured value."""

    analyzer = DistributionAnalyzer(distance_grouping_mode="round_digits", movement_distance_round_digits=2)
    rows = analyzer.build_input_rows(
        [
            _record(txt_offset=0, movement_distance=23.861, movement_commanded_distance=23.861),
            _record(txt_offset=20, movement_distance=23.862, movement_commanded_distance=23.862),
        ],
        "sample.txt",
    )
    grouped = analyzer.group_rows(rows)

    assert len(grouped) == 1
    assert next(iter(grouped.values()))[0].movement_distance_group_value == pytest.approx(23.86)


def test_bin_distance_display_uses_bin_precision() -> None:
    """Distance display should follow bin precision instead of a fixed two-decimal format."""

    analyzer = DistributionAnalyzer(distance_grouping_mode="bin", movement_distance_bin_size=0.05)
    rows = analyzer.build_input_rows(
        [_record(txt_offset=0, movement_distance=23.861, movement_commanded_distance=23.861)],
        "sample.txt",
    )
    grouped = analyzer.group_rows(rows)
    stats = analyzer.compute_group_stats(grouped)

    assert stats[0].movement_distance_group_value == pytest.approx(23.85)
    assert "D23.85" in stats[0].group_id


def test_chart_title_uses_group_distance_display_and_source() -> None:
    """Chart titles should show grouped distance precision and movement source."""

    analyzer = DistributionAnalyzer(distance_grouping_mode="exact")
    rows = analyzer.build_input_rows(
        [_record(txt_offset=0, movement_distance=23.861, movement_commanded_distance=23.861)],
        "sample.txt",
    )
    grouped = analyzer.group_rows(rows)
    stats = analyzer.compute_group_stats(grouped)

    title = NormalDistributionChartGenerator()._chart_title(stats[0])

    assert "Hardware Actual Distance 23.861 (HardwareActualDistance)" in title
    assert "Distance 23.86 " not in title


def test_unreliable_pwm_is_excluded_from_distribution() -> None:
    """Uncertain PWM evidence should not be used for same-PWM distribution grouping."""

    analyzer, rows, grouped, stats = _stats_for(
        [
            _record(
                pwm_match_status=PWM_STATUS_LATEST_BEFORE_TOO_FAR,
            )
        ]
    )

    assert rows == []
    assert grouped == {}
    assert stats == []
    assert analyzer.exclusion_counts["LatestBeforeStartTooFar"] == 1


def test_latest_before_too_far_exclusion_is_specific() -> None:
    """PWM time-window exclusions should not be reported as generic missing PWM."""

    analyzer = DistributionAnalyzer()
    record = _record(pwm_percent=None, pwm_match_status=PWM_STATUS_LATEST_BEFORE_TOO_FAR)
    record.pwm_missing_reason = "The latest PWM record before the activity for axis Z was too far away."
    record.pwm_time_delta_ms = 599000

    rows = analyzer.build_input_rows([record], "sample.txt")

    assert rows == []
    assert analyzer.exclusions[0].reason == "LatestBeforeStartTooFar"
    assert analyzer.exclusions[0].pwm_match_status == PWM_STATUS_LATEST_BEFORE_TOO_FAR
    assert analyzer.exclusions[0].pwm_time_delta_ms == 599000


def test_nearest_future_pwm_exclusion_is_specific() -> None:
    """Nearest future PWM should be clearly labeled when distribution rejects it."""

    analyzer = DistributionAnalyzer()
    record = _record(pwm_match_status=PWM_STATUS_MATCHED_NEAREST_FUTURE)

    rows = analyzer.build_input_rows([record], "sample.txt")

    assert rows == []
    assert analyzer.exclusions[0].reason == "NearestFuturePWMNotAllowedForDistribution"
    assert analyzer.exclusions[0].pwm_match_status == PWM_STATUS_MATCHED_NEAREST_FUTURE


def test_nearest_pwm_is_excluded_from_distribution_by_default() -> None:
    """Distribution grouping should require containing-file PWM unless explicitly relaxed."""

    analyzer, rows, grouped, stats = _stats_for(
        [
            _record(
                pwm_match_status=PWM_STATUS_MATCHED_NEAREST,
            )
        ]
    )

    assert rows == []
    assert grouped == {}
    assert stats == []
    assert analyzer.exclusion_counts["NearestLogFilePWMNotAllowedForDistribution"] == 1


def test_distribution_exclusion_keeps_reference_distance_reason_when_pwm_is_also_bad() -> None:
    """Search-reference distribution exclusions should show distance-not-applicable plus PWM context."""

    analyzer = DistributionAnalyzer()
    record = _record(
        rule_id="search_reference",
        start_event="start searching reference",
        end_event="reference found",
        start_value=None,
        movement_start_position=None,
        movement_target_position=None,
        movement_commanded_distance=None,
        movement_actual_distance=None,
        movement_distance=None,
        movement_distance_method="MissingValue",
        pwm_match_status=PWM_STATUS_MATCHED_NEAREST_FUTURE,
    )

    rows = analyzer.build_input_rows([record], "sample.txt")

    assert rows == []
    exclusion = analyzer.exclusions[0]
    assert exclusion.reason == "DistanceNotApplicableForReference"
    assert exclusion.movement_distance_exclusion_reason == "DistanceNotApplicableForReference"
    assert exclusion.pwm_exclusion_reason == "NearestFuturePWMNotAllowedForDistribution"
    assert "NearestFuturePWMNotAllowedForDistribution" in exclusion.secondary_reasons


def test_search_reference_with_evidence_exclusion_is_distance_not_applicable(tmp_path: Path) -> None:
    """Reference rows with valid Z/I hardware evidence should not look like missing S/E segments."""

    start = datetime(2026, 1, 1, 9, 0, 0)
    record = ActivityRecord(
        axis="Y",
        rule_id="search_reference",
        start_event="start searching reference",
        end_event="reference found",
        start_time=start,
        end_time=start + timedelta(seconds=3),
        match_status=STATUS_MATCHED,
        pwm_percent=80,
        pwm_match_status=PWM_STATUS_MATCHED_CONTAINING,
    )
    log_path = _write_lines(
        tmp_path / "reference.log",
        [
            "2026-01-01 09:00:01:000 [OUT] sample",
            "                              [N4:Y] TPOS 'Z' 76 (76)",
            "2026-01-01 09:00:02:000 [OUT] sample",
            "                              [N4:Y] TPOS 'I'",
        ],
    )
    log_result = DutyCycleLogParser(TextFileLoader()).parse(log_path)
    HardwareMotionAssociator().attach([record], [log_result])
    ActivityValidator().validate([record])

    analyzer = DistributionAnalyzer()
    rows = analyzer.build_input_rows([record], "sample.txt")

    assert rows == []
    assert analyzer.exclusions[0].reason == "DistanceNotApplicableForReference"
    assert analyzer.exclusions[0].movement_distance_exclusion_reason == "DistanceNotApplicableForReference"
    assert analyzer.exclusions[0].reason != "NoHardwareSegmentFound"


def test_nearest_pwm_can_be_explicitly_allowed_for_distribution() -> None:
    """Relaxed distribution settings should include nearest PWM only when requested."""

    analyzer = DistributionAnalyzer(
        allowed_pwm_match_statuses={PWM_STATUS_MATCHED_CONTAINING, PWM_STATUS_MATCHED_NEAREST},
    )
    rows = analyzer.build_input_rows(
        [
            _record(
                pwm_match_status=PWM_STATUS_MATCHED_NEAREST,
            )
        ],
        "sample.txt",
    )

    assert len(rows) == 1
    assert rows[0].pwm_match_status == PWM_STATUS_MATCHED_NEAREST


def test_distribution_statistics_use_sample_and_population_variance() -> None:
    """Stats should match pandas sample/population variance semantics."""

    _, _, _, stats = _stats_for(
        [
            _record(seconds=10.0, txt_offset=0),
            _record(seconds=12.0, txt_offset=20),
            _record(seconds=14.0, txt_offset=40),
        ]
    )

    result = stats[0]
    assert result.mean_s == pytest.approx(12.0)
    assert result.median_s == pytest.approx(12.0)
    assert result.sample_var_s2 == pytest.approx(4.0)
    assert result.sample_std_s == pytest.approx(2.0)
    assert result.population_var_ms2 / 1_000_000 == pytest.approx(8 / 3)
    assert result.population_std_ms / 1000 == pytest.approx(math.sqrt(8 / 3))


def test_single_sample_group_reports_insufficient_samples() -> None:
    """Single-sample groups should retain basic stats and blank sample SD/variance."""

    _, _, _, stats = _stats_for([_record(seconds=10.0)])

    result = stats[0]
    assert result.sample_count == 1
    assert result.mean_s == pytest.approx(10.0)
    assert result.sample_std_s is None
    assert result.sample_var_s2 is None
    assert result.distribution_status == "InsufficientSamples"


def test_distribution_chart_generation_creates_png(tmp_path: Path) -> None:
    """A normal-fit-ready group should generate a non-empty PNG chart."""

    _, rows, grouped, stats = _stats_for(
        [_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10, 11, 12, 13, 14])]
    )
    stats_by_group = {item.group_id: item for item in stats}

    generated = NormalDistributionChartGenerator(max_charts=10).generate_charts(grouped, stats_by_group, tmp_path)

    assert len(generated) == 1
    chart_path = next(iter(generated.values()))
    assert chart_path.exists()
    assert chart_path.stat().st_size > 0
    assert stats[0].chart_status == "ChartGenerated"
    assert rows


def test_zero_variance_chart_generation_does_not_crash(tmp_path: Path) -> None:
    """Identical durations should produce a safe chart and zero-variance status."""

    _, _, grouped, stats = _stats_for([_record(seconds=10, txt_offset=index * 20) for index in range(3)])
    stats_by_group = {item.group_id: item for item in stats}

    generated = NormalDistributionChartGenerator(max_charts=10).generate_charts(grouped, stats_by_group, tmp_path)

    assert stats[0].distribution_status == "ZeroVariance"
    assert "Zero variance" in stats[0].notes
    chart_path = next(iter(generated.values()))
    assert chart_path.exists()
    assert chart_path.stat().st_size > 0


def test_reference_duration_summary_and_chart_generation(tmp_path: Path) -> None:
    """Search-reference rows should have distance-free duration statistics and charts."""

    records = [
        _reference_record(seconds=value, txt_offset=index * 20, axis="Y")
        for index, value in enumerate([10.0, 12.0, 14.0])
    ]
    result = ReferenceDurationAnalyzer().analyze(records, "sample.txt")
    stats = result.stats[0]

    assert stats.axis == "Y"
    assert stats.rule_id == "search_reference"
    assert stats.sample_count == 3
    assert stats.mean_s == pytest.approx(12.0)
    assert stats.sample_std_s == pytest.approx(2.0)
    assert stats.sample_var_s2 == pytest.approx(4.0)
    assert stats.hardware_reference_evidence_found_count == 3
    assert stats.hardware_reference_evidence_missing_count == 0
    assert stats.chart_type == "Reference Duration Distribution"

    generated = NormalDistributionChartGenerator(max_charts=10).generate_reference_charts(
        result.grouped_rows,
        {stats.group_id: stats},
        tmp_path / "charts",
    )

    assert len(generated) == 1
    assert stats.chart_status == "ChartGenerated"
    assert next(iter(generated.values())).exists()


def test_axis_action_summary_combines_motion_and_reference_rows(tmp_path: Path) -> None:
    """Axis Action Summary should use the same motion/reference stats inputs as workbook sheets."""

    motion_records = [
        _record(seconds=value, txt_offset=index * 20, axis="X", rule_id="clear_motor", movement_distance=22.01)
        for index, value in enumerate([22.0, 25.0, 28.0])
    ] + [
        _record(
            seconds=value,
            txt_offset=200 + index * 20,
            axis="Y",
            rule_id="move_to_home",
            movement_distance=19.51,
        )
        for index, value in enumerate([20.0, 22.0, 24.0])
    ]
    reference_records = [
        _reference_record(seconds=value, txt_offset=400 + index * 20, axis="Y")
        for index, value in enumerate([10.0, 12.0, 14.0])
    ]
    distribution_result = DistributionAnalyzer().analyze(motion_records, "sample.txt")
    reference_result = ReferenceDurationAnalyzer().analyze(reference_records, "sample.txt")

    frames = SummaryGenerator().build_report_frames(
        motion_records + reference_records,
        [],
        distribution_result=distribution_result,
        reference_duration_result=reference_result,
    )

    rows = frames.axis_action_summary.to_dict("records")
    assert [row["Action"] for row in rows] == ["Clear Motor", "Search Reference", "Move to Home"]
    clear_row = next(row for row in rows if row["Axis"] == "X" and row["Action"] == "Clear Motor")
    ref_row = next(row for row in rows if row["Axis"] == "Y" and row["Action"] == "Search Reference")
    assert clear_row["n"] == 3
    assert clear_row["Mean (s)"] == pytest.approx(25.0)
    assert clear_row["SD (s)"] == pytest.approx(3.0)
    assert clear_row["Var (s^2)"] == pytest.approx(9.0)
    assert clear_row["CV (%)"] == pytest.approx(12.0)
    assert ref_row["n"] == 3
    assert ref_row["Mean (s)"] == pytest.approx(12.0)
    assert ref_row["SD (s)"] == pytest.approx(2.0)


def test_axis_action_summary_conditional_formatting_rules(tmp_path: Path) -> None:
    """Axis Action Summary should install the requested red/yellow CV+mean rules."""

    motion_records = [
        _record(seconds=value, txt_offset=index * 20, axis="X", rule_id="clear_motor", movement_distance=22.01)
        for index, value in enumerate([22.0, 25.0, 28.0])
    ] + [
        _record(
            seconds=value,
            txt_offset=200 + index * 20,
            axis="Y",
            rule_id="move_to_home",
            movement_distance=19.51,
        )
        for index, value in enumerate([20.0, 22.0, 24.0])
    ] + [
        _record(
            seconds=value,
            txt_offset=400 + index * 20,
            axis="Z",
            rule_id="clear_motor",
            movement_distance=50.0,
        )
        for index, value in enumerate([29.8, 30.0, 30.2])
    ]
    distribution_result = DistributionAnalyzer().analyze(motion_records, "sample.txt")
    frames = SummaryGenerator().build_report_frames(
        motion_records,
        [],
        distribution_result=distribution_result,
        reference_duration_result=ReferenceDurationAnalyzer().analyze(
            [_reference_record(seconds=value, txt_offset=600 + index * 20, axis="H") for index, value in enumerate([18.0, 19.0, 20.0])],
            "sample.txt",
        ),
    )
    output = tmp_path / "axis-summary.xlsx"
    ExcelExporter().export(
        output,
        frames,
        metadata={
            "txt_file_path": str(tmp_path / "sample.txt"),
            "log_folder_path": str(tmp_path / "logs"),
            "output_path": str(output),
            "association_strategy": "test",
            "txt_axis_event_count": 0,
            "boundary_event_count": 0,
            "log_file_count": 0,
            "pwm_profile_count": 0,
            "matched_count": len(motion_records),
            "unmatched_start_count": 0,
            "unmatched_end_count": 0,
            "parse_warning_count": 0,
            "closed_by_boundary_count": 0,
            "closed_by_new_start_count": 0,
            "initialization_failed_count": 0,
            "diagnostic_count": 0,
            "duration_warning_count": 0,
            "pwm_warning_count": 0,
            "distribution_enabled": True,
            "distribution_output_dir_full": "",
            "embed_distribution_charts": False,
        },
    )

    workbook = load_workbook(output)
    sheet = workbook["Axis Action Summary"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:9] == ["Axis", "Action", "n", "Mean (s)", "SD (s)", "Var (s^2)", "Median (s)", "Min-Max (s)", "CV (%)"]
    rules = list(sheet.conditional_formatting)
    formula_text = "\n".join(
        formula
        for conditional_range in rules
        for rule in sheet.conditional_formatting[conditional_range]
        for formula in rule.formula
    )
    assert "$I2>=2.0" in formula_text
    assert "$D2>=25.0" in formula_text
    assert "$D2>=20.0" in formula_text
    assert "$D2<25.0" in formula_text


def test_distribution_image_gallery_workbook_creation_inserts_images(tmp_path: Path) -> None:
    """The standalone gallery workbook should contain chart images plus stats and index sheets."""

    records: list[ActivityRecord] = []
    groups = [
        ("Z", "clear_motor", [10, 12, 14, 16, 18], 50),
        ("H", "move_to_home", [20, 22, 24, 26, 28], 24),
        ("X", "move_to_max", [30, 32, 34, 36, 38], 29),
    ]
    for group_index, (axis, rule_id, durations, distance) in enumerate(groups):
        records.extend(
            _record(
                seconds=value,
                txt_offset=group_index * 200 + sample_index * 20,
                axis=axis,
                rule_id=rule_id,
                start_value=None,
                movement_start_position=None,
                movement_target_position=None,
                movement_commanded_distance=distance,
                movement_distance=distance,
            )
            for sample_index, value in enumerate(durations)
        )

    _, export_result, workbook = _write_gallery_with_charts(tmp_path, records)

    assert export_result.output_path.exists()
    assert export_result.groups_with_chart_file_path_count == 3
    assert export_result.existing_chart_file_count == 3
    assert export_result.image_inserted_count == 3
    assert export_result.statistics_count == 3
    assert export_result.groups_without_charts_count == 0
    assert export_result.image_limit_skipped_count == 0
    assert {"Image Gallery", "Image Statistics", "Image Index"} <= set(workbook.sheetnames)
    assert len(workbook["Image Gallery"]._images) == 3
    gallery_values = {
        workbook["Image Gallery"].cell(row=row, column=1).value: workbook["Image Gallery"].cell(row=row, column=2).value
        for row in range(1, 25)
    }
    assert gallery_values["Movement Distance Grouping Mode"] == "bin"
    assert gallery_values["Movement Distance Bin Size"] == 0.1
    assert gallery_values["Hardware Distance Method"] == "TPOSStartEndRawDifference"
    assert gallery_values["Hardware Actual Distance Group Display"] is not None
    assert workbook["Image Statistics"].max_row == 4
    assert workbook["Image Index"].max_row == 4
    index_sheet = workbook["Image Index"]
    index_headers = [cell.value for cell in index_sheet[1]]
    anchors = [
        row[index_headers.index("Excel Anchor")]
        for row in index_sheet.iter_rows(min_row=2, values_only=True)
    ]
    anchor_rows = [int(anchor.rsplit("A", 1)[1]) for anchor in anchors]
    assert min(b - a for a, b in zip(anchor_rows, anchor_rows[1:])) >= 30
    stats_sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in stats_sheet[1]]
    z_stats = next(row for row in stats_sheet.iter_rows(min_row=2, values_only=True) if row[headers.index("Axis")] == "Z")
    assert z_stats[headers.index("Sample Count")] == 5
    assert z_stats[headers.index("Mean Duration (s)")] == pytest.approx(14.0)
    assert z_stats[headers.index("Sample SD Duration (s)")] == pytest.approx(math.sqrt(10))
    assert z_stats[headers.index("Sample Variance Duration (s^2)")] == pytest.approx(10.0)


def test_distribution_image_gallery_exports_mean_sd_and_variance(tmp_path: Path) -> None:
    """Image Statistics should expose computed mean, sample SD, and variance values."""

    _, _, workbook = _write_gallery_with_charts(
        tmp_path,
        [_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10.0, 12.0, 14.0])],
    )

    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == DISTRIBUTION_IMAGE_STATISTICS_COLUMNS
    assert "Movement Distance" not in headers
    assert "Selected Group Distance" in headers
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Sample Count")] == 3
    assert row[headers.index("Mean Duration (s)")] == pytest.approx(12.0)
    assert row[headers.index("Sample SD Duration (s)")] == pytest.approx(2.0)
    assert row[headers.index("Sample Variance Duration (s^2)")] == pytest.approx(4.0)
    assert row[headers.index("Normal Fit Mean (s)")] == pytest.approx(12.0)
    assert row[headers.index("Normal Fit Std Dev (s)")] == pytest.approx(2.0)
    assert row[headers.index("Normal Fit Variance (s^2)")] == pytest.approx(4.0)


def test_image_gallery_exports_reference_stats_without_paths(tmp_path: Path) -> None:
    """Reference charts should be labeled separately and avoid exposing full local paths."""

    records = [
        _reference_record(seconds=value, txt_offset=index * 20, axis="Y")
        for index, value in enumerate([10.0, 12.0, 14.0])
    ]
    result = ReferenceDurationAnalyzer().analyze(records, str(tmp_path / "sample.txt"))
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=10).generate_reference_charts(
        result.grouped_rows,
        stats_by_group,
        tmp_path / "charts",
    )

    export_result = DistributionImageGalleryExporter().export(tmp_path / "reference-gallery.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.image_inserted_count == 1
    stats_sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in stats_sheet[1]]
    row = next(stats_sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Chart Type")] == "Reference Duration Distribution"
    assert row[headers.index("Action")] == "Search Reference"
    assert row[headers.index("Distance")] == "Not Applicable"
    assert row[headers.index("Reference Evidence Status")] == "Found"
    assert row[headers.index("Mean Duration (s)")] == pytest.approx(12.0)
    assert row[headers.index("Sample SD Duration (s)")] == pytest.approx(2.0)
    chart_file = row[headers.index("Chart File")]
    assert chart_file.endswith(".png")
    assert "/" not in chart_file and "\\" not in chart_file

    index_sheet = workbook["Image Index"]
    index_headers = [cell.value for cell in index_sheet[1]]
    index_row = next(index_sheet.iter_rows(min_row=2, values_only=True))
    assert index_row[index_headers.index("Chart Type")] == "Reference Duration Distribution"
    assert index_row[index_headers.index("Distance")] == "Not Applicable"


def test_distribution_image_gallery_handles_skipped_chart_groups(tmp_path: Path) -> None:
    """Groups without generated charts should still appear in gallery statistics."""

    result = _distribution_result_for([_record(seconds=10.0)])
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    gallery_path = tmp_path / "skipped-gallery.xlsx"

    export_result = DistributionImageGalleryExporter().export(gallery_path, result)
    workbook = load_workbook(gallery_path, data_only=True)

    assert export_result.groups_with_chart_file_path_count == 0
    assert export_result.existing_chart_file_count == 0
    assert export_result.image_inserted_count == 0
    assert export_result.statistics_count == 1
    assert export_result.groups_without_charts_count == 1
    assert export_result.image_limit_skipped_count == 0
    assert len(workbook["Image Gallery"]._images) == 0
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Chart Status")] == "SkippedInsufficientSamples"
    assert row[headers.index("Gallery Image Status")] == "SkippedInsufficientSamples"
    gallery_summary = workbook["Image Gallery"]
    assert gallery_summary.cell(row=1, column=1).value == "Image Gallery Summary"
    assert gallery_summary.cell(row=2, column=2).value == 1
    assert "insufficient samples" in gallery_summary.cell(row=11, column=2).value


def test_distribution_image_gallery_defaults_to_message_when_no_images(tmp_path: Path) -> None:
    """Image Gallery should not become a long skipped-group sheet when no charts exist."""

    records = [
        _record(
            seconds=10,
            txt_offset=index * 20,
            start_value=None,
            movement_start_position=None,
            movement_target_position=None,
            movement_commanded_distance=10 + index,
            movement_distance=10 + index,
        )
        for index in range(10)
    ]
    result = _distribution_result_for(records)
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")

    export_result = DistributionImageGalleryExporter().export(tmp_path / "no-images.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.statistics_count == 10
    assert export_result.image_inserted_count == 0
    assert export_result.groups_without_charts_count == 10
    assert export_result.image_limit_skipped_count == 0
    gallery = workbook["Image Gallery"]
    assert gallery.cell(row=1, column=1).value == "Image Gallery Summary"
    assert gallery.cell(row=2, column=2).value == 10
    assert gallery.cell(row=5, column=2).value == 0
    assert gallery.cell(row=6, column=2).value == 10
    assert "insufficient samples" in gallery.cell(row=11, column=2).value
    assert sum(1 for row in gallery.iter_rows(values_only=True) for value in row if value == "Distribution Image") == 0
    assert workbook["Image Statistics"].max_row == 11


def test_distribution_image_gallery_no_image_message_explains_image_limit(tmp_path: Path) -> None:
    """When chart files exist but max_images is zero, the gallery summary should say so."""

    result = _distribution_result_for([_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10, 12, 14])])
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")

    export_result = DistributionImageGalleryExporter(max_images=0).export(tmp_path / "limit-zero.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.groups_with_chart_file_path_count == 1
    assert export_result.existing_chart_file_count == 1
    assert export_result.image_inserted_count == 0
    assert export_result.image_limit_skipped_count == 1
    gallery = workbook["Image Gallery"]
    assert "image limit is 0" in gallery.cell(row=11, column=2).value


def test_distribution_image_gallery_no_image_message_explains_embedding_unavailable(tmp_path: Path, monkeypatch) -> None:
    """When image support is unavailable, the gallery summary should explain that reason."""

    result = _distribution_result_for([_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10, 12, 14])])
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    exporter = DistributionImageGalleryExporter()
    monkeypatch.setattr(exporter, "_openpyxl_image_class", lambda: None)

    export_result = exporter.export(tmp_path / "embedding-unavailable.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.image_embedding_unavailable_count == 1
    assert "could not be embedded" in workbook["Image Gallery"].cell(row=11, column=2).value


def test_distribution_image_gallery_no_image_message_explains_missing_chart_files(tmp_path: Path) -> None:
    """Missing chart files should be reported without failing the gallery workbook."""

    result = _distribution_result_for([_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10, 12, 14])])
    result.stats[0].chart_file = str(tmp_path / "missing-chart.png")
    result.stats[0].chart_status = "ChartGenerated"

    export_result = DistributionImageGalleryExporter().export(tmp_path / "missing-chart.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.groups_with_chart_file_path_count == 1
    assert export_result.existing_chart_file_count == 0
    assert export_result.missing_chart_file_count == 1
    assert "could not be found" in workbook["Image Gallery"].cell(row=11, column=2).value
    stats_sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in stats_sheet[1]]
    row = next(stats_sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Gallery Image Status")] == "Chart File Missing"


def test_distribution_image_gallery_only_writes_generated_image_blocks_by_default(tmp_path: Path) -> None:
    """Generated image blocks should be present while skipped groups stay in Image Statistics."""

    chartable = []
    for group_index, axis in enumerate(["X", "Y", "Z"]):
        chartable.extend(
            _record(
                seconds=10 + sample_index,
                txt_offset=group_index * 100 + sample_index * 20,
                axis=axis,
                start_value=None,
                movement_start_position=None,
                movement_target_position=None,
                movement_commanded_distance=20 + group_index,
                movement_distance=20 + group_index,
            )
            for sample_index in range(3)
        )
    skipped = [
        _record(
            seconds=30,
            txt_offset=1000 + index * 20,
            start_value=None,
            movement_start_position=None,
            movement_target_position=None,
            movement_commanded_distance=40 + index,
            movement_distance=40 + index,
        )
        for index in range(7)
    ]

    _, export_result, workbook = _write_gallery_with_charts(tmp_path, chartable + skipped)

    assert export_result.groups_with_chart_file_path_count == 3
    assert export_result.existing_chart_file_count == 3
    assert export_result.image_inserted_count == 3
    assert export_result.statistics_count == 10
    gallery = workbook["Image Gallery"]
    assert len(gallery._images) == 3
    assert sum(1 for row in gallery.iter_rows(values_only=True) for value in row if value == "Distribution Image") == 3
    assert workbook["Image Statistics"].max_row == 11
    assert workbook["Image Index"].max_row == 4


def test_distribution_image_gallery_compact_grid_layout_inserts_multiple_images(tmp_path: Path) -> None:
    """Compact grid layout should collect generated images without vertical metadata blocks."""

    records: list[ActivityRecord] = []
    for group_index, axis in enumerate(["X", "Y", "Z", "H"]):
        records.extend(
            _record(
                seconds=10 + sample_index,
                txt_offset=group_index * 100 + sample_index * 20,
                axis=axis,
                movement_commanded_distance=20 + group_index,
                movement_distance=20 + group_index,
            )
            for sample_index in range(5)
        )

    _, export_result, workbook = _write_gallery_with_charts(tmp_path, records, layout="compact_grid")

    assert export_result.groups_with_chart_file_path_count == 4
    assert export_result.existing_chart_file_count == 4
    assert export_result.image_inserted_count == 4
    gallery = workbook["Image Gallery"]
    assert len(gallery._images) == 4
    assert sum(1 for row in gallery.iter_rows(values_only=True) for value in row if value == "Distribution Image") == 0
    index_sheet = workbook["Image Index"]
    headers = [cell.value for cell in index_sheet[1]]
    anchors = [
        row[headers.index("Excel Anchor")]
        for row in index_sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert len(set(anchors)) == 4
    assert all(anchor.startswith("Image Gallery!") for anchor in anchors)
    anchor_cells = [anchor.split("!", 1)[1] for anchor in anchors]
    assert anchor_cells[:2] == ["A3", "O3"]
    assert anchor_cells[2].startswith("A")
    assert anchor_cells[3].startswith("O")
    assert anchor_cells[2][1:] == anchor_cells[3][1:]
    assert int(anchor_cells[2][1:]) > 3
    assert gallery["A1"].value
    assert gallery["O1"].value


def test_distribution_image_gallery_can_include_skipped_blocks_when_configured(tmp_path: Path) -> None:
    """The optional compatibility mode should include skipped group blocks in Image Gallery."""

    records = [
        _record(
            seconds=10,
            txt_offset=index * 20,
            start_value=None,
            movement_start_position=None,
            movement_target_position=None,
            movement_commanded_distance=20 + index,
            movement_distance=20 + index,
        )
        for index in range(2)
    ]
    result = _distribution_result_for(records)
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")

    export_result = DistributionImageGalleryExporter(include_skipped_groups=True).export(
        tmp_path / "skipped-blocks.xlsx",
        result,
    )
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.image_inserted_count == 0
    gallery = workbook["Image Gallery"]
    assert sum(1 for row in gallery.iter_rows(values_only=True) for value in row if value == "Distribution Image") == 2


def test_distribution_image_gallery_marks_embedding_unavailable(tmp_path: Path, monkeypatch) -> None:
    """A valid chart file should not be labeled ImageInserted when image support is unavailable."""

    result = _distribution_result_for([_record(seconds=value, txt_offset=index * 20) for index, value in enumerate([10, 12, 14])])
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    exporter = DistributionImageGalleryExporter()
    monkeypatch.setattr(exporter, "_openpyxl_image_class", lambda: None)

    export_result = exporter.export(tmp_path / "no-embed.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.groups_with_chart_file_path_count == 1
    assert export_result.existing_chart_file_count == 1
    assert export_result.image_inserted_count == 0
    assert export_result.image_embedding_unavailable_count == 1
    assert export_result.image_insert_failed_count == 0
    assert len(workbook["Image Gallery"]._images) == 0
    sheet = workbook["Image Statistics"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Gallery Image Status")] == "ImageEmbeddingUnavailable"
    assert row[headers.index("Gallery Image Status")] != "ImageInserted"


def test_distribution_image_gallery_keeps_workbook_when_one_image_insert_fails(tmp_path: Path) -> None:
    """A corrupt PNG should mark one row failed without breaking other image insertions."""

    records = [
        *[
            _record(seconds=value, txt_offset=index * 20, axis="Z", rule_id="clear_motor")
            for index, value in enumerate([10, 12, 14])
        ],
        *[
            _record(
                seconds=value,
                txt_offset=100 + index * 20,
                axis="H",
                rule_id="move_to_home",
                movement_commanded_distance=24,
                movement_distance=24,
            )
            for index, value in enumerate([20, 22, 24])
        ],
    ]
    result = _distribution_result_for(records)
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    corrupt_path = Path(result.stats[0].chart_file)
    corrupt_path.write_bytes(b"not a png")

    export_result = DistributionImageGalleryExporter().export(tmp_path / "corrupt-image.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.groups_with_chart_file_path_count == 2
    assert export_result.existing_chart_file_count == 2
    assert export_result.image_inserted_count == 1
    assert export_result.image_insert_failed_count == 1
    assert len(workbook["Image Gallery"]._images) == 1
    index_sheet = workbook["Image Index"]
    headers = [cell.value for cell in index_sheet[1]]
    statuses = [
        row[headers.index("Gallery Image Status")]
        for row in index_sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert "ImageInserted" in statuses
    assert "ImageInsertFailed" in statuses
    failed_row = next(
        row
        for row in index_sheet.iter_rows(min_row=2, values_only=True)
        if row[headers.index("Gallery Image Status")] == "ImageInsertFailed"
    )
    assert failed_row[headers.index("Image Insert Error")]


def test_distribution_image_gallery_limits_inserted_images_but_keeps_all_stats(tmp_path: Path) -> None:
    """Image limits should skip extra insertions while keeping every statistics row."""

    records: list[ActivityRecord] = []
    for group_index in range(10):
        records.extend(
            _record(
                seconds=10 + sample_index,
                txt_offset=group_index * 100 + sample_index * 20,
                axis="Z",
                rule_id=f"clear_motor_{group_index}",
                movement_commanded_distance=50 + group_index,
                movement_distance=50 + group_index,
            )
            for sample_index in range(5)
        )

    _, export_result, workbook = _write_gallery_with_charts(tmp_path, records, max_images=3)

    assert export_result.groups_with_chart_file_path_count == 10
    assert export_result.existing_chart_file_count == 10
    assert export_result.image_inserted_count == 3
    assert export_result.statistics_count == 10
    assert export_result.groups_without_charts_count == 0
    assert export_result.image_limit_skipped_count == 7
    assert len(workbook["Image Gallery"]._images) == 3
    index_sheet = workbook["Image Index"]
    headers = [cell.value for cell in index_sheet[1]]
    assert headers == DISTRIBUTION_IMAGE_INDEX_COLUMNS
    statuses = [
        row[headers.index("Gallery Image Status")]
        for row in index_sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert statuses.count("ImageInserted") == 3
    assert statuses.count("SkippedDueToImageLimit") == 7


def test_distribution_image_gallery_reports_precise_chart_file_counts(tmp_path: Path) -> None:
    """Chart path, existing-file, missing-file, and image-limit counts should be distinct."""

    records = [
        _record(seconds=9, txt_offset=0, axis="N", rule_id="single_sample", movement_distance=9),
        *[
            _record(seconds=value, txt_offset=100 + index * 20, axis="Z", rule_id="group_missing", movement_distance=20)
            for index, value in enumerate([10, 12, 14])
        ],
        *[
            _record(seconds=value, txt_offset=300 + index * 20, axis="H", rule_id="group_inserted", movement_distance=30)
            for index, value in enumerate([20, 22, 24])
        ],
        *[
            _record(seconds=value, txt_offset=500 + index * 20, axis="X", rule_id="group_limited", movement_distance=40)
            for index, value in enumerate([30, 32, 34])
        ],
    ]
    result = _distribution_result_for(records)
    stats_by_group = {item.group_id: item for item in result.stats}
    NormalDistributionChartGenerator(max_charts=20).generate_charts(result.grouped_rows, stats_by_group, tmp_path / "charts")
    missing_stat = next(item for item in result.stats if item.rule_id == "group_missing")
    missing_path = Path(missing_stat.chart_file)
    missing_path.unlink()

    export_result = DistributionImageGalleryExporter(max_images=1).export(tmp_path / "precise-counts.xlsx", result)
    workbook = load_workbook(export_result.output_path, data_only=True)

    assert export_result.groups_without_charts_count == 1
    assert export_result.groups_with_chart_file_path_count == 3
    assert export_result.existing_chart_file_count == 2
    assert export_result.missing_chart_file_count == 1
    assert export_result.image_inserted_count == 1
    assert export_result.image_limit_skipped_count == 1
    assert len(workbook["Image Gallery"]._images) == 1


def test_excel_export_includes_distribution_sheets_and_chart_paths(tmp_path: Path) -> None:
    """A synthetic service run should write distribution worksheets and chart files."""

    txt_lines: list[str] = []
    hardware_lines = [
        "2026-01-01 00:00:00:000 [OUT] sample",
        "                              [N12:Z] RUN 0 80 (80)",
    ]
    for index, seconds in enumerate([10, 11, 12, 13, 14]):
        start = datetime(2026, 1, 1, 0, index, 0)
        start_motion = start + timedelta(milliseconds=1)
        end = start + timedelta(seconds=seconds)
        txt_lines.extend(
            [
                f"{start.strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[Z] min: -50.00",
                f"{start_motion.strftime('%Y-%m-%d %H:%M:%S')}:001 MCU   @[Z] start moving to home: -30.00",
                f"{end.strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[Z] motor homed",
            ]
        )
        hardware_lines.extend(_tpos_lines("Z", start_motion, end, 20.0))
    hardware_lines.append("2026-01-01 00:10:00:000 [OUT] sample")
    txt_path = _write_lines(tmp_path / "sample.txt", txt_lines)
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(log_folder / "z.log", hardware_lines)
    output_path = tmp_path / "analysis.xlsx"
    chart_dir = tmp_path / "charts"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        distribution_output_dir=chart_dir,
        max_distribution_charts=5,
    )

    assert result.distribution_group_count == 1
    assert result.distribution_raw_row_count == 5
    assert result.distribution_chart_count == 1
    assert result.distribution_image_gallery_path is not None
    assert result.distribution_image_gallery_path.exists()
    assert result.distribution_gallery_groups_with_chart_file_path_count == 1
    assert result.distribution_gallery_existing_chart_file_count == 1
    assert result.distribution_image_count == 1
    assert result.distribution_image_statistics_count == 1
    assert result.distribution_gallery_groups_without_charts_count == 0
    assert result.distribution_gallery_image_limit_skipped_count == 0
    workbook = load_workbook(output_path, data_only=True)
    assert {
        "Details",
        "Summary",
        "Distribution Summary",
        "Distribution Raw Data",
        "Distribution Charts",
        "Log Coverage Summary",
        "Log Coverage Gaps",
        "Distribution Eligibility",
        "Distribution Exclusion Summary",
        "Hardware Motion Segments",
        "Hardware Reference Events",
    } <= set(workbook.sheetnames)
    detail_headers = [cell.value for cell in workbook["Details"][1]]
    assert "Movement Actual Distance" not in detail_headers
    assert "Movement End Position" not in detail_headers
    assert "Software TXT Reported Distance" in detail_headers
    assert "Selected Movement Distance Source" in detail_headers
    segment_sheet = workbook["Hardware Motion Segments"]
    assert segment_sheet.max_row >= 6
    segment_headers = [cell.value for cell in segment_sheet[1]]
    assert "Start Line Text" in segment_headers
    assert "End Line Text" in segment_headers
    assert "Possible Duplicate Hardware Segment" in segment_headers
    assert "Effective Segment Used For Matching" in segment_headers
    assert "Matched TXT Activity Count" in segment_headers
    summary_sheet = workbook["Distribution Summary"]
    headers = [cell.value for cell in summary_sheet[1]]
    for expected in DISTRIBUTION_SUMMARY_COLUMNS:
        assert expected in headers
    header_index = {header: index for index, header in enumerate(headers)}
    row = next(summary_sheet.iter_rows(min_row=2, values_only=True))
    chart_file = row[header_index["Chart File"]]
    assert row[header_index["TXT Source File"]] == "sample.txt"
    assert row[header_index["PWM (%)"]] == 80
    assert row[header_index["Axis"]] == "Z"
    assert row[header_index["Example Hardware Start Position"]] == pytest.approx(0)
    assert row[header_index["Example Hardware Target Position"]] is None
    assert row[header_index["Example Hardware Commanded Distance"]] is None
    assert row[header_index["Hardware Actual Distance Group Value"]] == 20
    assert row[header_index["Selected Group Distance"]] == 20
    assert row[header_index["Hardware Actual Distance Rounded"]] == 20
    assert row[header_index["Hardware Distance Method"]] == "TPOSStartEndRawDifference"
    assert row[header_index["Hardware Distance Source"]] == "HardwareActualDistance"
    assert row[header_index["Rule ID"]] == "move_to_home"
    assert row[header_index["Sample Count"]] == 5
    assert row[header_index["Mean Duration (s)"]] == pytest.approx(11.999)
    assert row[header_index["Sample Std Dev Duration (s)"]] == pytest.approx(math.sqrt(2.5))
    assert row[header_index["Sample Variance Duration (s^2)"]] == pytest.approx(2.5)
    assert row[header_index["Normal Fit Mean (s)"]] == pytest.approx(11.999)
    assert row[header_index["Normal Fit Std Dev (s)"]] == pytest.approx(math.sqrt(2.5))
    assert row[header_index["Chart Status"]] == "ChartGenerated"
    assert "/" not in chart_file and "\\" not in chart_file
    assert (chart_dir / chart_file).exists()
    assert (chart_dir / chart_file).stat().st_size > 0
    gallery = load_workbook(result.distribution_image_gallery_path, data_only=True)
    assert {"Image Gallery", "Image Statistics", "Image Index"} <= set(gallery.sheetnames)
    assert len(gallery["Image Gallery"]._images) == 1
    summary_keys = [cell.value for cell in workbook["Summary"]["A"] if cell.value]
    assert "Distribution Images Inserted Into Gallery" not in summary_keys
    assert "Distribution Image Gallery Metadata Note" in summary_keys


def test_synthetic_multi_image_end_to_end_gallery_exports_images_and_sign_audits(tmp_path: Path) -> None:
    """Synthetic E2E data should create a real multi-image gallery with absolute PWM/distance grouping."""

    def fmt(value: datetime) -> str:
        return value.strftime("%Y-%m-%d %H:%M:%S:%f")[:-3]

    txt_lines: list[str] = []
    hardware_lines = [
        "2026-01-01 00:00:00:000 [OUT] synthetic",
        "                              [N12:Z] RUN 0 -80 (-80)",
        "                              [N6:H] RUN 0 80 (80)",
        "                              [N3:X] RUN 0 80 (80)",
        "2026-01-01 00:03:00:000 [OUT] synthetic",
        "                              [N12:Z] RUN 0 80 (80)",
    ]
    base_date = datetime(2026, 1, 1, 0, 0, 0)

    z_durations = [10, 12, 14, 16, 18]
    for index, seconds in enumerate(z_durations):
        base = base_date + timedelta(minutes=index, seconds=10)
        target = -29.51 if index % 2 == 0 else 29.51
        start_motion = base + timedelta(milliseconds=100)
        end = start_motion + timedelta(seconds=seconds)
        txt_lines.extend(
            [
                f"{fmt(base)} MCU   @[Z] max: 0.00",
                f"{fmt(start_motion)} MCU   @[Z] start moving to home: {target:.2f}",
                f"{fmt(end)} MCU   @[Z] motor homed",
            ]
        )
        hardware_lines.extend(_tpos_lines("Z", start_motion, end, abs(target)))

    h_durations = [20, 22, 24, 26, 28]
    for index, seconds in enumerate(h_durations):
        base = base_date + timedelta(minutes=6 + index, seconds=10)
        start_motion = base + timedelta(milliseconds=100)
        end = start_motion + timedelta(seconds=seconds)
        txt_lines.extend(
            [
                f"{fmt(base)} MCU   @[H] min: -49.03",
                f"{fmt(start_motion)} MCU   @[H] start moving to home: -25.16",
                f"{fmt(end)} MCU   @[H] motor homed",
            ]
        )
        hardware_lines.extend(_tpos_lines("H", start_motion, end, 23.87))

    x_durations = [30, 32, 34, 36, 38]
    for index, seconds in enumerate(x_durations):
        base = base_date + timedelta(minutes=12 + index, seconds=10)
        start_motion = base + timedelta(milliseconds=100)
        end = start_motion + timedelta(seconds=seconds)
        txt_lines.extend(
            [
                f"{fmt(base)} MCU   @[X] min: -32.01",
                f"{fmt(start_motion)} MCU   @[X] start moving to max pos: -3.00",
                f"{fmt(end)} MCU   @[X] motor reached max pos",
            ]
        )
        hardware_lines.extend(_tpos_lines("X", start_motion, end, 29.01))
    hardware_lines.extend(
        [
            "2026-01-01 00:20:00:000 [OUT] synthetic",
            "                              [N12:Z] RUN 0 80 (80)",
            "                              [N6:H] RUN 0 80 (80)",
            "                              [N3:X] RUN 0 80 (80)",
        ]
    )

    txt_path = _write_lines(tmp_path / "synthetic-multi.txt", txt_lines)
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(log_folder / "synthetic.log", hardware_lines)
    output_path = tmp_path / "synthetic-multi.xlsx"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        distribution_image_gallery_layout="compact_grid",
        max_distribution_charts=10,
    )

    assert result.distribution_group_count == 3
    assert result.distribution_raw_row_count == 15
    assert result.distribution_chart_count == 3
    assert result.distribution_gallery_groups_with_chart_file_path_count == 3
    assert result.distribution_gallery_existing_chart_file_count == 3
    assert result.distribution_image_count == 3
    assert result.distribution_gallery_image_insert_failed_count == 0
    assert result.distribution_image_gallery_path is not None

    workbook = load_workbook(output_path, data_only=True)
    gallery = load_workbook(result.distribution_image_gallery_path, data_only=True)
    assert len(gallery["Image Gallery"]._images) == 3
    assert gallery["Image Statistics"].max_row == 4
    assert gallery["Image Index"].max_row == 4
    stats_sheet = gallery["Image Statistics"]
    headers = [cell.value for cell in stats_sheet[1]]
    z_row = next(row for row in stats_sheet.iter_rows(min_row=2, values_only=True) if row[headers.index("Axis")] == "Z")
    assert z_row[headers.index("PWM (%)")] == 80
    assert z_row[headers.index("PWM Raw Values Seen")] == "-80; 80"
    assert z_row[headers.index("PWM Directions Seen")] == "Forward; Reverse"
    assert z_row[headers.index("PWM Direction Mixed")] is True
    assert z_row[headers.index("Hardware Actual Distance Group Display")] == "29.5"
    assert z_row[headers.index("Hardware Actual Distance Min")] == pytest.approx(29.51, abs=0.01)
    assert z_row[headers.index("Hardware Actual Distance Max")] == pytest.approx(29.51, abs=0.01)
    assert z_row[headers.index("Mean Duration (s)")] == pytest.approx(14.0)
    assert z_row[headers.index("Sample SD Duration (s)")] == pytest.approx(math.sqrt(10))
    assert z_row[headers.index("Sample Variance Duration (s^2)")] == pytest.approx(10.0)
    assert z_row[headers.index("Normal Fit Mean (s)")] == pytest.approx(14.0)
    assert z_row[headers.index("Normal Fit Std Dev (s)")] == pytest.approx(math.sqrt(10))
    z_chart_file = z_row[headers.index("Chart File")]
    assert "/" not in z_chart_file and "\\" not in z_chart_file
    assert result.distribution_output_dir is not None
    assert (result.distribution_output_dir / z_chart_file).exists()
    summary_keys = [cell.value for cell in workbook["Summary"]["A"] if cell.value]
    assert "Distribution Images Inserted Into Gallery" not in summary_keys
    assert "Distribution Image Gallery Metadata Note" in summary_keys


def test_distribution_image_gallery_path_cannot_overwrite_main_workbook(tmp_path: Path) -> None:
    """Service should reject a gallery path that equals the main workbook path."""

    txt_path = _write_lines(
        tmp_path / "collision.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:10:00:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "analysis.xlsx"

    with pytest.raises(ValueError, match="must be different from the main output workbook path"):
        LogAnalysisService().run_analysis(
            txt_file_path=txt_path,
            log_folder_path=log_folder,
            output_path=output_path,
            distribution_image_gallery_output=output_path,
        )
    assert not output_path.exists()


def test_distribution_image_gallery_output_adds_xlsx_suffix(tmp_path: Path) -> None:
    """A custom gallery path without extension should be normalized to .xlsx."""

    txt_path = _write_lines(
        tmp_path / "suffix.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:10:00:000 [OUT] sample",
        ],
    )

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=tmp_path / "analysis.xlsx",
        distribution_image_gallery_output=tmp_path / "custom-gallery",
    )

    assert result.distribution_image_gallery_path == (tmp_path / "custom-gallery.xlsx").resolve()
    assert result.distribution_image_gallery_path.exists()
    load_workbook(result.distribution_image_gallery_path, data_only=True)


def test_distribution_image_gallery_is_not_exported_when_main_workbook_fails(tmp_path: Path, monkeypatch) -> None:
    """Gallery export should not create a success artifact before the main workbook succeeds."""

    txt_path = _write_lines(
        tmp_path / "main-fail.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:03:000 MCU   @[Z] motor homed",
            "2026-01-01 00:01:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:01:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:01:04:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:10:00:000 [OUT] sample",
        ],
    )
    service = LogAnalysisService()
    output_path = tmp_path / "analysis.xlsx"
    gallery_path = tmp_path / "analysis_distribution_image_gallery.xlsx"
    gallery_called = {"value": False}

    def fail_main_export(*_args, **_kwargs):
        raise RuntimeError("main workbook failed")

    def record_gallery_call(*_args, **_kwargs):
        gallery_called["value"] = True
        raise AssertionError("gallery export should not run")

    monkeypatch.setattr(service._exporter, "export", fail_main_export)
    monkeypatch.setattr(DistributionImageGalleryExporter, "export", record_gallery_call)

    with pytest.raises(RuntimeError, match="main workbook failed"):
        service.run_analysis(txt_path, log_folder, output_path)

    assert gallery_called["value"] is False
    assert not gallery_path.exists()


def test_distribution_image_gallery_failure_keeps_main_workbook(tmp_path: Path, monkeypatch) -> None:
    """A gallery failure after main export should be reported without deleting the main workbook."""

    txt_path = _write_lines(
        tmp_path / "gallery-fail.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:03:000 MCU   @[Z] motor homed",
            "2026-01-01 00:01:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:01:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:01:04:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:10:00:000 [OUT] sample",
        ],
    )

    def fail_gallery_export(*_args, **_kwargs):
        raise RuntimeError("gallery workbook failed")

    monkeypatch.setattr(DistributionImageGalleryExporter, "export", fail_gallery_export)
    output_path = tmp_path / "analysis.xlsx"

    result = LogAnalysisService().run_analysis(txt_path, log_folder, output_path)

    assert output_path.exists()
    load_workbook(output_path, data_only=True)
    assert result.distribution_image_gallery_error == "gallery workbook failed"
    assert result.distribution_image_count == 0


def test_service_rejects_software_distance_source_options(tmp_path: Path) -> None:
    """Software distance source options should fail fast and avoid workbook output."""

    txt_lines: list[str] = []
    for index in range(2):
        base = datetime(2026, 1, 1, 0, index, 0)
        txt_lines.extend(
            [
                f"{base.strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[Y] max: 0.00",
                f"{(base + timedelta(milliseconds=1)).strftime('%Y-%m-%d %H:%M:%S')}:001 MCU   @[Y] start clearing: -19.50",
                f"{(base + timedelta(seconds=5)).strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[Y] motor cleared",
                f"{(base + timedelta(seconds=5, milliseconds=1)).strftime('%Y-%m-%d %H:%M:%S')}:001 MCU   @[Y] min: -29.51",
            ]
        )
    txt_path = _write_lines(tmp_path / "commanded.txt", txt_lines)
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "y.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N4:Y] RUN 0 80 (80)",
            "2026-01-01 00:02:00:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "commanded.xlsx"

    with pytest.raises(ValueError, match="Software distance sources are no longer supported"):
        LogAnalysisService().run_analysis(
            txt_file_path=txt_path,
            log_folder_path=log_folder,
            output_path=output_path,
            distribution_distance_source="commanded_only",
            max_distribution_charts=2,
        )

    assert not output_path.exists()


def test_service_reports_hardware_warning_counts_in_run_result_and_summary(tmp_path: Path) -> None:
    """Matched rows without reliable selected hardware distance should be visible in counts."""

    txt_path = _write_lines(
        tmp_path / "hardware-warning.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:10:00:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "hardware-warning.xlsx"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=1,
    )

    assert result.hardware_no_segment_found_count == 1
    assert result.hardware_warning_count == 1
    assert result.matched_count == 1
    assert result.matched_hardware_overlap_count == 0
    assert result.distribution_eligible_row_count == 0
    assert result.matched_missing_hardware_distance_count == 1
    assert result.matched_hardware_warning_count == 1
    assert result.distribution_excluded_missing_distance_count == 1
    assert result.distribution_excluded_multiple_reasons_count == 0

    workbook = load_workbook(output_path, data_only=True)
    details_index, details_rows = _workbook_rows(workbook, "Details")
    detail_row = next(row for row in details_rows if row[details_index["Match Status"]] == STATUS_MATCHED)
    assert detail_row[details_index["Overall Status"]] == OVERALL_STATUS_HARDWARE_WARNING
    assert detail_row[details_index["Selected Movement Distance"]] is None

    summary_rows = list(workbook["Summary"].iter_rows(min_row=2, max_col=2, values_only=True))
    metadata = {key: value for key, value in summary_rows if key}
    assert metadata["Hardware No Segment Found"] == 1
    assert metadata["Hardware Warning Count"] == 1
    assert metadata["Matched + Hardware-Overlap Rows"] == 0
    assert metadata["Matched Missing Hardware Distance"] == 1


def test_distribution_charts_sheet_explains_no_generated_charts(tmp_path: Path) -> None:
    """When all groups are below the chart threshold, Excel should explain why charts are absent."""

    txt_path = _write_lines(
        tmp_path / "single-sample.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:00:01:000 [OUT] sample",
            "                              [N12:Z] TPOS 'S' 0 0 0 0 (0)",
            "2026-01-01 00:00:02:000 [OUT] sample",
            "                              [N12:Z] TPOS 'E' 0 0 0 0 (-1761280)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "single-sample.xlsx"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=5,
    )

    workbook = load_workbook(output_path, data_only=True)
    assert result.distribution_group_count == 1
    assert result.distribution_chart_count == 0
    assert "Distribution Charts" in workbook.sheetnames
    message = workbook["Distribution Charts"].cell(row=1, column=1).value
    assert "No distribution charts were generated because all valid groups had fewer than 2 samples." == message


def test_distribution_chart_anchors_match_exported_blocks(tmp_path: Path) -> None:
    """Chart anchors in Distribution Summary should point at real chart metadata blocks."""

    txt_lines: list[str] = []
    hardware_lines = [
        "2026-01-01 00:00:00:000 [OUT] sample",
        "                              [N3:X] RUN 0 80 (80)",
        "                              [N4:Y] RUN 0 80 (80)",
        "                              [N12:Z] RUN 0 80 (80)",
    ]
    axes = [("X", -32.0, -3.0, "motor reached max pos", "start moving to max pos"), ("Y", -29.0, -10.0, "motor homed", "start moving to home"), ("Z", -50.0, -30.0, "motor homed", "start moving to home")]
    for axis_index, (axis, start_position, target, end_label, start_label) in enumerate(axes):
        for sample_index in range(2):
            base = datetime(2026, 1, 1, 0, axis_index * 10 + sample_index, 0)
            start_motion = base + timedelta(milliseconds=1)
            end_time = base + timedelta(seconds=4 + sample_index)
            txt_lines.extend(
                [
                    f"{base.strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[{axis}] min: {start_position:.2f}",
                    f"{start_motion.strftime('%Y-%m-%d %H:%M:%S')}:001 MCU   @[{axis}] {start_label}: {target:.2f}",
                    f"{end_time.strftime('%Y-%m-%d %H:%M:%S')}:000 MCU   @[{axis}] {end_label}",
                ]
            )
            hardware_lines.extend(_tpos_lines(axis, start_motion, end_time, abs(target - start_position)))
    hardware_lines.append("2026-01-01 00:30:00:000 [OUT] sample")
    txt_path = _write_lines(tmp_path / "anchors.txt", txt_lines)
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(log_folder / "all.log", hardware_lines)
    output_path = tmp_path / "anchors.xlsx"

    LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=3,
    )

    workbook = load_workbook(output_path, data_only=True)
    summary_sheet = workbook["Distribution Summary"]
    headers = [cell.value for cell in summary_sheet[1]]
    header_index = {header: index for index, header in enumerate(headers)}
    anchors = [
        row[header_index["Chart Sheet Anchor / Image ID"]]
        for row in summary_sheet.iter_rows(min_row=2, values_only=True)
        if row[header_index["Chart File"]]
    ]
    assert anchors == [
        f"Distribution Charts!A{1 + index * DISTRIBUTION_CHART_BLOCK_HEIGHT}"
        for index in range(3)
    ]
    chart_sheet = workbook["Distribution Charts"]
    for index in range(3):
        assert chart_sheet.cell(row=1 + index * DISTRIBUTION_CHART_BLOCK_HEIGHT, column=1).value == "Group ID"


def test_log_coverage_summary_warns_when_control_logs_cover_only_a_short_span(tmp_path: Path) -> None:
    """Coverage summary should make partial control-log folders visible."""

    txt_path = _write_lines(
        tmp_path / "long.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
            "2026-01-01 01:00:00:000 MCU   @[H] min: -49.03",
            "2026-01-01 01:00:01:000 MCU   @[H] start moving to home: -25.16",
            "2026-01-01 01:00:02:000 MCU   @[H] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "short.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "coverage.xlsx"

    LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=1,
    )

    workbook = load_workbook(output_path, data_only=True)
    assert "Log Coverage Summary" in workbook.sheetnames
    sheet = workbook["Log Coverage Summary"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == LOG_COVERAGE_SUMMARY_COLUMNS
    values = next(sheet.iter_rows(min_row=2, values_only=True))
    row = {header: values[index] for index, header in enumerate(headers)}
    assert row["Control Log File Count"] == 1
    assert row["Coverage Ratio (%)"] < 1
    assert row["Rows With Distribution-Accepted PWM"] == 1
    assert row["Rows With Containing-File PWM"] == 1
    assert row["Rows With No Same-Axis PWM In Folder"] >= 1
    assert row["Rows Excluded Due To Missing Hardware Distance"] >= 1
    assert row["Rows Excluded Due To Multiple Reasons"] >= 1
    assert row["Rows Excluded From Distribution Due To PWM Reliability"] == 0
    assert "Control log coverage appears incomplete" in row["Notes"]
    assert "Log Coverage Gaps" in workbook.sheetnames
    exclusion_sheet = workbook["Distribution Exclusion Summary"]
    exclusion_headers = [cell.value for cell in exclusion_sheet[1]]
    exclusion_rows = list(exclusion_sheet.iter_rows(min_row=2, values_only=True))
    exclusion_index = {header: index for index, header in enumerate(exclusion_headers)}
    assert any(
        row[exclusion_index["PWM Exclusion Reason"]] == "NoSameAxisPWMInFolder"
        and row[exclusion_index["Axis"]] == "H"
        and row[exclusion_index["PWM Match Status"]] == "NoSameAxisPWMInFolder"
        for row in exclusion_rows
    )


def test_log_coverage_summary_uses_interval_union_not_outer_range(tmp_path: Path) -> None:
    """Separated control logs should not be treated as one continuous covered window."""

    txt_path = _write_lines(
        tmp_path / "gap.txt",
        [
            "2026-01-01 09:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 09:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 09:00:02:000 MCU   @[Z] motor homed",
            "2026-01-01 10:01:00:000 MCU   @[Z] min: -50.00",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "a.log",
        [
            "2026-01-01 09:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 09:01:00:000 [OUT] sample",
        ],
    )
    _write_lines(
        log_folder / "b.log",
        [
            "2026-01-01 10:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 10:01:00:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "gap.xlsx"

    LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=1,
    )

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook["Log Coverage Summary"]
    headers = [cell.value for cell in sheet[1]]
    values = next(sheet.iter_rows(min_row=2, values_only=True))
    row = {header: values[index] for index, header in enumerate(headers)}
    assert row["Merged Control Log Interval Count"] == 2
    assert row["Covered Duration (s)"] == pytest.approx(120.0)
    assert row["Coverage Gap Count"] == 1
    assert row["Coverage Ratio (%)"] < 4
    gaps_sheet = workbook["Log Coverage Gaps"]
    gap_headers = [cell.value for cell in gaps_sheet[1]]
    assert gap_headers == LOG_COVERAGE_GAPS_COLUMNS
    gap_rows = list(gaps_sheet.iter_rows(min_row=2, values_only=True))
    assert len(gap_rows) == row["Coverage Gap Count"]
    gap_index = {header: index for index, header in enumerate(gap_headers)}
    assert gap_rows[0][gap_index["Gap Duration (s)"]] > 3500
    assert gap_rows[0][gap_index["Gap Type"]] == "BetweenControlLogs"
    assert gap_rows[0][gap_index["Nearest Previous Control Log File"]] == "a.log"
    assert gap_rows[0][gap_index["Nearest Next Control Log File"]] == "b.log"


def test_log_coverage_gaps_sheet_lists_multiple_uncovered_intervals(tmp_path: Path) -> None:
    """Multiple gaps from interval-union coverage should be individually auditable."""

    txt_path = _write_lines(
        tmp_path / "many-gaps.txt",
        [
            "2026-01-01 09:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 09:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 09:00:02:000 MCU   @[Z] motor homed",
            "2026-01-01 11:00:00:000 MCU   @[Z] min: -50.00",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "a.log",
        [
            "2026-01-01 09:10:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 09:20:00:000 [OUT] sample",
        ],
    )
    _write_lines(
        log_folder / "b.log",
        [
            "2026-01-01 10:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 10:10:00:000 [OUT] sample",
        ],
    )
    output_path = tmp_path / "many-gaps.xlsx"

    LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=1,
    )

    workbook = load_workbook(output_path, data_only=True)
    summary_sheet = workbook["Log Coverage Summary"]
    summary_headers = [cell.value for cell in summary_sheet[1]]
    summary_row = next(summary_sheet.iter_rows(min_row=2, values_only=True))
    summary = {header: summary_row[index] for index, header in enumerate(summary_headers)}
    gaps_sheet = workbook["Log Coverage Gaps"]
    gap_rows = list(gaps_sheet.iter_rows(min_row=2, values_only=True))
    gap_headers = [cell.value for cell in gaps_sheet[1]]
    gap_index = {header: index for index, header in enumerate(gap_headers)}
    gap_types = [row[gap_index["Gap Type"]] for row in gap_rows]

    assert summary["Coverage Gap Count"] == len(gap_rows)
    assert len(gap_rows) >= 3
    assert gap_types[0] == "BeforeFirstControlLog"
    assert "BetweenControlLogs" in gap_types
    assert gap_types[-1] == "AfterLastControlLog"
    assert gap_rows[0][gap_index["Nearest Next Control Log File"]] == "a.log"
    assert gap_rows[-1][gap_index["Nearest Previous Control Log File"]] == "b.log"


def test_verbose_without_trace_lines_suppresses_matcher_line_debug(tmp_path: Path) -> None:
    """CLI --verbose should keep matcher per-line debug logs hidden unless --trace-lines is supplied."""

    txt_path = _write_lines(
        tmp_path / "verbose.txt",
        [
            "2026-01-01 00:00:00:000 MCU   @[Z] min: -50.00",
            "2026-01-01 00:00:01:000 MCU   @[Z] start moving to home: -30.00",
            "2026-01-01 00:00:02:000 MCU   @[Z] motor homed",
        ],
    )
    log_folder = tmp_path / "logs"
    log_folder.mkdir()
    _write_lines(
        log_folder / "z.log",
        [
            "2026-01-01 00:00:00:000 [OUT] sample",
            "                              [N12:Z] RUN 0 80 (80)",
            "2026-01-01 00:00:10:000 [OUT] sample",
        ],
    )

    quiet_result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.log_activity_tool",
            "--txt-file",
            str(txt_path),
            "--log-folder",
            str(log_folder),
            "--output",
            str(tmp_path / "quiet.xlsx"),
            "--no-gui",
            "--verbose",
        ],
        cwd=Path(__file__).resolve().parents[1],
        text=True,
        capture_output=True,
        check=True,
    )
    trace_result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.log_activity_tool",
            "--txt-file",
            str(txt_path),
            "--log-folder",
            str(log_folder),
            "--output",
            str(tmp_path / "trace.xlsx"),
            "--no-gui",
            "--verbose",
            "--trace-lines",
        ],
        cwd=Path(__file__).resolve().parents[1],
        text=True,
        capture_output=True,
        check=True,
    )

    quiet_output = quiet_result.stdout + quiet_result.stderr
    trace_output = trace_result.stdout + trace_result.stderr
    assert "Classifying activity message" not in quiet_output
    assert "Registering start event" not in quiet_output
    assert "Resolving end value" not in quiet_output
    assert "Validating record for axis" not in quiet_output
    assert "Validating duration for axis" not in quiet_output
    assert "Validating PWM association for axis" not in quiet_output
    assert "Classifying activity message" in trace_output
    assert "Validating record for axis" in trace_output


def test_cli_rejects_old_software_distance_source(tmp_path: Path) -> None:
    """CLI should reject deprecated software distance source values before writing workbooks."""

    output_path = tmp_path / "old-source.xlsx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.log_activity_tool",
            "--txt-file",
            str(tmp_path / "missing.txt"),
            "--log-folder",
            str(tmp_path / "missing_logs"),
            "--output",
            str(output_path),
            "--distribution-distance-source",
            "commanded_only",
            "--no-gui",
        ],
        cwd=Path(__file__).resolve().parents[1],
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode != 0
    assert "invalid choice" in (result.stdout + result.stderr)
    assert "hardware_actual" in (result.stdout + result.stderr)
    assert not output_path.exists()


def _workbook_rows(workbook, sheet_name: str) -> tuple[dict[str, int], list[tuple[object, ...]]]:
    """Return workbook rows keyed by header name for regression checks."""

    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return {header: index for index, header in enumerate(headers)}, list(sheet.iter_rows(min_row=2, values_only=True))


def test_april30_partial_log_regression_keeps_distances_and_explains_coverage(tmp_path: Path) -> None:
    """The April 30 one-log sample should keep corrected distances and show partial coverage."""

    repo_root = Path(__file__).resolve().parents[1]
    txt_path = repo_root / "Log" / "UroBiopsy_20260430.txt"
    source_log = repo_root / "Log" / "RobotMovingValues" / "20260430" / "20260430_095017586_initialization.log"
    if not txt_path.exists() or not source_log.exists():
        pytest.skip("April 30 local regression sample is not available.")
    log_folder = tmp_path / "april30_one_log"
    log_folder.mkdir()
    shutil.copy2(source_log, log_folder / source_log.name)
    output_path = tmp_path / "april30-partial.xlsx"

    result = LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=20,
    )

    workbook = load_workbook(output_path, data_only=True)
    assert {
        "Details",
        "Distribution Summary",
        "Distribution Raw Data",
        "Distribution Exclusion Summary",
        "Distribution Eligibility",
        "Log Coverage Summary",
        "Log Coverage Gaps",
        "Distribution Charts",
        "Reference Duration Summary",
        "Reference Duration Charts",
        "Axis Action Summary",
        "Hardware Motion Segments",
    } <= set(workbook.sheetnames)
    details_index, details_rows = _workbook_rows(workbook, "Details")

    y_clear = next(
        row
        for row in details_rows
        if row[details_index["Axis"]] == "Y"
        and "start clearing: -19.50" in str(row[details_index["Source TXT Start Line Text"]])
        and row[details_index["Match Status"]] == STATUS_MATCHED
    )
    assert y_clear[details_index["Software TXT Commanded Distance"]] == pytest.approx(19.50)
    assert y_clear[details_index["Software TXT Reported Distance"]] == pytest.approx(29.51)
    assert y_clear[details_index["Hardware Actual Distance"]] == pytest.approx(19.51, abs=0.02)
    assert y_clear[details_index["Selected Movement Distance"]] == pytest.approx(19.51, abs=0.02)
    assert y_clear[details_index["Selected Movement Distance Source"]] == "HardwareActualDistance"

    x_clear = next(
        row
        for row in details_rows
        if row[details_index["Axis"]] == "X"
        and "start clearing: -22.00" in str(row[details_index["Source TXT Start Line Text"]])
        and row[details_index["Match Status"]] == STATUS_MATCHED
    )
    assert x_clear[details_index["Software TXT Commanded Distance"]] == pytest.approx(22.00)
    assert x_clear[details_index["Software TXT Reported Distance"]] == pytest.approx(32.01)
    assert x_clear[details_index["Hardware Actual Distance"]] == pytest.approx(22.01, abs=0.02)
    assert x_clear[details_index["Selected Movement Distance"]] == pytest.approx(22.01, abs=0.02)

    h_home = next(
        row
        for row in details_rows
        if row[details_index["Axis"]] == "H"
        and "start moving to home: -25.16" in str(row[details_index["Source TXT Start Line Text"]])
        and row[details_index["Match Status"]] == STATUS_MATCHED
    )
    assert h_home[details_index["Selected Movement Distance"]] == pytest.approx(23.87, abs=0.05)

    assert result.distribution_group_count == 18
    assert result.distribution_raw_row_count == 18
    assert result.distribution_chart_count == 0
    assert result.distribution_image_gallery_path is not None
    gallery = load_workbook(result.distribution_image_gallery_path, data_only=True)
    expected_gallery_stats_rows = 1 + result.distribution_group_count + result.reference_distribution_group_count
    assert gallery["Image Statistics"].max_row == expected_gallery_stats_rows
    assert len(gallery["Image Gallery"]._images) == result.reference_distribution_chart_count
    assert result.reference_distribution_group_count == 8
    assert result.reference_distribution_chart_count == 8
    chart_message = workbook["Distribution Charts"].cell(row=1, column=1).value
    assert "No distribution charts were generated because all valid groups had fewer than 2 samples." == chart_message

    coverage_index, coverage_rows = _workbook_rows(workbook, "Log Coverage Summary")
    coverage = coverage_rows[0]
    assert coverage[coverage_index["Coverage Ratio (%)"]] < 10
    gaps_index, gaps_rows = _workbook_rows(workbook, "Log Coverage Gaps")
    assert len(gaps_rows) >= 2
    assert gaps_rows[0][gaps_index["Gap Type"]] == "BeforeFirstControlLog"
    assert gaps_rows[-1][gaps_index["Gap Type"]] == "AfterLastControlLog"


def test_april10_bad_duration_pair_remains_rejected(tmp_path: Path) -> None:
    """The old April 10 Z clear across-initialization false match must stay fixed."""

    repo_root = Path(__file__).resolve().parents[1]
    txt_path = repo_root / "Log" / "UroBiopsy_20260410.txt"
    log_folder = repo_root / "Log" / "RobotMovingValues" / "20260410"
    if not txt_path.exists() or not log_folder.exists():
        pytest.skip("April 10 local regression sample is not available.")
    output_path = tmp_path / "april10.xlsx"

    LogAnalysisService().run_analysis(
        txt_file_path=txt_path,
        log_folder_path=log_folder,
        output_path=output_path,
        max_distribution_charts=10,
    )

    workbook = load_workbook(output_path, data_only=True)
    details_index, details_rows = _workbook_rows(workbook, "Details")
    bad_matches = [
        row
        for row in details_rows
        if row[details_index["Match Status"]] == STATUS_MATCHED
        and "2026-04-10 08:37:43:959 MCU   @[Z] start clearing: -50.00"
        in str(row[details_index["Source TXT Start Line Text"]])
        and "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared"
        in str(row[details_index["Source TXT End Line Text"]])
    ]
    assert bad_matches == []

    correct = next(
        row
        for row in details_rows
        if row[details_index["Match Status"]] == STATUS_MATCHED
        and "2026-04-10 09:20:23:541 MCU   @[Z] start clearing: -50.00"
        in str(row[details_index["Source TXT Start Line Text"]])
        and "2026-04-10 09:20:46:317 MCU   @[Z] motor cleared"
        in str(row[details_index["Source TXT End Line Text"]])
    )
    assert correct[details_index["Duration (ms)"]] == 22776
    assert correct[details_index["Duration (s)"]] == pytest.approx(22.776)
