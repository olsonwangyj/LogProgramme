"""Tests for DHR batch discovery, selection, and cross-iSRT summaries."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from backend.log_axis_activity_analyzer.batch_runner import DHRBatchProcessor, DHRCaseProcessingRecord, safe_case_id


def _write_biopsy_log(path: Path, initialization_count: int) -> Path:
    """Write a small UroBiopsy TXT with repeated completed initialization markers."""

    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        f"2025-01-01 00:{index // 60:02d}:{index % 60:02d}:000 MCU   @robot initialization done"
        for index in range(initialization_count)
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _log_folder(dhr_root: Path, relative_case: str) -> Path:
    """Return the canonical Log folder for a fixture case."""

    return dhr_root / relative_case / "Biobot" / "System" / "Log"


def _write_main_workbook_with_raw_samples(path: Path, samples: list[tuple[str, str, str, float]]) -> Path:
    """Write a minimal main workbook with Distribution Raw Data rows."""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Distribution Raw Data"
    headers = ["Axis", "Rule ID", "Action Label", "Duration (s)"]
    sheet.append(headers)
    for axis, rule_id, action_label, duration_s in samples:
        sheet.append([axis, rule_id, action_label, duration_s])
    workbook.save(path)
    return path


def _processed_record(case_id: str, main_workbook: Path | str) -> DHRCaseProcessingRecord:
    """Build a processed case record for cross-summary export tests."""

    return DHRCaseProcessingRecord(
        case_id=case_id,
        relative_case_folder=case_id,
        processing_status="Processed",
        main_workbook=str(main_workbook),
    )


def test_batch_selects_newest_valid_urobiopsy_txt(tmp_path: Path) -> None:
    """When the newest TXT has enough data, it should be selected."""

    dhr_root = tmp_path / "DHR"
    log_folder = _log_folder(dhr_root, "iSRT000001/C")
    _write_biopsy_log(log_folder / "UroBiopsy_20250101.txt", 10)
    newest = _write_biopsy_log(log_folder / "UroBiopsy_20250102.txt", 25)
    processor = DHRBatchProcessor()

    candidate = processor.discover_cases(dhr_root)[0]
    selection, skip_reason = processor.select_input_for_case(candidate)

    assert skip_reason == ""
    assert selection is not None
    assert selection.txt_file == newest.resolve()
    assert selection.initialization_count == 25


def test_batch_selects_older_valid_when_newest_is_insufficient(tmp_path: Path) -> None:
    """The selector should try older files when the newest TXT lacks enough initialization cycles."""

    dhr_root = tmp_path / "DHR"
    log_folder = _log_folder(dhr_root, "iSRT000001/C")
    older_valid = _write_biopsy_log(log_folder / "UroBiopsy_20250102.txt", 25)
    _write_biopsy_log(log_folder / "UroBiopsy_20250103.txt", 5)
    processor = DHRBatchProcessor()

    candidate = processor.discover_cases(dhr_root)[0]
    selection, skip_reason = processor.select_input_for_case(candidate)

    assert skip_reason == ""
    assert selection is not None
    assert selection.txt_file == older_valid.resolve()
    assert selection.initialization_count == 25


def test_batch_skips_case_when_no_txt_has_enough_initializations(tmp_path: Path) -> None:
    """Cases below the initialization threshold should be skipped with a clear reason."""

    dhr_root = tmp_path / "DHR"
    log_folder = _log_folder(dhr_root, "iSRT000001/C")
    _write_biopsy_log(log_folder / "UroBiopsy_20250101.txt", 10)
    _write_biopsy_log(log_folder / "UroBiopsy_20250102.txt", 19)
    processor = DHRBatchProcessor()

    candidate = processor.discover_cases(dhr_root)[0]
    selection, skip_reason = processor.select_input_for_case(candidate)

    assert selection is None
    assert "initialization_count >= 20" in skip_reason
    result = processor.run(dhr_root, tmp_path / "out")
    assert result.records[0].processing_status == "SkippedInsufficientInitializationCount"


def test_cross_summary_uses_pooled_mean_not_mean_of_means(tmp_path: Path) -> None:
    """Cross-iSRT summary should pool raw durations before calculating the mean."""

    output_root = tmp_path / "DHR_output"
    case_a = _write_main_workbook_with_raw_samples(
        output_root / "case_a" / "case_a.xlsx",
        [
            ("Z", "clear_motor", "start clearing -> motor cleared", 10.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 10.0),
        ],
    )
    case_b = _write_main_workbook_with_raw_samples(
        output_root / "case_b" / "case_b.xlsx",
        [
            ("Z", "clear_motor", "start clearing -> motor cleared", 20.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 20.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 20.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 20.0),
        ],
    )
    records = [
        _processed_record("case_a", case_a.relative_to(output_root)),
        _processed_record("case_b", case_b.relative_to(output_root)),
    ]

    cross_path = DHRBatchProcessor().export_cross_summary(output_root / "cross_iSRT_summary.xlsx", output_root, records)
    sheet = load_workbook(cross_path, data_only=True)["Overall Axis Action Summary"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))

    assert row[headers.index("Action")] == "Clear Motor"
    assert row[headers.index("n")] == 6
    assert row[headers.index("Mean (s)")] == pytest.approx(16.6666667)
    assert row[headers.index("Mean (s)")] != pytest.approx(15.0)


def test_cross_summary_pools_median_and_iqr(tmp_path: Path) -> None:
    """Median and IQR should be calculated from all pooled samples."""

    output_root = tmp_path / "DHR_output"
    case_a = _write_main_workbook_with_raw_samples(
        output_root / "case_a" / "case_a.xlsx",
        [
            ("Z", "clear_motor", "start clearing -> motor cleared", 1.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 100.0),
        ],
    )
    case_b = _write_main_workbook_with_raw_samples(
        output_root / "case_b" / "case_b.xlsx",
        [
            ("Z", "clear_motor", "start clearing -> motor cleared", 2.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 3.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 4.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 5.0),
            ("Z", "clear_motor", "start clearing -> motor cleared", 6.0),
        ],
    )
    records = [
        _processed_record("case_a", case_a.relative_to(output_root)),
        _processed_record("case_b", case_b.relative_to(output_root)),
    ]

    cross_path = DHRBatchProcessor().export_cross_summary(output_root / "cross_iSRT_summary.xlsx", output_root, records)
    sheet = load_workbook(cross_path, data_only=True)["Overall Axis Action Summary"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))

    assert row[headers.index("n")] == 7
    assert row[headers.index("Median (s)")] == pytest.approx(4.0)
    assert row[headers.index("IQR (s)")] == pytest.approx(3.0)
    assert row[headers.index("Min–Max (s)")] == "1.000–100.000"


def test_cross_summary_excludes_search_reference_and_keeps_move_to_home(tmp_path: Path) -> None:
    """The cross-iSRT clean summary should filter Search Reference only."""

    output_root = tmp_path / "DHR_output"
    main = _write_main_workbook_with_raw_samples(
        output_root / "case_a" / "case_a.xlsx",
        [
            ("Z", "search_reference", "start searching reference -> reference found", 1.0),
            ("Z", "move_to_home", "start moving to home -> motor homed", 21.0),
            ("Z", "move_to_home", "start moving to home -> motor homed", 23.0),
        ],
    )
    records = [_processed_record("case_a", main.relative_to(output_root))]

    cross_path = DHRBatchProcessor().export_cross_summary(output_root / "cross_iSRT_summary.xlsx", output_root, records)
    sheet = load_workbook(cross_path, data_only=True)["Overall Axis Action Summary"]
    headers = [cell.value for cell in sheet[1]]
    actions = [row[headers.index("Action")] for row in sheet.iter_rows(min_row=2, values_only=True)]

    assert actions == ["Move to Home"]


def test_cross_summary_workbook_structure(tmp_path: Path) -> None:
    """The pooled summary must be the first sheet, with audit sheets after it."""

    output_root = tmp_path / "DHR_output"
    main = _write_main_workbook_with_raw_samples(
        output_root / "case_a" / "case_a.xlsx",
        [("N", "move_to_max", "start moving to max pos -> motor reached max pos", 12.0)],
    )
    records = [_processed_record("case_a", main.relative_to(output_root))]

    cross_path = DHRBatchProcessor().export_cross_summary(output_root / "cross_iSRT_summary.xlsx", output_root, records)
    workbook = load_workbook(cross_path, data_only=True)

    assert workbook.sheetnames == [
        "Overall Axis Action Summary",
        "Case Processing Summary",
        "Case Contribution Details",
    ]
    sheet = workbook["Overall Axis Action Summary"]
    assert [cell.value for cell in sheet[1]] == [
        "Axis",
        "Action",
        "n",
        "Mean (s)",
        "SD (s)",
        "Var (s²)",
        "Median (s)",
        "IQR (s)",
        "Min–Max (s)",
        "CV (%)",
    ]
    assert len(sheet.conditional_formatting) > 0


def test_batch_run_writes_pooled_cross_summary(tmp_path: Path) -> None:
    """Batch mode should create a pooled cross-iSRT summary workbook."""

    dhr_root = tmp_path / "DHR"
    for relative_case in ("iSRT000001/C", "iSRT000002/23Apr/C"):
        log_folder = _log_folder(dhr_root, relative_case)
        _write_biopsy_log(log_folder / "UroBiopsy_20250102.txt", 25)

    class FakeService:
        def run_analysis(self, **kwargs):
            output_path = Path(kwargs["output_path"])
            output_path.parent.mkdir(parents=True, exist_ok=True)
            summary_path = output_path.with_name(f"{output_path.stem}_summary.xlsx")
            gallery_path = output_path.with_name(f"{output_path.stem}_distribution_image_gallery.xlsx")
            chart_folder = Path(kwargs["distribution_output_dir"])
            chart_folder.mkdir(parents=True, exist_ok=True)

            _write_main_workbook_with_raw_samples(
                output_path,
                [("P", "move_to_home", "start moving to home -> motor homed", 21.0)],
            )
            gallery = Workbook()
            gallery.active.title = "Image Gallery"
            gallery.create_sheet("Image Statistics")
            gallery.create_sheet("Image Index")
            gallery.save(gallery_path)
            _write_main_workbook_with_raw_samples(summary_path, [])
            return SimpleNamespace(
                output_path=output_path,
                summary_output_path=summary_path,
                distribution_image_gallery_path=gallery_path,
                distribution_output_dir=chart_folder,
            )

    processor = DHRBatchProcessor(service_factory=lambda: FakeService())
    result = processor.run(dhr_root, tmp_path / "DHR_output")

    assert result.discovered_count == 2
    assert result.processed_count == 2
    assert result.skipped_count == 0
    assert result.cross_summary_path.exists()
    workbook = load_workbook(result.cross_summary_path, data_only=True)
    assert workbook.sheetnames[:3] == [
        "Overall Axis Action Summary",
        "Case Processing Summary",
        "Case Contribution Details",
    ]
    summary_sheet = workbook["Overall Axis Action Summary"]
    headers = [cell.value for cell in summary_sheet[1]]
    row = next(summary_sheet.iter_rows(min_row=2, values_only=True))

    assert row[headers.index("Action")] == "Move to Home"
    assert row[headers.index("n")] == 2


def test_safe_case_id_uses_relative_case_folder_parts() -> None:
    """Output folder names should be stable and collision-resistant across nested case paths."""

    assert safe_case_id(Path("iSRT300040") / "23Apr" / "C") == "iSRT300040_23Apr_C"


def test_cli_dhr_root_runs_batch_mode(monkeypatch, tmp_path: Path) -> None:
    """--dhr-root should call the batch processor instead of single-case path resolution."""

    from app import log_activity_tool

    dhr_root = tmp_path / "DHR"
    output_root = tmp_path / "DHR_output"
    dhr_root.mkdir()
    calls: dict[str, object] = {}

    class FakeBatchProcessor:
        def __init__(self, logger=None):
            calls["logger"] = logger

        def run(self, **kwargs):
            calls["run_kwargs"] = kwargs
            return SimpleNamespace(
                dhr_root=dhr_root,
                output_root=output_root,
                cross_summary_path=output_root / "cross_iSRT_summary.xlsx",
                discovered_count=1,
                processed_count=1,
                skipped_count=0,
                failed_count=0,
                records=[],
            )

    monkeypatch.setattr(log_activity_tool, "DHRBatchProcessor", FakeBatchProcessor)

    assert log_activity_tool.main(
        [
            "--dhr-root",
            str(dhr_root),
            "--batch-output",
            str(output_root),
            "--distribution-image-gallery",
            "--no-file-picker",
        ]
    ) == 0
    assert calls["run_kwargs"]["dhr_root"] == str(dhr_root)
    assert calls["run_kwargs"]["batch_output"] == str(output_root)
    assert calls["run_kwargs"]["recursive"] is True
