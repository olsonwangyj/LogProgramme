"""Batch processing for DHR iSRT folders and cross-case summary export."""

from __future__ import annotations

import logging
import math
import re
import statistics
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import (
    AXIS_SUMMARY_MEAN_RED_THRESHOLD,
    AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD,
    DISTRIBUTION_IMAGE_GALLERY_LAYOUT,
    DISTRIBUTION_DISTANCE_GROUPING_MODE,
    ENABLE_DISTRIBUTION_ANALYSIS,
    EMBED_DISTRIBUTION_CHARTS_IN_EXCEL,
    EXPORT_DISTRIBUTION_IMAGE_GALLERY,
    HARDWARE_DISTANCE_SOURCE,
    MOVEMENT_DISTANCE_BIN_SIZE,
    MOVEMENT_DISTANCE_ROUND_DIGITS,
    NORMAL_DISTRIBUTION_MAX_CHARTS,
    OVERALL_AXIS_ACTION_SUMMARY_COLUMNS,
    OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME,
)
from .file_loader import TextFileLoader
from .service import LogAnalysisService


CASE_PROCESSING_SUMMARY_COLUMNS = [
    "Case ID",
    "Relative Case Folder",
    "Selected TXT File",
    "Selected TXT Date",
    "Hardware Log Folder",
    "Initialization Count",
    "Processing Status",
    "Skip Reason",
    "Main Workbook",
    "Summary Workbook",
    "Gallery Workbook",
    "Chart Folder",
]

CASE_CONTRIBUTION_DETAILS_COLUMNS = [
    "Case ID",
    *OVERALL_AXIS_ACTION_SUMMARY_COLUMNS,
]

INITIALIZATION_DONE_TOKEN = "@robot initialization done"
DEFAULT_MIN_INITIALIZATION_COUNT = 20
UROBIOPSY_DATE_PATTERN = re.compile(r"UroBiopsy_(?P<date>\d{8})(?:_(?P<time>\d{6}))?\.txt$", re.IGNORECASE)


@dataclass(frozen=True)
class DHRCaseCandidate:
    """One logical iSRT case folder and its associated hardware log folders."""

    case_folder: Path
    relative_case_folder: Path
    log_folders: tuple[Path, ...]

    @property
    def case_id(self) -> str:
        """Return a filesystem-safe case ID from the DHR-relative path."""

        return safe_case_id(self.relative_case_folder)


@dataclass(frozen=True)
class SelectedDHRInput:
    """Selected TXT input and associated hardware log folder for one case."""

    txt_file: Path
    txt_date: str
    hardware_log_folder: Path
    initialization_count: int


@dataclass
class DHRCaseProcessingRecord:
    """Batch processing status for one discovered case."""

    case_id: str
    relative_case_folder: str
    selected_txt_file: str = ""
    selected_txt_date: str = ""
    hardware_log_folder: str = ""
    initialization_count: int = 0
    processing_status: str = ""
    skip_reason: str = ""
    main_workbook: str = ""
    summary_workbook: str = ""
    gallery_workbook: str = ""
    chart_folder: str = ""

    def to_row(self) -> dict[str, object]:
        """Return this processing record as a workbook row mapping."""

        return {
            "Case ID": self.case_id,
            "Relative Case Folder": self.relative_case_folder,
            "Selected TXT File": self.selected_txt_file,
            "Selected TXT Date": self.selected_txt_date,
            "Hardware Log Folder": self.hardware_log_folder,
            "Initialization Count": self.initialization_count,
            "Processing Status": self.processing_status,
            "Skip Reason": self.skip_reason,
            "Main Workbook": self.main_workbook,
            "Summary Workbook": self.summary_workbook,
            "Gallery Workbook": self.gallery_workbook,
            "Chart Folder": self.chart_folder,
        }


@dataclass
class DHRBatchRunResult:
    """Summary of a full DHR batch run."""

    dhr_root: Path
    output_root: Path
    cross_summary_path: Path
    discovered_count: int
    processed_count: int
    skipped_count: int
    failed_count: int
    records: list[DHRCaseProcessingRecord] = field(default_factory=list)


class DHRBatchProcessor:
    """Run the existing single-case analysis pipeline over discovered DHR cases."""

    def __init__(
        self,
        service_factory: Callable[[], object] | None = None,
        logger: logging.Logger | None = None,
    ) -> None:
        """Initialize the batch processor with optional service injection for tests."""

        self._logger = logger or logging.getLogger(self.__class__.__name__)
        self._service_factory = service_factory or (
            lambda: LogAnalysisService(self._logger.getChild("service"))
        )
        self._loader = TextFileLoader(self._logger.getChild("loader"))

    def discover_cases(self, dhr_root: Path | str) -> list[DHRCaseCandidate]:
        """Discover logical iSRT case folders beneath a DHR root.

        A case is the nearest folder above ``Biobot/System*/Log``. Multiple
        historical ``System`` snapshots under one case are grouped together so
        each iSRT case emits one output set.
        """

        root = Path(dhr_root).expanduser().resolve()
        if not root.is_dir():
            raise FileNotFoundError(f"DHR root does not exist: {root}")
        grouped: dict[Path, list[Path]] = {}
        for log_dir in root.rglob("Log"):
            if not log_dir.is_dir():
                continue
            system_dir = log_dir.parent
            biobot_dir = system_dir.parent
            if biobot_dir.name.casefold() != "biobot":
                continue
            if not system_dir.name.casefold().startswith("system"):
                continue
            case_folder = biobot_dir.parent.resolve()
            grouped.setdefault(case_folder, []).append(log_dir.resolve())
        candidates = [
            DHRCaseCandidate(
                case_folder=case_folder,
                relative_case_folder=case_folder.relative_to(root),
                log_folders=tuple(sorted(log_folders, key=lambda value: str(value).casefold())),
            )
            for case_folder, log_folders in grouped.items()
        ]
        return sorted(candidates, key=lambda candidate: str(candidate.relative_case_folder).casefold())

    def select_input_for_case(
        self,
        candidate: DHRCaseCandidate,
        min_initialization_count: int = DEFAULT_MIN_INITIALIZATION_COUNT,
        encoding_txt: str | None = None,
    ) -> tuple[SelectedDHRInput | None, str]:
        """Select the newest UroBiopsy TXT with enough completed initialization cycles."""

        txt_files = self._candidate_txt_files(candidate.log_folders)
        if not txt_files:
            return None, "No UroBiopsy_*.txt files found in associated Log folders."

        best_count = 0
        best_txt = txt_files[0]
        for txt_file in txt_files:
            initialization_count = self.count_initialization_cycles(txt_file, encoding_txt)
            if initialization_count > best_count:
                best_count = initialization_count
                best_txt = txt_file
            if initialization_count >= min_initialization_count:
                return (
                    SelectedDHRInput(
                        txt_file=txt_file,
                        txt_date=txt_date_from_name(txt_file),
                        hardware_log_folder=txt_file.parent,
                        initialization_count=initialization_count,
                    ),
                    "",
                )
        return (
            None,
            (
                f"No UroBiopsy TXT reached initialization_count >= {min_initialization_count}; "
                f"best was {best_count} in {best_txt.name}."
            ),
        )

    def count_initialization_cycles(self, txt_file: Path | str, encoding_txt: str | None = None) -> int:
        """Count completed initialization cycles using ``@robot initialization done`` lines."""

        iterator, _encoding = self._loader.iter_lines(txt_file, encoding_txt)
        token = INITIALIZATION_DONE_TOKEN.casefold()
        return sum(1 for _line_number, line in iterator if token in line.casefold())

    def run(
        self,
        dhr_root: Path | str,
        batch_output: Path | str | None = None,
        *,
        min_initialization_count: int = DEFAULT_MIN_INITIALIZATION_COUNT,
        encoding_txt: str | None = None,
        encoding_logs: str | None = None,
        association_strategy: str = "same_file_then_nearest",
        recursive: bool = True,
        enable_distribution_analysis: bool = ENABLE_DISTRIBUTION_ANALYSIS,
        max_distribution_charts: int = NORMAL_DISTRIBUTION_MAX_CHARTS,
        embed_distribution_charts: bool = EMBED_DISTRIBUTION_CHARTS_IN_EXCEL,
        movement_distance_round_digits: int = MOVEMENT_DISTANCE_ROUND_DIGITS,
        distribution_distance_source: str = HARDWARE_DISTANCE_SOURCE,
        distribution_distance_grouping_mode: str = DISTRIBUTION_DISTANCE_GROUPING_MODE,
        movement_distance_bin_size: float = MOVEMENT_DISTANCE_BIN_SIZE,
        distribution_allow_nearest_pwm: bool = False,
        distribution_allow_latest_before_pwm: bool = False,
        distribution_allow_carry_forward_pwm: bool = False,
        distribution_allow_nearest_hardware_segment: bool = False,
        distribution_allow_ambiguous_hardware_segment: bool = False,
        allow_pwm_carry_forward: bool = False,
        export_distribution_image_gallery: bool | None = EXPORT_DISTRIBUTION_IMAGE_GALLERY,
        distribution_image_gallery_layout: str = DISTRIBUTION_IMAGE_GALLERY_LAYOUT,
    ) -> DHRBatchRunResult:
        """Process every discovered DHR case and write a cross-iSRT summary workbook."""

        root = Path(dhr_root).expanduser().resolve()
        output_root = (
            Path(batch_output).expanduser().resolve()
            if batch_output is not None
            else (root.parent / f"{root.name}_output").resolve()
        )
        output_root.mkdir(parents=True, exist_ok=True)
        candidates = self.discover_cases(root)
        records: list[DHRCaseProcessingRecord] = []

        for candidate in candidates:
            record = DHRCaseProcessingRecord(
                case_id=candidate.case_id,
                relative_case_folder=self._relative_to_root(candidate.case_folder, root),
            )
            selection, skip_reason = self.select_input_for_case(
                candidate,
                min_initialization_count=min_initialization_count,
                encoding_txt=encoding_txt,
            )
            if selection is None:
                record.processing_status = (
                    "SkippedNoUroBiopsyTxt"
                    if skip_reason.startswith("No UroBiopsy_*.txt")
                    else "SkippedInsufficientInitializationCount"
                )
                record.skip_reason = skip_reason
                records.append(record)
                self._logger.info("Skipping %s: %s", candidate.case_id, skip_reason)
                continue

            case_output_dir = output_root / candidate.case_id
            case_output_dir.mkdir(parents=True, exist_ok=True)
            main_output = case_output_dir / f"{candidate.case_id}.xlsx"
            chart_folder = case_output_dir / f"{candidate.case_id}_distribution_charts"
            record.selected_txt_file = self._relative_to_root(selection.txt_file, root)
            record.selected_txt_date = selection.txt_date
            record.hardware_log_folder = self._relative_to_root(selection.hardware_log_folder, root)
            record.initialization_count = selection.initialization_count
            record.main_workbook = self._relative_to_root(main_output, output_root)
            record.summary_workbook = self._relative_to_root(case_output_dir / f"{candidate.case_id}_summary.xlsx", output_root)
            record.gallery_workbook = self._relative_to_root(
                case_output_dir / f"{candidate.case_id}_distribution_image_gallery.xlsx",
                output_root,
            )
            record.chart_folder = self._relative_to_root(chart_folder, output_root)

            try:
                self._logger.info(
                    "Processing %s with %s (%s initialization cycles)",
                    candidate.case_id,
                    selection.txt_file,
                    selection.initialization_count,
                )
                service = self._service_factory()
                result = service.run_analysis(
                    txt_file_path=selection.txt_file,
                    log_folder_path=selection.hardware_log_folder,
                    output_path=main_output,
                    encoding_txt=encoding_txt,
                    encoding_logs=encoding_logs,
                    association_strategy=association_strategy,
                    recursive=recursive,
                    enable_distribution_analysis=enable_distribution_analysis,
                    distribution_output_dir=chart_folder,
                    max_distribution_charts=max_distribution_charts,
                    embed_distribution_charts=embed_distribution_charts,
                    movement_distance_round_digits=movement_distance_round_digits,
                    distribution_distance_source=distribution_distance_source,
                    distribution_distance_grouping_mode=distribution_distance_grouping_mode,
                    movement_distance_bin_size=movement_distance_bin_size,
                    distribution_allow_nearest_pwm=distribution_allow_nearest_pwm,
                    distribution_allow_latest_before_pwm=distribution_allow_latest_before_pwm,
                    distribution_allow_carry_forward_pwm=distribution_allow_carry_forward_pwm,
                    distribution_allow_nearest_hardware_segment=distribution_allow_nearest_hardware_segment,
                    distribution_allow_ambiguous_hardware_segment=distribution_allow_ambiguous_hardware_segment,
                    allow_pwm_carry_forward=allow_pwm_carry_forward,
                    export_distribution_image_gallery=export_distribution_image_gallery,
                    distribution_image_gallery_layout=distribution_image_gallery_layout,
                )
                record.processing_status = "Processed"
                record.main_workbook = self._relative_to_root(result.output_path, output_root)
                if getattr(result, "summary_output_path", None):
                    record.summary_workbook = self._relative_to_root(result.summary_output_path, output_root)
                if getattr(result, "distribution_image_gallery_path", None):
                    record.gallery_workbook = self._relative_to_root(result.distribution_image_gallery_path, output_root)
                if getattr(result, "distribution_output_dir", None):
                    record.chart_folder = self._relative_to_root(result.distribution_output_dir, output_root)
            except Exception as exc:  # pragma: no cover - covered through integration and fake-service tests.
                self._logger.exception("Failed to process %s: %s", candidate.case_id, exc)
                record.processing_status = "Failed"
                record.skip_reason = str(exc)
            records.append(record)

        cross_summary_path = output_root / "cross_iSRT_summary.xlsx"
        self.export_cross_summary(cross_summary_path, output_root, records)
        processed_count = sum(record.processing_status == "Processed" for record in records)
        skipped_count = sum(record.processing_status.startswith("Skipped") for record in records)
        failed_count = sum(record.processing_status == "Failed" for record in records)
        return DHRBatchRunResult(
            dhr_root=root,
            output_root=output_root,
            cross_summary_path=cross_summary_path,
            discovered_count=len(candidates),
            processed_count=processed_count,
            skipped_count=skipped_count,
            failed_count=failed_count,
            records=records,
        )

    def export_cross_summary(
        self,
        output_path: Path | str,
        output_root: Path | str,
        records: list[DHRCaseProcessingRecord],
    ) -> Path:
        """Write the cross-iSRT summary workbook from pooled raw duration samples."""

        path = Path(output_path).expanduser().resolve()
        root = Path(output_root).expanduser().resolve()
        samples = self._pooled_axis_action_samples(root, records)
        overall_rows = self._summary_rows_from_samples(samples)
        contribution_rows = self._case_contribution_rows_from_samples(samples)

        workbook = Workbook()
        overall_sheet = workbook.active
        overall_sheet.title = OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME
        self._write_rows(overall_sheet, OVERALL_AXIS_ACTION_SUMMARY_COLUMNS, overall_rows)

        processing_sheet = workbook.create_sheet("Case Processing Summary")
        self._write_rows(processing_sheet, CASE_PROCESSING_SUMMARY_COLUMNS, [record.to_row() for record in records])

        contribution_sheet = workbook.create_sheet("Case Contribution Details")
        self._write_rows(contribution_sheet, CASE_CONTRIBUTION_DETAILS_COLUMNS, contribution_rows)

        for sheet in workbook.worksheets:
            self._format_sheet(
                sheet,
                mean_conditional_formatting=sheet.title == OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME,
            )
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def _candidate_txt_files(self, log_folders: tuple[Path, ...]) -> list[Path]:
        """Return candidate UroBiopsy TXT files newest first."""

        txt_files: list[Path] = []
        for log_folder in log_folders:
            txt_files.extend(path.resolve() for path in log_folder.glob("UroBiopsy_*.txt") if path.is_file())
        return sorted(txt_files, key=self._txt_sort_key, reverse=True)

    def _txt_sort_key(self, txt_file: Path) -> tuple[datetime, float, str]:
        """Sort UroBiopsy TXT files by name date/time, then mtime."""

        match = UROBIOPSY_DATE_PATTERN.match(txt_file.name)
        parsed = datetime.min
        if match is not None:
            date_text = match.group("date")
            time_text = match.group("time") or "000000"
            try:
                parsed = datetime.strptime(f"{date_text}{time_text}", "%Y%m%d%H%M%S")
            except ValueError:
                parsed = datetime.min
        try:
            mtime = txt_file.stat().st_mtime
        except OSError:
            mtime = 0.0
        return parsed, mtime, str(txt_file).casefold()

    def _pooled_axis_action_samples(
        self,
        output_root: Path,
        records: list[DHRCaseProcessingRecord],
    ) -> list[dict[str, object]]:
        """Read processed case workbooks and return raw duration samples for pooling."""

        samples: list[dict[str, object]] = []
        for record in records:
            if record.processing_status != "Processed" or not record.main_workbook:
                continue
            main_path = output_root / record.main_workbook
            if not main_path.is_file():
                continue
            workbook = load_workbook(main_path, data_only=True, read_only=True)
            if "Distribution Raw Data" not in workbook.sheetnames:
                workbook.close()
                continue
            sheet = workbook["Distribution Raw Data"]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            header_index = {header: index for index, header in enumerate(headers) if header}
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not values:
                    continue
                axis = self._row_value(values, header_index, "Axis")
                rule_id = self._row_value(values, header_index, "Rule ID")
                action_label = self._row_value(values, header_index, "Action Label")
                duration_s = self._to_float_or_none(self._row_value(values, header_index, "Duration (s)"))
                action = self._human_action(rule_id, action_label)
                if not axis or not action or duration_s is None:
                    continue
                if self._should_exclude_action(action, rule_id, action_label):
                    continue
                samples.append(
                    {
                        "Case ID": record.case_id,
                        "Axis": str(axis),
                        "Action": action,
                        "Duration (s)": duration_s,
                    }
                )
            workbook.close()
        return samples

    def _summary_rows_from_samples(self, samples: list[dict[str, object]]) -> list[dict[str, object]]:
        """Return pooled cross-case statistics grouped by axis/action."""

        grouped: dict[tuple[str, str], list[float]] = {}
        for sample in samples:
            key = (str(sample["Axis"]), str(sample["Action"]))
            grouped.setdefault(key, []).append(float(sample["Duration (s)"]))
        return [
            self._summary_row(axis, action, durations)
            for (axis, action), durations in sorted(grouped.items(), key=lambda item: self._axis_action_sort_key(item[0]))
        ]

    def _case_contribution_rows_from_samples(self, samples: list[dict[str, object]]) -> list[dict[str, object]]:
        """Return per-case axis/action statistics from raw samples for audit."""

        grouped: dict[tuple[str, str, str], list[float]] = {}
        for sample in samples:
            key = (str(sample["Case ID"]), str(sample["Axis"]), str(sample["Action"]))
            grouped.setdefault(key, []).append(float(sample["Duration (s)"]))
        rows: list[dict[str, object]] = []
        for (case_id, axis, action), durations in sorted(
            grouped.items(),
            key=lambda item: (item[0][0], *self._axis_action_sort_key((item[0][1], item[0][2]))),
        ):
            rows.append({"Case ID": case_id, **self._summary_row(axis, action, durations)})
        return rows

    def _summary_row(self, axis: str, action: str, durations: list[float]) -> dict[str, object]:
        """Build one summary row from pooled raw duration values."""

        stats = self._duration_stats_seconds(durations)
        return {
            "Axis": axis,
            "Action": action,
            "n": stats["n"],
            "Mean (s)": stats["mean"],
            "SD (s)": stats["sd"],
            "Var (s²)": stats["var"],
            "Median (s)": stats["median"],
            "IQR (s)": stats["iqr"],
            "Min–Max (s)": self._min_max_display(stats["min"], stats["max"]),
            "CV (%)": stats["cv"],
        }

    def _write_rows(self, sheet, columns: list[str], rows: list[dict[str, object]]) -> None:
        """Write a rectangular table to a worksheet."""

        for column_index, column_name in enumerate(columns, start=1):
            sheet.cell(row=1, column=column_index, value=column_name)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column_name in enumerate(columns, start=1):
                sheet.cell(row=row_index, column=column_index, value=row.get(column_name))

    def _duration_stats_seconds(self, values: list[float]) -> dict[str, float | int | None]:
        """Compute summary statistics from second-based raw duration values."""

        numeric = [float(value) for value in values if self._to_float_or_none(value) is not None]
        if not numeric:
            return {
                "n": 0,
                "mean": None,
                "sd": None,
                "var": None,
                "median": None,
                "iqr": None,
                "min": None,
                "max": None,
                "cv": None,
            }
        count = len(numeric)
        mean = statistics.mean(numeric)
        sd = statistics.stdev(numeric) if count > 1 else None
        var = statistics.variance(numeric) if count > 1 else None
        q1 = self._linear_quantile(numeric, 0.25)
        q3 = self._linear_quantile(numeric, 0.75)
        return {
            "n": count,
            "mean": mean,
            "sd": sd,
            "var": var,
            "median": statistics.median(numeric),
            "iqr": q3 - q1,
            "min": min(numeric),
            "max": max(numeric),
            "cv": sd / mean * 100.0 if sd is not None and mean else None,
        }

    def _linear_quantile(self, values: list[float], quantile: float) -> float:
        """Return pandas-compatible linear interpolation quantile."""

        ordered = sorted(values)
        if len(ordered) == 1:
            return ordered[0]
        position = (len(ordered) - 1) * quantile
        lower_index = int(math.floor(position))
        upper_index = int(math.ceil(position))
        if lower_index == upper_index:
            return ordered[lower_index]
        fraction = position - lower_index
        return ordered[lower_index] + (ordered[upper_index] - ordered[lower_index]) * fraction

    def _format_sheet(self, sheet, mean_conditional_formatting: bool = False) -> None:
        """Apply simple report formatting to a cross-summary worksheet."""

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        alternate_fill = PatternFill("solid", fgColor="F3F4F6")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row_index in range(2, sheet.max_row + 1):
            if row_index % 2 == 0:
                for column_index in range(1, sheet.max_column + 1):
                    sheet.cell(row=row_index, column=column_index).fill = alternate_fill
        headers = {cell.value: cell.column for cell in sheet[1]}
        centered_headers = {
            "n",
            "Mean (s)",
            "SD (s)",
            "Var (s²)",
            "Median (s)",
            "IQR (s)",
            "Min–Max (s)",
            "CV (%)",
        }
        number_formats = {
            "Mean (s)": "0.000",
            "SD (s)": "0.0000",
            "Var (s²)": "0.0000",
            "Median (s)": "0.000",
            "IQR (s)": "0.0000",
            "CV (%)": "0.00",
        }
        for row_index in range(2, sheet.max_row + 1):
            for header in centered_headers:
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).alignment = Alignment(horizontal="center")
            for header, number_format in number_formats.items():
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).number_format = number_format
        if mean_conditional_formatting:
            self._apply_mean_conditional_formatting(sheet, headers)
        for column_cells in sheet.columns:
            header = column_cells[0].value
            values = [str(cell.value) for cell in column_cells if cell.value is not None]
            width = max((len(value) for value in values), default=10) + 2
            if header in {"n", "Mean (s)", "SD (s)", "Var (s²)", "Median (s)", "IQR (s)", "CV (%)"}:
                for cell in column_cells[1:]:
                    cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width, 10), 60)

    def _apply_mean_conditional_formatting(self, sheet, headers: dict[str, int]) -> None:
        """Apply red/yellow highlighting to Mean cells only."""

        if "Mean (s)" not in headers or sheet.max_row < 2:
            return
        mean_col = headers["Mean (s)"]
        mean_letter = sheet.cell(row=1, column=mean_col).column_letter
        target_range = f"{mean_letter}2:{mean_letter}{sheet.max_row}"
        red_formula = f"${mean_letter}2>={AXIS_SUMMARY_MEAN_RED_THRESHOLD}"
        yellow_formula = (
            f"AND(${mean_letter}2>={AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD},"
            f"${mean_letter}2<{AXIS_SUMMARY_MEAN_RED_THRESHOLD})"
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[red_formula],
                fill=PatternFill("solid", fgColor="FFC7CE"),
                stopIfTrue=True,
            ),
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[yellow_formula], fill=PatternFill("solid", fgColor="FFEB9C")),
        )

    def _relative_to_root(self, path: Path | str, root: Path) -> str:
        """Return a stable relative path string when possible."""

        resolved = Path(path).expanduser().resolve()
        try:
            return str(resolved.relative_to(root))
        except ValueError:
            return str(resolved)

    def _row_value(self, values: tuple[object, ...], header_index: dict[str, int], header: str) -> object:
        """Return a row value by header name."""

        index = header_index.get(header)
        if index is None or index >= len(values):
            return None
        return values[index]

    def _to_float_or_none(self, value: object) -> float | None:
        """Return a finite float or None."""

        if value is None:
            return None
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return None
        return numeric if math.isfinite(numeric) else None

    def _human_action(self, rule_id: object, fallback: object = "") -> str:
        """Return compact user-facing action text."""

        labels = {
            "search_reference": "Search Reference",
            "clear_motor": "Clear Motor",
            "move_to_home": "Move to Home",
            "move_to_max": "Move to Max",
        }
        key = str(rule_id or "").strip()
        return labels.get(key, str(fallback or key).strip())

    def _should_exclude_action(self, action: object, rule_id: object, action_label: object = "") -> bool:
        """Return whether a raw sample should be excluded from clean summaries."""

        candidates = {str(value or "").strip().casefold() for value in (action, rule_id, action_label)}
        return bool(
            candidates
            & {
                "search reference",
                "search_reference",
                "start searching reference -> reference found",
            }
        )

    def _axis_action_sort_key(self, key: tuple[str, str]) -> tuple[str, int, str]:
        """Sort summary rows by axis and standard action order."""

        axis, action = key
        action_order = {
            "Search Reference": 0,
            "Clear Motor": 1,
            "Move to Home": 2,
            "Move to Max": 3,
        }
        return (axis, action_order.get(action, 99), action)

    def _min_max_display(self, minimum: float | None, maximum: float | None) -> str:
        """Format a compact min-max seconds display."""

        if minimum is None or maximum is None:
            return ""
        return f"{minimum:.3f}\u2013{maximum:.3f}"


def safe_case_id(relative_case_folder: Path | str) -> str:
    """Return a safe output folder name from a DHR-relative case path."""

    parts = Path(relative_case_folder).parts
    cleaned = [re.sub(r"[^A-Za-z0-9]+", "_", part).strip("_") for part in parts]
    cleaned = [part for part in cleaned if part]
    return "_".join(cleaned) or "case"


def txt_date_from_name(txt_file: Path | str) -> str:
    """Return an ISO date/time label parsed from a UroBiopsy filename."""

    path = Path(txt_file)
    match = UROBIOPSY_DATE_PATTERN.match(path.name)
    if match is None:
        return ""
    date_text = match.group("date")
    time_text = match.group("time")
    try:
        if time_text:
            return datetime.strptime(f"{date_text}{time_text}", "%Y%m%d%H%M%S").isoformat(sep=" ")
        return datetime.strptime(date_text, "%Y%m%d").date().isoformat()
    except ValueError:
        return date_text
