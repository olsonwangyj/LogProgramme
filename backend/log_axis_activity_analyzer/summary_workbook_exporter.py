"""Standalone Excel export for the clean overall axis/action summary."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import (
    AXIS_SUMMARY_MEAN_RED_THRESHOLD,
    AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD,
    OVERALL_AXIS_ACTION_SUMMARY_COLUMNS,
    OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME,
)


DEFAULT_EXCLUDED_SUMMARY_ACTIONS = {"Move to Home", "move_to_home"}


class SummaryWorkbookExporter:
    """Write the standalone one-sheet overall axis/action summary workbook."""

    def __init__(self, logger: logging.Logger | None = None) -> None:
        """Initialize the exporter with an optional logger."""

        self._logger = logger or logging.getLogger(self.__class__.__name__)

    def export(
        self,
        output_path: Path | str,
        axis_action_summary,
        exclude_actions: set[str] | None = None,
    ) -> Path:
        """Write the summary-only workbook and return the resolved output path."""

        path = Path(output_path).expanduser().resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        excluded = exclude_actions if exclude_actions is not None else DEFAULT_EXCLUDED_SUMMARY_ACTIONS
        rows = [
            self._overall_axis_action_row(row)
            for row in self._summary_rows(axis_action_summary)
            if not self._should_exclude_row(row, excluded)
        ]

        self._logger.info("Exporting standalone summary workbook to %s", path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME
        self._write_table(sheet, rows)
        self._format_sheet(sheet)
        workbook.save(path)
        return path

    def _summary_rows(self, axis_action_summary) -> list[dict[str, object]]:
        """Return axis/action summary rows from a DataFrame or row iterable."""

        if axis_action_summary is None:
            return []
        if hasattr(axis_action_summary, "to_dict"):
            return list(axis_action_summary.to_dict(orient="records"))
        return [dict(row) for row in axis_action_summary]

    def _should_exclude_row(self, row: dict[str, object], exclude_actions: Iterable[str]) -> bool:
        """Return whether a row should be left out of the standalone summary report."""

        excluded = {self._normalize_action(value) for value in exclude_actions}
        action_candidates = (
            row.get("Action"),
            row.get("Rule ID"),
            row.get("Rule ID / Action"),
            row.get("Action Label"),
            row.get("Start Event"),
            row.get("Event Type / Start Event"),
        )
        if any(self._normalize_action(value) in excluded for value in action_candidates):
            return True
        action = self._normalize_action(row.get("Action"))
        if action == "start moving to home -> motor homed":
            return True
        start_event = self._normalize_action(row.get("Start Event") or row.get("Event Type / Start Event"))
        end_event = self._normalize_action(row.get("End Event"))
        return start_event.startswith("start moving to home") and end_event == "motor homed"

    def _overall_axis_action_row(self, row: dict[str, object]) -> dict[str, object]:
        """Map internal ASCII summary columns to final report-facing labels."""

        return {
            "Axis": row.get("Axis"),
            "Action": row.get("Action"),
            "n": row.get("n"),
            "Mean (s)": row.get("Mean (s)"),
            "SD (s)": row.get("SD (s)"),
            "Var (s²)": row.get("Var (s^2)", row.get("Var (s²)")),
            "Median (s)": row.get("Median (s)"),
            "IQR (s)": row.get("IQR (s)"),
            "Min–Max (s)": self._display_min_max(row.get("Min-Max (s)", row.get("Min–Max (s)"))),
            "CV (%)": row.get("CV (%)"),
        }

    def _write_table(self, sheet, rows: list[dict[str, object]]) -> None:
        """Write the one summary table."""

        for column_index, column_name in enumerate(OVERALL_AXIS_ACTION_SUMMARY_COLUMNS, start=1):
            sheet.cell(row=1, column=column_index, value=column_name)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column_name in enumerate(OVERALL_AXIS_ACTION_SUMMARY_COLUMNS, start=1):
                sheet.cell(row=row_index, column=column_index, value=row.get(column_name))

    def _format_sheet(self, sheet) -> None:
        """Apply final-report formatting and Mean-only conditional formatting."""

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        alternate_fill = PatternFill("solid", fgColor="F3F4F6")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        headers = {cell.value: cell.column for cell in sheet[1]}
        centered_headers = {"n", "Mean (s)", "SD (s)", "Var (s²)", "Median (s)", "IQR (s)", "Min–Max (s)", "CV (%)"}
        number_formats = {
            "Mean (s)": "0.000",
            "SD (s)": "0.0000",
            "Var (s²)": "0.0000",
            "Median (s)": "0.000",
            "IQR (s)": "0.0000",
            "CV (%)": "0.00",
        }
        for row_index in range(2, sheet.max_row + 1):
            if row_index % 2 == 0:
                for column in range(1, sheet.max_column + 1):
                    sheet.cell(row=row_index, column=column).fill = alternate_fill
            for header in centered_headers:
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).alignment = Alignment(horizontal="center")
            for header, number_format in number_formats.items():
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).number_format = number_format
        self._apply_conditional_formatting(sheet, headers)
        self._auto_fit_columns(sheet)

    def _apply_conditional_formatting(self, sheet, headers: dict[str, int]) -> None:
        """Highlight only Mean cells with the configured thresholds."""

        if "Mean (s)" not in headers or sheet.max_row < 2:
            return
        mean_col = headers["Mean (s)"]
        mean_letter = sheet.cell(row=1, column=mean_col).column_letter
        target_range = f"{mean_letter}2:{mean_letter}{sheet.max_row}"
        red_formula = f"${mean_letter}2>={AXIS_SUMMARY_MEAN_RED_THRESHOLD}"
        yellow_formula = (
            f"AND(${mean_letter}2>={AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD},"
            f"${mean_letter}2<{AXIS_SUMMARY_MEAN_RED_THRESHOLD})"
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[red_formula],
                fill=PatternFill("solid", fgColor="FFC7CE"),
                stopIfTrue=True,
            ),
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[yellow_formula], fill=PatternFill("solid", fgColor="FFEB9C")),
        )

    def _auto_fit_columns(self, sheet) -> None:
        """Set readable worksheet column widths."""

        minimum_widths = {
            "Axis": 10,
            "Action": 18,
            "n": 8,
            "Mean (s)": 12,
            "SD (s)": 12,
            "Var (s²)": 12,
            "Median (s)": 12,
            "IQR (s)": 12,
            "Min–Max (s)": 16,
            "CV (%)": 10,
        }
        for column_cells in sheet.columns:
            header = column_cells[0].value
            values = [str(cell.value) for cell in column_cells if cell.value is not None]
            width = max(max((len(value) for value in values), default=10) + 2, minimum_widths.get(header, 10))
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 40)

    def _display_min_max(self, value: object) -> object:
        """Return min-max display text with the report-facing en dash."""

        if value is None:
            return None
        return str(value).replace("-", "–")

    def _normalize_action(self, value: object) -> str:
        """Normalize row labels for robust action filtering."""

        return str(value or "").strip().casefold()
