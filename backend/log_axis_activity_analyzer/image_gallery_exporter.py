"""Separate workbook export for distribution chart image galleries."""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import (
    AXIS_ACTION_SUMMARY_COLUMNS,
    AXIS_SUMMARY_MEAN_RED_THRESHOLD,
    AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD,
    DISTRIBUTION_IMAGE_BLOCK_HEIGHT_ROWS,
    DISTRIBUTION_IMAGE_GALLERY_LAYOUT,
    DISTRIBUTION_IMAGE_GALLERY_LAYOUT_OPTIONS,
    DISTRIBUTION_IMAGE_GALLERY_METADATA_KEYS,
    DISTRIBUTION_IMAGE_GALLERY_INCLUDE_SKIPPED_GROUPS,
    DISTRIBUTION_IMAGE_GALLERY_SHEET_NAME,
    DISTRIBUTION_IMAGE_GRID_BLOCK_WIDTH_COLUMNS,
    DISTRIBUTION_IMAGE_GRID_COLUMNS,
    DISTRIBUTION_IMAGE_HEIGHT_PX,
    DISTRIBUTION_IMAGE_INDEX_COLUMNS,
    DISTRIBUTION_IMAGE_INDEX_SHEET_NAME,
    DISTRIBUTION_IMAGE_MAX_IMAGES,
    DISTRIBUTION_IMAGE_STATISTICS_COLUMNS,
    DISTRIBUTION_IMAGE_STATISTICS_SHEET_NAME,
    DISTRIBUTION_IMAGE_WIDTH_PX,
    OVERALL_AXIS_ACTION_SUMMARY_COLUMNS,
    OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME,
)
from .distribution import (
    DistributionAnalysisResult,
    DistributionInputRow,
    DistributionStats,
    ReferenceDurationInputRow,
    ReferenceDurationStats,
)

StatsLike = DistributionStats | ReferenceDurationStats
RowLike = DistributionInputRow | ReferenceDurationInputRow


@dataclass(frozen=True)
class DistributionImageGalleryExportResult:
    """Summary of one image gallery workbook export."""

    output_path: Path
    groups_with_chart_file_path_count: int
    existing_chart_file_count: int
    image_inserted_count: int
    statistics_count: int
    groups_without_charts_count: int
    image_limit_skipped_count: int
    missing_chart_file_count: int
    image_embedding_unavailable_count: int
    image_insert_failed_count: int

    @property
    def chart_file_count(self) -> int:
        """Backward-compatible alias for existing chart files."""

        return self.existing_chart_file_count


class DistributionImageGalleryExporter:
    """Write generated distribution PNGs and key statistics into a separate workbook."""

    def __init__(
        self,
        max_images: int = DISTRIBUTION_IMAGE_MAX_IMAGES,
        image_width_px: int = DISTRIBUTION_IMAGE_WIDTH_PX,
        image_height_px: int = DISTRIBUTION_IMAGE_HEIGHT_PX,
        block_height_rows: int = DISTRIBUTION_IMAGE_BLOCK_HEIGHT_ROWS,
        include_skipped_groups: bool = DISTRIBUTION_IMAGE_GALLERY_INCLUDE_SKIPPED_GROUPS,
        layout: str = DISTRIBUTION_IMAGE_GALLERY_LAYOUT,
        grid_columns: int = DISTRIBUTION_IMAGE_GRID_COLUMNS,
        grid_block_width_columns: int = DISTRIBUTION_IMAGE_GRID_BLOCK_WIDTH_COLUMNS,
        logger: logging.Logger | None = None,
    ) -> None:
        """Initialize gallery workbook layout options."""

        self.max_images = max_images
        self.image_width_px = image_width_px
        self.image_height_px = image_height_px
        image_rows = math.ceil(image_height_px / 18) + 6
        self.block_height_rows = max(block_height_rows, image_rows)
        self.include_skipped_groups = include_skipped_groups
        self.layout = layout if layout in DISTRIBUTION_IMAGE_GALLERY_LAYOUT_OPTIONS else DISTRIBUTION_IMAGE_GALLERY_LAYOUT
        self.grid_columns = max(1, int(grid_columns))
        self.grid_block_width_columns = max(4, int(grid_block_width_columns))
        self._logger = logger or logging.getLogger(self.__class__.__name__)

    def export(
        self,
        output_path: Path | str,
        distribution_result: DistributionAnalysisResult,
        axis_action_summary=None,
    ) -> DistributionImageGalleryExportResult:
        """Create the image gallery workbook and return export counts."""

        path = Path(output_path).expanduser().resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        self._logger.info("Exporting distribution image gallery workbook to %s", path)

        workbook = Workbook()
        overall_sheet = workbook.active
        overall_sheet.title = OVERALL_AXIS_ACTION_SUMMARY_SHEET_NAME
        gallery_sheet = workbook.create_sheet(DISTRIBUTION_IMAGE_GALLERY_SHEET_NAME)
        statistics_sheet = workbook.create_sheet(DISTRIBUTION_IMAGE_STATISTICS_SHEET_NAME)
        index_sheet = workbook.create_sheet(DISTRIBUTION_IMAGE_INDEX_SHEET_NAME)
        axis_action_sheet = workbook.create_sheet("Axis Action Summary")

        rows_by_group = distribution_result.grouped_rows
        gallery_rows, index_rows, counts = self._build_gallery_and_index(
            gallery_sheet,
            distribution_result.stats,
            rows_by_group,
        )
        statistics_rows = [
            self._statistics_row(stats, rows_by_group.get(stats.group_id, []), gallery_rows[stats.group_id])
            for stats in distribution_result.stats
        ]
        self._write_table(statistics_sheet, DISTRIBUTION_IMAGE_STATISTICS_COLUMNS, statistics_rows)
        self._write_table(index_sheet, DISTRIBUTION_IMAGE_INDEX_COLUMNS, index_rows)
        self._write_overall_axis_action_summary(overall_sheet, axis_action_summary)
        self._write_axis_action_summary(axis_action_sheet, axis_action_summary)
        self._format_gallery_sheet(gallery_sheet)
        self._format_table_sheet(statistics_sheet)
        self._format_table_sheet(index_sheet)
        self._format_axis_action_summary_sheet(overall_sheet)
        self._format_axis_action_summary_sheet(axis_action_sheet)

        workbook.save(path)
        self._logger.info(
            "Exported distribution image gallery with %s inserted images and %s statistics rows",
            counts["image_inserted_count"],
            len(statistics_rows),
        )
        return DistributionImageGalleryExportResult(
            output_path=path,
            groups_with_chart_file_path_count=counts["groups_with_chart_file_path_count"],
            existing_chart_file_count=counts["existing_chart_file_count"],
            image_inserted_count=counts["image_inserted_count"],
            statistics_count=len(statistics_rows),
            groups_without_charts_count=counts["groups_without_charts_count"],
            image_limit_skipped_count=counts["image_limit_skipped_count"],
            missing_chart_file_count=counts["missing_chart_file_count"],
            image_embedding_unavailable_count=counts["image_embedding_unavailable_count"],
            image_insert_failed_count=counts["image_insert_failed_count"],
        )

    def _build_gallery_and_index(
        self,
        sheet,
        stats_list: list[StatsLike],
        rows_by_group: dict[str, list[RowLike]],
    ) -> tuple[dict[str, dict[str, object]], list[dict[str, object]], dict[str, int]]:
        """Write gallery blocks and build index rows."""

        gallery_rows: dict[str, dict[str, object]] = {}
        index_rows: list[dict[str, object]] = []
        current_row = 1
        image_class = self._openpyxl_image_class()
        image_class_available = image_class is not None
        counts = {
            "groups_with_chart_file_path_count": 0,
            "existing_chart_file_count": 0,
            "image_inserted_count": 0,
            "groups_without_charts_count": 0,
            "image_limit_skipped_count": 0,
            "missing_chart_file_count": 0,
            "image_embedding_unavailable_count": 0,
            "image_insert_failed_count": 0,
        }
        wrote_gallery_block = False
        compact_item_count = 0
        for stats in stats_list:
            first_row = self._first_row(rows_by_group, stats.group_id)
            chart_file = str(getattr(stats, "chart_file", "") or "")
            chart_path = Path(chart_file) if chart_file else None
            image_status = self._gallery_image_status(
                stats,
                chart_path,
                counts["image_inserted_count"],
                image_class_available,
            )
            image_insert_error = ""
            image_number = None
            anchor = f"A{current_row + 3}"
            compact_position: tuple[int, int, str] | None = None
            if image_status == "ImageInserted":
                if self.layout == "compact_grid":
                    compact_position = self._compact_grid_position(compact_item_count)
                    anchor = compact_position[2]
                try:
                    image = image_class(str(chart_path))
                    image.width = self.image_width_px
                    image.height = self.image_height_px
                    sheet.add_image(image, anchor)
                    image_number = counts["image_inserted_count"] + 1
                except Exception as exc:
                    image_status = "ImageInsertFailed"
                    image_insert_error = str(exc)
                    image_number = None
            self._increment_gallery_counts(counts, image_status)
            gallery_row = self._gallery_metadata_row(stats, first_row, image_status, image_insert_error, image_number)
            gallery_rows[stats.group_id] = gallery_row
            if self._should_write_gallery_block(image_status):
                wrote_gallery_block = True
                if self.layout == "compact_grid":
                    start_row, start_column, anchor = compact_position or self._compact_grid_position(compact_item_count)
                    compact_item_count += 1
                    self._write_compact_grid_block(sheet, start_row, start_column, gallery_row)
                else:
                    self._write_gallery_block(sheet, current_row, gallery_row)
                if image_status != "ImageInserted" and self.layout != "compact_grid":
                    sheet.cell(
                        row=current_row + 3,
                        column=1,
                        value=image_status,
                    )
                if self.layout != "compact_grid":
                    current_row += self.block_height_rows
            if self._should_write_index_row(stats, image_status):
                group_rows = rows_by_group.get(stats.group_id, [])
                index_rows.append(
                    self._index_row(
                        stats,
                        group_rows,
                        image_number,
                        self._chart_file_display(chart_file),
                        anchor if image_status == "ImageInserted" else "",
                        image_status,
                        image_insert_error,
                    )
                )
        if not wrote_gallery_block:
            self._write_no_image_summary(sheet, stats_list, counts, len(stats_list))
        return gallery_rows, index_rows, counts

    def _write_no_image_summary(
        self,
        sheet,
        stats_list: list[StatsLike],
        counts: dict[str, int],
        statistics_count: int,
    ) -> None:
        """Write a reason-specific summary when no images are inserted into the gallery sheet."""

        summary_rows = [
            ("Image Gallery Summary", ""),
            ("Statistics rows", statistics_count),
            ("Groups with chart file path", counts["groups_with_chart_file_path_count"]),
            ("Existing chart files", counts["existing_chart_file_count"]),
            ("Images inserted", counts["image_inserted_count"]),
            ("Groups without charts", counts["groups_without_charts_count"]),
            ("Images skipped due to image limit", counts["image_limit_skipped_count"]),
            ("Chart files missing", counts["missing_chart_file_count"]),
            ("Image embedding unavailable", counts["image_embedding_unavailable_count"]),
            ("Image insert failures", counts["image_insert_failed_count"]),
            ("Reason", self._no_image_reason(stats_list, counts)),
        ]
        for row_index, (key, value) in enumerate(summary_rows, start=1):
            sheet.cell(row=row_index, column=1, value=key)
            sheet.cell(row=row_index, column=2, value=value)
            sheet.cell(row=row_index, column=1).font = Font(bold=True)

    def _no_image_reason(self, stats_list: list[StatsLike], counts: dict[str, int]) -> str:
        """Return a clear reason for an image-less gallery sheet."""

        if counts["existing_chart_file_count"] and counts["image_limit_skipped_count"] == counts["existing_chart_file_count"]:
            return f"Chart files were generated, but no images were inserted because the image limit is {self.max_images}."
        if (
            counts["existing_chart_file_count"]
            and counts["image_embedding_unavailable_count"] == counts["existing_chart_file_count"]
        ):
            return "Chart files were generated, but images could not be embedded because openpyxl/Pillow image support is unavailable."
        if counts["groups_with_chart_file_path_count"] and counts["missing_chart_file_count"] == counts["groups_with_chart_file_path_count"]:
            return "Chart metadata exists, but chart image files could not be found."
        if counts["existing_chart_file_count"] and counts["image_insert_failed_count"] == counts["existing_chart_file_count"]:
            return "Chart files were generated, but image insertion failed for every available chart file."
        statuses = {self._display_chart_status(stats) for stats in stats_list}
        if statuses and statuses <= {"SkippedInsufficientSamples", "InsufficientSamples", "InsufficientSamplesForNormalFit"}:
            return "No chart images were generated because all groups had insufficient samples. See Image Statistics."
        if counts["groups_without_charts_count"] == len(stats_list):
            return "No chart images were generated because no distribution group produced a chart file. See Image Statistics."
        return "No chart images were inserted. See Image Statistics for Distribution Status, Chart Status, and Gallery Image Status."

    def _compact_grid_position(self, item_index: int) -> tuple[int, int, str]:
        """Return title row, start column, and image anchor for one compact-grid item."""

        grid_row = item_index // self.grid_columns
        grid_column = item_index % self.grid_columns
        start_row = grid_row * self.block_height_rows + 1
        start_column = grid_column * self.grid_block_width_columns + 1
        image_anchor = f"{get_column_letter(start_column)}{start_row + 2}"
        return start_row, start_column, image_anchor

    def _write_compact_grid_block(self, sheet, start_row: int, start_column: int, metadata: dict[str, object]) -> None:
        """Write a short title block for compact image gallery layout."""

        if metadata.get("Chart Type") == "Reference Duration Distribution":
            title = (
                f"Reference Duration Distribution | Axis {metadata.get('Axis', '')} | "
                f"{metadata.get('Action', '')} | Distance Not Applicable"
            )
        else:
            title = (
                f"Motion Duration Distribution | Axis {metadata.get('Axis', '')} | "
                f"PWM {metadata.get('PWM (%)', '')}% | "
                f"Hardware Actual Distance {metadata.get('Hardware Actual Distance Group Display', '')}"
            )
        subtitle = (
            f"N={metadata.get('Sample Count', '')}, Mean={metadata.get('Mean Duration (s)', '')}, "
            f"SD={metadata.get('Sample SD Duration (s)', '')}, Status={metadata.get('Gallery Image Status', '')}"
        )
        title_cell = sheet.cell(row=start_row, column=start_column, value=title)
        title_cell.font = Font(bold=True)
        sheet.cell(row=start_row + 1, column=start_column, value=subtitle)

    def _write_gallery_block(self, sheet, start_row: int, metadata: dict[str, object]) -> None:
        """Write one compact chart summary row above the image."""

        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        value_fill = PatternFill("solid", fgColor="F9FAFB")
        for column_index, key in enumerate(DISTRIBUTION_IMAGE_GALLERY_METADATA_KEYS, start=1):
            header_cell = sheet.cell(row=start_row, column=column_index, value=key)
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = Alignment(horizontal="center")
            value_cell = sheet.cell(row=start_row + 1, column=column_index, value=metadata.get(key))
            value_cell.fill = value_fill
            value_cell.alignment = Alignment(horizontal="center", vertical="center")

    def _statistics_row(
        self,
        stats: StatsLike,
        rows: list[RowLike],
        gallery_row: dict[str, object],
    ) -> dict[str, object]:
        """Build one Image Statistics row from computed stats, not chart text."""

        first_row = rows[0] if rows else None
        population_std_s = self._population_std_s(stats)
        population_var_s2 = self._population_var_s2(stats)
        return {
            "Group ID": stats.group_id,
            "Chart Type": self._chart_type(stats),
            "TXT Source File": self._basename(getattr(stats, "txt_source_file", "")),
            "PWM (%)": self._abs_optional(getattr(stats, "pwm_percent", None)),
            "PWM Raw Values Seen": self._numeric_values_seen(getattr(row, "pwm_raw_value", None) for row in rows),
            "PWM Directions Seen": self._string_values_seen(getattr(row, "pwm_direction", "") for row in rows),
            "PWM Direction Mixed": self._values_mixed(getattr(row, "pwm_direction", "") for row in rows),
            "PWM Raw Value Example": getattr(first_row, "pwm_raw_value", None) if first_row is not None else None,
            "PWM Direction Example": getattr(first_row, "pwm_direction", "") if first_row is not None else "",
            "Axis": stats.axis,
            "Action": self._action(stats),
            "Selected Group Distance": self._abs_optional(getattr(stats, "movement_distance", None)),
            "Hardware Actual Distance Group Value": self._abs_optional(
                getattr(stats, "movement_distance_group_value", None)
            ),
            "Hardware Actual Distance Group Display": self._hardware_distance_group_display(stats),
            "Distance": self._distance_display(stats),
            "Reference Evidence Status": self._reference_evidence_status(stats),
            "Movement Distance Grouping Mode": getattr(stats, "movement_distance_grouping_mode", ""),
            "Movement Distance Bin Size": getattr(stats, "movement_distance_bin_size", None),
            "Hardware Actual Distance Raw Example": self._abs_optional(
                getattr(stats, "movement_distance_raw_example", None)
            ),
            "Hardware Actual Distance Min": self._abs_optional(getattr(stats, "movement_distance_min", None)),
            "Hardware Actual Distance Max": self._abs_optional(getattr(stats, "movement_distance_max", None)),
            "Hardware Motion Source File": self._basename(
                getattr(first_row, "hardware_motion_source_file", "") if first_row is not None else ""
            ),
            "Hardware Raw Start Position": getattr(first_row, "hardware_raw_start_position", None)
            if first_row is not None
            else None,
            "Hardware Raw End Position": getattr(first_row, "hardware_raw_end_position", None)
            if first_row is not None
            else None,
            "Hardware Start Position": getattr(first_row, "hardware_start_position", None)
            if first_row is not None
            else None,
            "Hardware End Position": getattr(first_row, "hardware_end_position", None)
            if first_row is not None
            else None,
            "Hardware Actual Distance": getattr(first_row, "hardware_actual_distance", None)
            if first_row is not None
            else None,
            "Position Values Mixed": getattr(stats, "position_values_mixed", None),
            "Hardware Distance Source": getattr(stats, "movement_distance_source", ""),
            "Hardware Distance Method": getattr(stats, "movement_distance_method", ""),
            "Rule ID": stats.rule_id,
            "Action Label": stats.action_label,
            "Sample Count": stats.sample_count,
            "Mean Duration (ms)": stats.mean_ms,
            "Mean Duration (s)": stats.mean_s,
            "Median Duration (ms)": stats.median_ms,
            "Median Duration (s)": stats.median_s,
            "Sample SD Duration (ms)": stats.sample_std_ms,
            "Sample SD Duration (s)": stats.sample_std_s,
            "Sample Variance Duration (ms^2)": stats.sample_var_ms2,
            "Sample Variance Duration (s^2)": stats.sample_var_s2,
            "Normal Fit Mean (s)": stats.normal_fit_mean_s,
            "Normal Fit Std Dev (s)": stats.normal_fit_std_s,
            "Normal Fit Variance (s^2)": stats.normal_fit_variance_s2,
            "Population SD Duration (ms)": getattr(stats, "population_std_ms", None),
            "Population SD Duration (s)": population_std_s,
            "Population Variance Duration (ms^2)": getattr(stats, "population_var_ms2", None),
            "Population Variance Duration (s^2)": population_var_s2,
            "Min Duration (ms)": stats.min_ms,
            "Max Duration (ms)": stats.max_ms,
            "P05 Duration (ms)": getattr(stats, "p05_ms", None),
            "P25 Duration (ms)": getattr(stats, "p25_ms", None),
            "P75 Duration (ms)": getattr(stats, "p75_ms", None),
            "P95 Duration (ms)": getattr(stats, "p95_ms", None),
            "Outlier Count": getattr(stats, "outlier_count", 0),
            "Outlier Values": getattr(stats, "outlier_values", ""),
            "Chart Uses Outlier-Trimmed Axis": getattr(stats, "chart_uses_outlier_trimmed_axis", False),
            "Coefficient of Variation (%)": stats.cv_percent,
            "Distribution Status": stats.distribution_status,
            "Chart Status": gallery_row["Chart Status"],
            "Gallery Image Status": gallery_row["Gallery Image Status"],
            "Image Insert Error": gallery_row.get("Image Insert Error", ""),
            "Chart File": self._chart_file_display(getattr(stats, "chart_file", None)),
            "Notes": stats.notes,
        }

    def _gallery_metadata_row(
        self,
        stats: StatsLike,
        first_row: RowLike | None,
        image_status: str,
        image_insert_error: str = "",
        image_number: int | None = None,
    ) -> dict[str, object]:
        """Build one Image Gallery metadata block row."""

        return {
            "Image #": image_number,
            "Group ID": stats.group_id,
            "Chart Type": self._chart_type(stats),
            "TXT Source File": self._basename(getattr(stats, "txt_source_file", "")),
            "PWM (%)": self._abs_optional(getattr(stats, "pwm_percent", None)),
            "Axis": stats.axis,
            "Action": self._action(stats),
            "n": stats.sample_count,
            "Mean (s)": stats.mean_s,
            "SD (s)": stats.sample_std_s,
            "Var (s^2)": stats.sample_var_s2,
            "Median (s)": stats.median_s,
            "Min-Max (s)": self._min_max_display(stats),
            "CV (%)": stats.cv_percent,
            "Hardware Actual Distance": self._abs_optional(getattr(stats, "movement_distance", None)),
            "Hardware Actual Distance Group Value": self._abs_optional(
                getattr(stats, "movement_distance_group_value", None)
            ),
            "Hardware Actual Distance Group Display": self._hardware_distance_group_display(stats),
            "Distance": self._distance_display(stats),
            "Reference Evidence Status": self._reference_evidence_status(stats),
            "Hardware Distance Source": getattr(stats, "movement_distance_source", ""),
            "Hardware Distance Method": getattr(stats, "movement_distance_method", ""),
            "Movement Distance Grouping Mode": getattr(stats, "movement_distance_grouping_mode", ""),
            "Movement Distance Bin Size": getattr(stats, "movement_distance_bin_size", None),
            "Rule ID": stats.rule_id,
            "Action Label": stats.action_label,
            "Sample Count": stats.sample_count,
            "Mean Duration (s)": stats.mean_s,
            "Sample SD Duration (s)": stats.sample_std_s,
            "Sample Variance Duration (s^2)": stats.sample_var_s2,
            "Outlier Count": getattr(stats, "outlier_count", 0),
            "Outlier Values": getattr(stats, "outlier_values", ""),
            "Chart Uses Outlier-Trimmed Axis": getattr(stats, "chart_uses_outlier_trimmed_axis", False),
            "Chart File": self._chart_file_display(getattr(stats, "chart_file", None)),
            "Distribution Status": stats.distribution_status,
            "Chart Status": self._display_chart_status(stats),
            "Gallery Image Status": image_status,
            "Image Insert Error": image_insert_error,
            "PWM Raw Value Example": getattr(first_row, "pwm_raw_value", None) if first_row is not None else None,
        }

    def _index_row(
        self,
        stats: StatsLike,
        rows: list[RowLike],
        image_number: int | None,
        chart_file: str,
        anchor: str,
        image_status: str,
        image_insert_error: str = "",
    ) -> dict[str, object]:
        """Build one Image Index row."""

        return {
            "Image Number": image_number,
            "Group ID": stats.group_id,
            "Chart File": chart_file,
            "Excel Anchor": f"{DISTRIBUTION_IMAGE_GALLERY_SHEET_NAME}!{anchor}" if anchor else "",
            "Chart Type": self._chart_type(stats),
            "Axis": stats.axis,
            "Action": self._action(stats),
            "PWM (%)": self._abs_optional(getattr(stats, "pwm_percent", None)),
            "Hardware Actual Distance Group Value": self._abs_optional(
                getattr(stats, "movement_distance_group_value", None)
            ),
            "Hardware Actual Distance Group Display": self._hardware_distance_group_display(stats),
            "Distance": self._distance_display(stats),
            "Reference Evidence Status": self._reference_evidence_status(stats),
            "Hardware Distance Source": getattr(stats, "movement_distance_source", ""),
            "Hardware Distance Method": getattr(stats, "movement_distance_method", ""),
            "Movement Distance Grouping Mode": getattr(stats, "movement_distance_grouping_mode", ""),
            "Movement Distance Bin Size": getattr(stats, "movement_distance_bin_size", None),
            "PWM Raw Values Seen": self._numeric_values_seen(getattr(row, "pwm_raw_value", None) for row in rows),
            "PWM Directions Seen": self._string_values_seen(getattr(row, "pwm_direction", "") for row in rows),
            "PWM Direction Mixed": self._values_mixed(getattr(row, "pwm_direction", "") for row in rows),
            "Rule ID": stats.rule_id,
            "Sample Count": stats.sample_count,
            "Mean Duration (s)": stats.mean_s,
            "Sample SD Duration (s)": stats.sample_std_s,
            "Normal Fit Mean (s)": stats.normal_fit_mean_s,
            "Normal Fit Std Dev (s)": stats.normal_fit_std_s,
            "Normal Fit Variance (s^2)": stats.normal_fit_variance_s2,
            "Outlier Count": getattr(stats, "outlier_count", 0),
            "Outlier Values": getattr(stats, "outlier_values", ""),
            "Chart Uses Outlier-Trimmed Axis": getattr(stats, "chart_uses_outlier_trimmed_axis", False),
            "Chart Status": self._display_chart_status(stats),
            "Gallery Image Status": image_status,
            "Image Insert Error": image_insert_error,
        }

    def _gallery_image_status(
        self,
        stats: StatsLike,
        chart_path: Path | None,
        inserted_count: int,
        image_class_available: bool,
    ) -> str:
        """Return image insertion status for one group."""

        if not getattr(stats, "chart_file", None):
            return self._display_chart_status(stats)
        if chart_path is None or not chart_path.is_file():
            return "Chart File Missing"
        if inserted_count >= self.max_images:
            return "SkippedDueToImageLimit"
        if not image_class_available:
            return "ImageEmbeddingUnavailable"
        return "ImageInserted"

    def _increment_gallery_counts(self, counts: dict[str, int], image_status: str) -> None:
        """Update gallery export counters from one image status."""

        if image_status == "Chart File Missing":
            counts["groups_with_chart_file_path_count"] += 1
            counts["missing_chart_file_count"] += 1
            return
        if image_status in {
            "ImageInserted",
            "SkippedDueToImageLimit",
            "ImageEmbeddingUnavailable",
            "ImageInsertFailed",
        }:
            counts["groups_with_chart_file_path_count"] += 1
            counts["existing_chart_file_count"] += 1
        else:
            counts["groups_without_charts_count"] += 1
            return
        if image_status == "ImageInserted":
            counts["image_inserted_count"] += 1
        if image_status == "SkippedDueToImageLimit":
            counts["image_limit_skipped_count"] += 1
        if image_status == "ImageEmbeddingUnavailable":
            counts["image_embedding_unavailable_count"] += 1
        if image_status == "ImageInsertFailed":
            counts["image_insert_failed_count"] += 1

    def _should_write_gallery_block(self, image_status: str) -> bool:
        """Return whether a full block should appear on Image Gallery."""

        return image_status == "ImageInserted" or self.include_skipped_groups

    def _should_write_index_row(self, stats: StatsLike, image_status: str) -> bool:
        """Return whether a group should appear in Image Index."""

        return bool(getattr(stats, "chart_file", None)) or self.include_skipped_groups or image_status == "ImageInserted"

    def _display_chart_status(self, stats: StatsLike) -> str:
        """Normalize chart status for the standalone gallery workbook."""

        if stats.chart_status == "SkippedInsufficientSamplesForChart":
            return "SkippedInsufficientSamples"
        return stats.chart_status

    def _write_table(self, sheet, columns: list[str], rows: list[dict[str, object]]) -> None:
        """Write a simple headered table."""

        for column_index, column_name in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=column_index, value=column_name)
            cell.font = Font(bold=True)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column_name in enumerate(columns, start=1):
                sheet.cell(row=row_index, column=column_index, value=row.get(column_name))

    def _write_axis_action_summary(self, sheet, axis_action_summary) -> None:
        """Write the clean axis/action summary into the chart/statistics workbook."""

        for column_index, column_name in enumerate(AXIS_ACTION_SUMMARY_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=column_index, value=column_name)
            cell.font = Font(bold=True)
        rows = []
        if axis_action_summary is not None:
            if hasattr(axis_action_summary, "to_dict"):
                rows = axis_action_summary.to_dict(orient="records")
            else:
                rows = list(axis_action_summary)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column_name in enumerate(AXIS_ACTION_SUMMARY_COLUMNS, start=1):
                sheet.cell(row=row_index, column=column_index, value=row.get(column_name))

    def _write_overall_axis_action_summary(self, sheet, axis_action_summary) -> None:
        """Write the consolidated front summary sheet with final report labels."""

        for column_index, column_name in enumerate(OVERALL_AXIS_ACTION_SUMMARY_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=column_index, value=column_name)
            cell.font = Font(bold=True)
        rows = []
        if axis_action_summary is not None:
            if hasattr(axis_action_summary, "to_dict"):
                rows = axis_action_summary.to_dict(orient="records")
            else:
                rows = list(axis_action_summary)
        if not rows:
            sheet.cell(row=2, column=1, value="No axis/action summary data was available.")
            return
        for row_index, row in enumerate(rows, start=2):
            normalized = self._overall_axis_action_row(row)
            for column_index, column_name in enumerate(OVERALL_AXIS_ACTION_SUMMARY_COLUMNS, start=1):
                sheet.cell(row=row_index, column=column_index, value=normalized.get(column_name))

    def _overall_axis_action_row(self, row: dict[str, object]) -> dict[str, object]:
        """Map internal ASCII summary columns to user-facing front-summary labels."""

        return {
            "Axis": row.get("Axis"),
            "Action": row.get("Action"),
            "n": row.get("n"),
            "Mean (s)": row.get("Mean (s)"),
            "SD (s)": row.get("SD (s)"),
            "Var (s²)": row.get("Var (s^2)", row.get("Var (s²)")),
            "Median (s)": row.get("Median (s)"),
            "Min–Max (s)": self._display_min_max(row.get("Min-Max (s)", row.get("Min–Max (s)"))),
            "CV (%)": row.get("CV (%)"),
        }

    def _format_gallery_sheet(self, sheet) -> None:
        """Apply readable gallery worksheet sizing."""

        sheet.freeze_panes = "A2"
        widths = {
            "A": 10,
            "B": 30,
            "C": 10,
            "D": 18,
            "E": 8,
            "F": 11,
            "G": 11,
            "H": 12,
            "I": 12,
            "J": 16,
            "K": 10,
            "L": 10,
            "M": 14,
            "N": 22,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        if self.layout == "compact_grid":
            for column_index in range(1, self.grid_columns * self.grid_block_width_columns + 1):
                sheet.column_dimensions[get_column_letter(column_index)].width = 14

    def _format_table_sheet(self, sheet) -> None:
        """Apply basic formatting to an index/statistics worksheet."""

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [str(cell.value) for cell in column if cell.value is not None]
            width = min(max((len(value) for value in values), default=10) + 2, 72)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    def _format_axis_action_summary_sheet(self, sheet) -> None:
        """Apply summary table styling and Mean-only conditional formatting."""

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F2937")
        header_font = Font(color="FFFFFF", bold=True)
        alternate_fill = PatternFill("solid", fgColor="F3F4F6")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row_index in range(2, sheet.max_row + 1):
            if row_index % 2 == 0:
                for column in range(1, sheet.max_column + 1):
                    sheet.cell(row=row_index, column=column).fill = alternate_fill
            for header in ("n", "Mean (s)", "SD (s)", "Var (s^2)", "Var (s²)", "Median (s)", "CV (%)"):
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).alignment = Alignment(horizontal="center")
            for header, number_format in {
                "Mean (s)": "0.000",
                "SD (s)": "0.0000",
                "Var (s^2)": "0.0000",
                "Var (s²)": "0.0000",
                "Median (s)": "0.000",
                "CV (%)": "0.00",
            }.items():
                if header in headers:
                    sheet.cell(row=row_index, column=headers[header]).number_format = number_format
        self._apply_axis_action_conditional_formatting(sheet, headers)
        self._format_table_sheet(sheet)

    def _apply_axis_action_conditional_formatting(self, sheet, headers: dict[str, int]) -> None:
        """Highlight only Mean cells based on the latest report thresholds."""

        if "Mean (s)" not in headers:
            return
        if sheet.max_row < 2:
            return
        mean_col = headers["Mean (s)"]
        mean_letter = sheet.cell(row=1, column=mean_col).column_letter
        target_range = f"{mean_letter}2:{mean_letter}{sheet.max_row}"
        red_formula = f"${mean_letter}2>={AXIS_SUMMARY_MEAN_RED_THRESHOLD}"
        yellow_formula = (
            f"AND(${mean_letter}2>={AXIS_SUMMARY_MEAN_YELLOW_THRESHOLD},"
            f"${mean_letter}2<{AXIS_SUMMARY_MEAN_RED_THRESHOLD})"
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[red_formula], fill=PatternFill("solid", fgColor="FFC7CE"), stopIfTrue=True),
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(formula=[yellow_formula], fill=PatternFill("solid", fgColor="FFEB9C")),
        )

    def _first_row(
        self,
        rows_by_group: dict[str, list[RowLike]],
        group_id: str,
    ) -> RowLike | None:
        """Return the first raw row for one group."""

        rows = rows_by_group.get(group_id) or []
        return rows[0] if rows else None

    def _openpyxl_image_class(self):
        """Return openpyxl Image class when image support is available."""

        try:
            from openpyxl.drawing.image import Image
        except ImportError as exc:  # pragma: no cover - depends on Pillow availability.
            self._logger.warning("Image embedding unavailable for gallery workbook: %s", exc)
            return None
        return Image

    def _abs_optional(self, value: object) -> float | None:
        """Return an absolute finite float or None."""

        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return None
        if not math.isfinite(numeric):
            return None
        return abs(numeric)

    def _numeric_values_seen(self, values) -> str:
        """Return sorted unique numeric values in a compact audit string."""

        unique_values = sorted(
            {
                float(value)
                for value in values
                if value is not None and self._is_finite_number(value)
            }
        )
        return "; ".join(f"{value:g}" for value in unique_values)

    def _string_values_seen(self, values) -> str:
        """Return sorted unique non-empty strings in a compact audit string."""

        unique_values = sorted({str(value) for value in values if value})
        return "; ".join(unique_values)

    def _values_mixed(self, values) -> bool:
        """Return whether more than one non-empty value appears."""

        return len({str(value) for value in values if value}) > 1

    def _is_finite_number(self, value: object) -> bool:
        """Return whether a value can be represented as a finite float."""

        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return False
        return math.isfinite(numeric)

    def _chart_type(self, stats: StatsLike) -> str:
        """Return the user-facing chart type for a stats row."""

        return getattr(stats, "chart_type", "Motion Duration Distribution") or "Motion Duration Distribution"

    def _action(self, stats: StatsLike) -> str:
        """Return compact user-facing action text."""

        labels = {
            "search_reference": "Search Reference",
            "clear_motor": "Clear Motor",
            "move_to_home": "Move to Home",
            "move_to_max": "Move to Max",
        }
        return labels.get(stats.rule_id, stats.action_label or stats.rule_id)

    def _reference_evidence_status(self, stats: StatsLike) -> str:
        """Return a compact reference evidence status when applicable."""

        if self._chart_type(stats) != "Reference Duration Distribution":
            return ""
        found = getattr(stats, "hardware_reference_evidence_found_count", 0)
        missing = getattr(stats, "hardware_reference_evidence_missing_count", 0)
        if found and not missing:
            return "Found"
        if found and missing:
            return "Mixed"
        return "Missing"

    def _distance_display(self, stats: StatsLike) -> str:
        """Return the chart/statistics distance display."""

        if self._chart_type(stats) == "Reference Duration Distribution":
            return "Not Applicable"
        return self._hardware_distance_group_display(stats)

    def _hardware_distance_group_display(self, stats: StatsLike) -> str:
        """Return a readable grouped hardware distance display."""

        from .distribution import format_movement_distance_group_value

        value = self._abs_optional(getattr(stats, "movement_distance_group_value", None))
        if value is None:
            return ""
        return format_movement_distance_group_value(
            value,
            getattr(stats, "movement_distance_grouping_mode", ""),
            getattr(stats, "movement_distance_bin_size", None),
            getattr(stats, "movement_distance_round_digits", 0),
        )

    def _population_std_s(self, stats: StatsLike) -> float | None:
        """Return population standard deviation in seconds."""

        value = getattr(stats, "population_std_s", None)
        if value is not None:
            return value
        value = getattr(stats, "population_std_ms", None)
        return value / 1000.0 if value is not None else None

    def _population_var_s2(self, stats: StatsLike) -> float | None:
        """Return population variance in seconds squared."""

        value = getattr(stats, "population_var_s2", None)
        if value is not None:
            return value
        value = getattr(stats, "population_var_ms2", None)
        return value / 1_000_000.0 if value is not None else None

    def _min_max_display(self, stats: StatsLike) -> str:
        """Return compact duration min-max text for gallery summary rows."""

        min_s = getattr(stats, "min_s", None)
        max_s = getattr(stats, "max_s", None)
        if min_s is None:
            min_ms = getattr(stats, "min_ms", None)
            min_s = min_ms / 1000.0 if min_ms is not None else None
        if max_s is None:
            max_ms = getattr(stats, "max_ms", None)
            max_s = max_ms / 1000.0 if max_ms is not None else None
        if min_s is None or max_s is None:
            return ""
        return f"{float(min_s):.3f}-{float(max_s):.3f}"

    def _display_min_max(self, value: object) -> object:
        """Return min-max display text with the report-facing en dash."""

        if value is None:
            return None
        return str(value).replace("-", "–")

    def _chart_file_display(self, chart_file: str | Path | None) -> str:
        """Return a user-facing chart image filename."""

        return self._basename(chart_file)

    def _basename(self, value: str | Path | None) -> str:
        """Return filename-only display text so workbooks do not expose local paths."""

        if not value:
            return ""
        return str(value).replace("\\", "/").split("/")[-1]
